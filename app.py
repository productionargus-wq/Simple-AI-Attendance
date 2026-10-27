import os
import sys
import time
import re
import io
import csv
import json
import uuid
import gzip
import secrets
import base64
import hashlib
import threading
import urllib.parse
import requests
import werkzeug.utils
import zipfile
import traceback
from datetime import datetime, timedelta, date, timezone
from dotenv import load_dotenv
from PIL import Image, ImageOps

# Load environment variables (supports frozen PyInstaller bundle or local file)
if getattr(sys, 'frozen', False):
    exe_dir = os.path.dirname(sys.executable)
    load_dotenv(os.path.join(exe_dir, '.env'))
    if hasattr(sys, '_MEIPASS'):
        load_dotenv(os.path.join(sys._MEIPASS, '.env'))
else:
    load_dotenv()

from flask import Flask, render_template, request, jsonify, redirect, url_for, send_file, send_from_directory, Response, session, make_response, flash
from flask.sessions import SecureCookieSessionInterface
from functools import wraps
from flask_cors import CORS
import database

GOOGLE_CLIENT_ID = os.environ.get('GOOGLE_CLIENT_ID', '').strip()
GOOGLE_CLIENT_SECRET = os.environ.get('GOOGLE_CLIENT_SECRET', '').strip()

try:
    import pdf_generator
except Exception as e:
    pdf_generator = None
    print(f"Warning: pdf_generator import failed: {e}")

try:
    import numpy as np
except ImportError:
    np = None

try:
    import cv2
except ImportError:
    cv2 = None

# Lazy import face_engine to avoid numpy/opencv crash on serverless
face_engine = None
def get_face_engine():
    global face_engine
    if face_engine is None:
        try:
            import face_engine as _fe
            face_engine = _fe
        except Exception as e:
            print(f"Warning: face_engine import failed: {e}")
    return face_engine

class PersistentSessionInterface(SecureCookieSessionInterface):
    """
    Guarantees session cookies are persisted to the browser's persistent on-disk cookie store
    across tab and browser restarts by explicitly setting both Max-Age (in seconds) and Expires.
    Supports candidate cookie names and direct DB remember token auto-rehydration seamlessly.
    """
    def open_session(self, app, request):
        s = self.get_signing_serializer(app)
        if s is None:
            return None
        max_age = int(app.permanent_session_lifetime.total_seconds())

        # 1. Try candidate signed session cookies in order
        candidates = [self.get_cookie_name(app), 'session']
        for cname in candidates:
            val = request.cookies.get(cname)
            if val:
                try:
                    data = s.loads(val, max_age=max_age)
                    if data and isinstance(data, dict):
                        return self.session_class(data)
                except Exception:
                    continue

        # 2. Try direct rehydration from persistent DB remember token if cookies were lost or expired
        rem_token = request.cookies.get('argus_remember_token')
        if rem_token:
            try:
                s_doc = database.get_user_session_by_token(rem_token)
                if s_doc:
                    role = s_doc.get('role')
                    sess_data = {
                        '_permanent': True,
                        'role': role,
                        'company_id': s_doc.get('company_id', ''),
                        'company_name': s_doc.get('company_name', ''),
                        'remember_token': rem_token
                    }
                    if role in ['super_admin', 'company_admin']:
                        sess_data['admin_logged_in'] = True
                        sess_data['admin_username'] = s_doc.get('username') or ('Company Admin' if role == 'company_admin' else 'Admin')
                        sess_data['admin_email'] = s_doc.get('email', '')
                    elif role == 'employee':
                        sess_data['employee_logged_in'] = True
                        sess_data['employee_id'] = s_doc.get('employee_id', '')
                        sess_data['employee_name'] = s_doc.get('employee_name', '')
                        sess_data['employee_email'] = s_doc.get('email', '')
                    return self.session_class(sess_data)
            except Exception:
                pass

        return self.session_class()

    def should_set_cookie(self, app, session):
        if session.modified:
            return True
        if session.get('admin_logged_in') or session.get('employee_logged_in') or session.permanent:
            return True
        return super().should_set_cookie(app, session)

    def save_session(self, app, session, response):
        name = self.get_cookie_name(app)
        domain = self.get_cookie_domain(app)
        path = self.get_cookie_path(app)
        secure = self.get_cookie_secure(app)
        samesite = self.get_cookie_samesite(app)
        httponly = self.get_cookie_httponly(app)

        if session.accessed:
            response.vary.add("Cookie")

        if not session:
            if session.modified:
                for cname in [name, 'session']:
                    response.delete_cookie(
                        cname,
                        domain=domain,
                        path=path,
                        secure=secure,
                        samesite=samesite,
                        httponly=httponly,
                    )
                response.vary.add("Cookie")
            return

        if not self.should_set_cookie(app, session):
            return

        is_auth = bool(session.get('admin_logged_in') or session.get('employee_logged_in') or session.permanent)
        expires = (datetime.now(timezone.utc) + timedelta(days=365)) if is_auth else self.get_expiration_time(app, session)
        max_age = int(app.permanent_session_lifetime.total_seconds()) if is_auth else None
        val = self.get_signing_serializer(app).dumps(dict(session))

        # Write to both cookie names to ensure universal cross-tab persistence
        for cname in [name, 'session']:
            response.set_cookie(
                cname,
                val,
                max_age=max_age,
                expires=expires,
                httponly=httponly,
                domain=domain,
                path=path,
                secure=secure,
                samesite=samesite,
            )
        response.vary.add("Cookie")

def finalize_login_session(resp, role, user_info=None):
    """
    Sets up permanent session state and generates a persistent 365-day database-backed
    remember token cookie alongside standard session cookies.
    """
    session.permanent = True
    session.modified = True
    session['logged_in'] = True
    token = database.create_user_session({
        'role': role,
        'company_id': session.get('company_id', ''),
        'company_name': session.get('company_name', ''),
        'username': session.get('admin_username') or session.get('employee_name', ''),
        'email': session.get('admin_email') or session.get('company_email') or session.get('employee_email', ''),
        'employee_id': session.get('employee_id', ''),
        'employee_name': session.get('employee_name', '')
    })
    if token:
        session['remember_token'] = token
        exp = datetime.now(timezone.utc) + timedelta(days=365)
        resp.set_cookie(
            'argus_remember_token',
            token,
            max_age=31536000,
            expires=exp,
            path='/',
            httponly=False,
            samesite='Lax'
        )
    return resp

def _resolve_bundle_path(rel_path):
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, rel_path)
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), rel_path)

app = Flask(
    __name__,
    template_folder=_resolve_bundle_path('templates'),
    static_folder=_resolve_bundle_path('static')
)
CORS(app)
app.secret_key = 'argus-tech-secret-key-2026'

# Session persistence: keep user logged in until explicit logout (365 days)
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=365)
app.config['SESSION_COOKIE_NAME'] = 'argus_session'
app.config['SESSION_COOKIE_HTTPONLY'] = True
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_REFRESH_EACH_REQUEST'] = True
app.session_interface = PersistentSessionInterface()

@app.route('/health')
@app.route('/ping')
def health_check():
    """Ultra-lightweight keep-alive endpoint for uptime monitors and cloud health checks."""
    return jsonify({
        'status': 'healthy',
        'service': 'ARGUS Attendance Engine',
        'timestamp': int(time.time())
    }), 200

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
app.config['UPLOAD_FOLDER'] = os.environ.get('UPLOAD_FOLDER', os.path.join(BASE_DIR, 'static', 'uploads'))
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16 MB max
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 43200  # 12 hours static caching
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp'}

try:
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
except Exception as e:
    print(f"Warning: Could not create upload directory: {e}")

def process_uploaded_photo(file_bytes, original_filename='photo.jpg'):
    """
    Optimizes and prepares photo for dual-storage (disk + permanent MongoDB Atlas base64):
    1. Auto-orients portrait/landscape via EXIF metadata.
    2. Resizes to max 600x600 px using high-quality LANCZOS resampling.
    3. Converts to standard RGB JPEG (~30-50 KB).
    4. Generates base64 data URI for 100% persistent MongoDB storage (never lost on refresh/redeploy).
    5. Saves optimized copy to UPLOAD_FOLDER on disk.
    """
    try:
        img = Image.open(io.BytesIO(file_bytes))
        img = ImageOps.exif_transpose(img)
        if img.mode in ('RGBA', 'P', 'LA'):
            img = img.convert('RGB')
        img.thumbnail((600, 600), Image.Resampling.LANCZOS)
        out_buf = io.BytesIO()
        img.save(out_buf, format='JPEG', quality=85, optimize=True)
        optimized_bytes = out_buf.getvalue()
    except Exception as e:
        print(f"Warning: PIL photo optimization failed, using raw bytes: {e}")
        optimized_bytes = file_bytes

    b64_str = base64.b64encode(optimized_bytes).decode('utf-8')
    photo_data_url = f"data:image/jpeg;base64,{b64_str}"

    clean_name = werkzeug.utils.secure_filename(original_filename) or 'face_capture.jpg'
    if not clean_name.lower().endswith(('.jpg', '.jpeg', '.png', '.webp')):
        clean_name += '.jpg'
    unique_filename = f"{int(database.time.time())}_{clean_name}"

    try:
        upload_dir = app.config['UPLOAD_FOLDER']
        os.makedirs(upload_dir, exist_ok=True)
        file_path = os.path.join(upload_dir, unique_filename)
        with open(file_path, 'wb') as f:
            f.write(optimized_bytes)
    except Exception as io_err:
        print(f"Warning: could not write photo to disk: {io_err}")

    return unique_filename, photo_data_url, optimized_bytes

def _background_startup():
    try:
        fe = get_face_engine()
        if fe:
            fe.auto_sync_stored_employee_embeddings()
    except Exception as e:
        print(f"Warning: background face sync error: {e}")

try:
    database.init_db()
    # Pre-warm face recognition models synchronously before accepting requests
    fe = get_face_engine()
    if fe:
        fe.ensure_models_available()
        fe.get_detector(320, 240)
        fe.get_recognizer()
    threading.Thread(target=_background_startup, daemon=True, name="FaceSyncStartup").start()
    import report_scheduler
    report_scheduler.start_scheduler()

    def _keep_alive_worker():
        """Periodically pings public web service URL on Render every 10 minutes to prevent sleep."""
        import urllib.request
        import urllib.error
        external_url = (os.environ.get('RENDER_EXTERNAL_URL') or os.environ.get('APP_URL') or '').strip()
        if not external_url:
            return
        ping_url = f"{external_url.rstrip('/')}/health"
        print(f"[KeepAlive] Render self-keepalive worker active for: {ping_url}")
        time.sleep(90)
        while True:
            try:
                req = urllib.request.Request(ping_url, headers={'User-Agent': 'ArgusKeepAliveWorker/1.0'})
                with urllib.request.urlopen(req, timeout=15) as resp:
                    pass
            except Exception:
                pass
            time.sleep(600)  # Ping every 10 minutes (safely inside Render's 15m window)

    threading.Thread(target=_keep_alive_worker, daemon=True, name="RenderKeepAlive").start()
except Exception as e:
    print(f"Warning: database/face/scheduler init error: {e}")

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

DOCUMENTS_FOLDER = os.path.join(app.config['UPLOAD_FOLDER'], 'employee_documents')
try:
    os.makedirs(DOCUMENTS_FOLDER, exist_ok=True)
except Exception as e:
    print(f"Warning: Could not create employee_documents directory: {e}")

ALLOWED_DOCUMENT_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp', 'pdf', 'doc', 'docx', 'xls', 'xlsx', 'txt'}

def allowed_document_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_DOCUMENT_EXTENSIONS

def process_uploaded_employee_documents(req_form, req_files, emp_id_prefix='doc'):
    """Processes uploaded employee documents from form data and files."""
    documents = []
    
    # 1. Process indexed rows: doc_title_X and doc_file_X
    try:
        doc_count = int(req_form.get('doc_count', 100))
    except Exception:
        doc_count = 100

    for i in range(doc_count):
        file_key = f'doc_file_{i}'
        title_key = f'doc_title_{i}'
        if file_key in req_files:
            f = req_files[file_key]
            if f and f.filename and allowed_document_file(f.filename):
                title = req_form.get(title_key, '').strip()
                clean_fname = werkzeug.utils.secure_filename(f.filename) or 'document'
                ext = clean_fname.rsplit('.', 1)[1].lower() if '.' in clean_fname else 'dat'
                uniq_name = f"{emp_id_prefix}_{int(time.time())}_{uuid.uuid4().hex[:6]}_{clean_fname}"
                save_path = os.path.join(DOCUMENTS_FOLDER, uniq_name)
                f.save(save_path)
                documents.append({
                    'id': str(uuid.uuid4())[:8],
                    'document_name': title or clean_fname.rsplit('.', 1)[0].replace('_', ' ').title(),
                    'filename': uniq_name,
                    'file_url': f"/static/uploads/employee_documents/{uniq_name}",
                    'file_type': ext,
                    'uploaded_at': database.get_ist_now().strftime('%d/%m/%Y %I:%M %p')
                })

    # 2. Fallback check for list format (only if no indexed files found)
    if not documents:
        list_files = req_files.getlist('doc_files') or req_files.getlist('doc_files[]')
        list_titles = req_form.getlist('doc_titles') or req_form.getlist('doc_titles[]')
        for idx, f in enumerate(list_files):
            if f and f.filename and allowed_document_file(f.filename):
                title = list_titles[idx].strip() if idx < len(list_titles) else ''
                clean_fname = werkzeug.utils.secure_filename(f.filename) or 'document'
                ext = clean_fname.rsplit('.', 1)[1].lower() if '.' in clean_fname else 'dat'
                uniq_name = f"{emp_id_prefix}_{int(time.time())}_{uuid.uuid4().hex[:6]}_{clean_fname}"
                save_path = os.path.join(DOCUMENTS_FOLDER, uniq_name)
                f.save(save_path)
                documents.append({
                    'id': str(uuid.uuid4())[:8],
                    'document_name': title or clean_fname.rsplit('.', 1)[0].replace('_', ' ').title(),
                    'filename': uniq_name,
                    'file_url': f"/static/uploads/employee_documents/{uniq_name}",
                    'file_type': ext,
                    'uploaded_at': database.get_ist_now().strftime('%d/%m/%Y %I:%M %p')
                })

    return documents

@app.route('/uploads/employee_documents/<path:filename>')
def uploaded_employee_document(filename):
    return send_from_directory(DOCUMENTS_FOLDER, filename)

# ----------------- MULTI-TENANT RBAC HELPERS & DECORATORS ----------------- #

def get_current_company_id():
    """Returns the company_id associated with the currently logged-in user."""
    role = session.get('role')
    comp_id = session.get('company_id')
    if role == 'super_admin':
        if request:
            try:
                arg_comp = request.args.get('company_id')
                if arg_comp and str(arg_comp).strip() and str(arg_comp).strip() not in ['None', 'null']:
                    return str(arg_comp).strip()
            except Exception:
                pass
        return comp_id or 'ARGUS_MASTER'

    if role == 'company_admin':
        if comp_id and comp_id != 'ARGUS_MASTER' and database.get_company_by_id(comp_id):
            return comp_id
        # Self-heal from email
        c_email = session.get('company_email') or session.get('admin_email')
        if c_email:
            comp_rec = database.get_company_by_email(c_email)
            if comp_rec and comp_rec.get('id'):
                session['company_id'] = comp_rec['id']
                session['company_name'] = comp_rec.get('company_name', session.get('company_name', ''))
                session.modified = True
                return comp_rec['id']
        # Fallback to first active company
        db = database.get_db()
        first_c = db.company_admin.find_one({'id': {'$nin': ['ALL', 'ARGUS_MASTER']}, 'status': {'$nin': ['Deactive', 'Inactive']}})
        if first_c:
            session['company_id'] = first_c['id']
            session['company_name'] = first_c.get('company_name', '')
            session.modified = True
            return first_c['id']

    return comp_id or 'ARGUS_MASTER'

def get_current_branch_id():
    """Returns active branch_id from request args or session. Returns None if ALL, unassigned, or invalid."""
    bid = request.args.get('branch_id')
    if bid is None:
        bid = session.get('active_branch_id')
    if not bid or str(bid).strip().upper() in ['ALL', 'NONE', '', 'NULL']:
        return None
    return str(bid).strip()

def get_current_company_info():
    """Returns dynamic company dictionary for PDF generation and headers."""
    comp_id = get_current_company_id()
    branch_id = get_current_branch_id()
    branch_doc = None
    if comp_id and branch_id and str(branch_id).strip().upper() not in ['HOME', 'HQ', 'MAIN', 'UNASSIGNED']:
        try:
            branch_doc = database.get_branch_by_id(branch_id, comp_id)
        except Exception:
            branch_doc = None

    if comp_id and comp_id != 'ARGUS_MASTER':
        comp = database.get_company_by_id(comp_id)
        if comp:
            addr = comp.get('address') or comp.get('location') or ''
            if not addr and comp.get('latitude') and comp.get('longitude'):
                addr = f"Location: Lat {comp.get('latitude')}, Lng {comp.get('longitude')}"
            comp_name = comp.get('company_name') or session.get('company_name', 'ARGUS TECHNOLOGIES')
            if branch_doc and branch_doc.get('branch_name'):
                comp_name = f"{comp_name} - {branch_doc.get('branch_name')}"
                if branch_doc.get('location'):
                    addr = branch_doc.get('location')
            return {
                'id': comp.get('id', comp_id),
                'company_name': comp_name,
                'company_address': addr,
                'email': comp.get('email', ''),
                'phone': comp.get('phone', ''),
                'gstin': comp.get('gstin', ''),
                'branch_name': branch_doc.get('branch_name') if branch_doc else '',
                'branch_id': branch_id or ''
            }
    # Master / Default ARGUS info
    master_comp = database.get_company_by_id('ARGUS_MASTER')
    if master_comp:
        return {
            'id': 'ARGUS_MASTER',
            'company_name': master_comp.get('company_name', 'ARGUS TECHNOLOGIES'),
            'company_address': master_comp.get('address', 'SF NO. 515, Bharathiyar Road, Maniyakaranpalayam, Ganapathy (PO), Coimbatore - 641 006'),
            'email': master_comp.get('email', 'technologiesargus@gmail.com'),
            'phone': master_comp.get('phone', '+91 98765 43210'),
            'gstin': master_comp.get('gstin', '33AABCA0000A1Z5')
        }
    return {
        'id': 'ARGUS_MASTER',
        'company_name': 'ARGUS TECHNOLOGIES',
        'company_address': 'SF NO. 515, Bharathiyar Road, Maniyakaranpalayam, Ganapathy (PO), Coimbatore - 641 006',
        'email': 'technologiesargus@gmail.com',
        'phone': '+91 98765 43210',
        'gstin': '33AABCA0000A1Z5'
    }

def login_required(f):
    """Requires Super Admin or Company Admin session."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('admin_logged_in') or session.get('role') not in ['super_admin', 'company_admin']:
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Authentication required. Your session may have expired. Please refresh and log in again.'}), 401
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

def super_admin_required(f):
    """Requires Super Admin session (Platform Owner)."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('admin_logged_in') or session.get('role') != 'super_admin':
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Super Admin privileges required.'}), 403
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

def employee_required(f):
    """Requires Employee self-service session."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('employee_logged_in') or session.get('role') != 'employee':
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Employee authentication required.'}), 401
            return redirect(url_for('employee_login'))
        return f(*args, **kwargs)
    return decorated_function

def check_user_can_delete_entries():
    """Checks whether the currently logged-in user can delete attendance/report entries."""
    if not session.get('admin_logged_in'):
        return False
    if session.get('role') == 'super_admin':
        return True
    if session.get('role') == 'company_admin':
        comp_id = session.get('company_id')
        if comp_id and database.can_company_delete_entries(comp_id):
            return True
    return False

@app.context_processor
def inject_delete_permissions():
    """Injects delete permissions, system admin status, support access, pending leave count, and dynamic company logo into all templates with sub-millisecond cached lookups."""
    role = session.get('role')
    is_super = role == 'super_admin'
    target_comp_id = session.get('company_id') or ('ARGUS_MASTER' if is_super else None)

    # 1. Single cached company lookup (takes <0.1ms from in-memory TTL cache)
    comp_doc = database.get_company_by_id(target_comp_id, include_counts=False) if target_comp_id else None

    # Derive permissions directly from comp_doc
    if is_super:
        comp_can_del = True
        company_support_enabled = True
    else:
        comp_can_del = bool(comp_doc.get('can_delete_entries', False)) if comp_doc else False
        company_support_enabled = bool(comp_doc.get('support_enabled', True)) if comp_doc else False

    active_bid = get_current_branch_id()
    active_bname = session.get('active_branch_name')
    if active_bid:
        if str(active_bid).strip().upper() in ['HOME', 'HQ', 'MAIN', 'UNASSIGNED']:
            active_bid = 'HOME'
            active_bname = 'Home Branch'
            session['active_branch_id'] = 'HOME'
            session['active_branch_name'] = 'Home Branch'
        elif not active_bname or session.get('active_branch_id') != active_bid:
            b_doc = database.get_branch_by_id(active_bid, company_id='ALL')
            if b_doc:
                active_bname = b_doc.get('branch_name', '')
                session['active_branch_id'] = str(active_bid)
                session['active_branch_name'] = active_bname
    else:
        active_bname = None

    # 2. Cached badge counts (served from in-memory cache)
    if is_super:
        pending_cnt = database.get_pending_leave_requests_count(None, branch_id=active_bid)
        admin_support_count = database.get_support_badge_count(None)
        support_open_count = 0
    else:
        pending_cnt = database.get_pending_leave_requests_count(target_comp_id, branch_id=active_bid) if target_comp_id else 0
        support_open_count = database.get_support_badge_count(target_comp_id) if target_comp_id else 0
        admin_support_count = 0

    # 3. Resolve logo dynamically from cached comp_doc
    has_custom_logo = False
    logo_url = None
    if comp_doc:
        if comp_doc.get('logo_data'):
            logo_url = comp_doc['logo_data']
            has_custom_logo = True
        elif comp_doc.get('logo'):
            logo_url = f"/uploads/{comp_doc['logo']}"
            has_custom_logo = True

    if not has_custom_logo:
        logo_url = url_for('static', filename='images/default_profile.svg')

    # 4. Resolve company modules from cached comp_doc
    if is_super:
        company_modules = {'single_company': True, 'branch_company': True, 'field_worker': True}
    elif comp_doc and comp_doc.get('modules'):
        company_modules = comp_doc['modules']
    elif target_comp_id:
        company_modules = database.get_company_modules(target_comp_id)
    else:
        company_modules = {'single_company': True, 'branch_company': False, 'field_worker': False}

    # 5. Fast cached branch list for navbar (skip headcount aggregation for 0.01ms speed)
    branches_list = []
    if target_comp_id and company_modules.get('branch_company', True):
        branches_list = database.get_company_branches(target_comp_id, include_headcount=False)

    # 6. Multi-tenant company list for System Admin tab filters
    system_admin_companies = []
    active_comp_id = target_comp_id or 'ARGUS_MASTER'
    if is_super:
        if request:
            try:
                req_comp = request.args.get('company_id')
                if req_comp and str(req_comp).strip() and str(req_comp).strip() not in ['None', 'null']:
                    active_comp_id = str(req_comp).strip()
            except Exception:
                pass
        try:
            comps_res = database.get_all_companies(limit=1000)
            raw_comps = comps_res.get('data', []) if isinstance(comps_res, dict) else []
            system_admin_companies = [{'id': 'ARGUS_MASTER', 'company_name': 'System Administration'}]
            for c in raw_comps:
                if c.get('id') != 'ARGUS_MASTER':
                    system_admin_companies.append({
                        'id': c['id'],
                        'company_name': c.get('company_name', c['id'])
                    })
        except Exception:
            system_admin_companies = [{'id': 'ARGUS_MASTER', 'company_name': 'System Administration'}]

    return {
        'is_system_admin': is_super,
        'system_admin_companies': system_admin_companies,
        'active_company_id': active_comp_id,
        'can_delete_entries': comp_can_del,
        'company_support_enabled': company_support_enabled,
        'pending_leave_count': pending_cnt,
        'support_open_count': support_open_count,
        'admin_support_count': admin_support_count,
        'current_company_logo_url': logo_url,
        'has_company_logo': has_custom_logo,
        'company_modules': company_modules,
        'total_branches_count': len(branches_list),
        'user_branches': branches_list,
        'active_branch_id': active_bid,
        'active_branch_name': active_bname or ''
    }

import traceback
from werkzeug.exceptions import HTTPException

DEFAULT_AVATAR_SVG = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" width="100%" height="100%">
  <circle cx="12" cy="12" r="12" fill="#f1f5f9"/>
  <path fill="#64748b" fill-rule="evenodd" d="M7.5 6a4.5 4.5 0 1 1 9 0 4.5 4.5 0 0 1-9 0ZM3.751 20.105a8.25 8.25 0 0 1 16.498 0 .75.75 0 0 1-.437.695A18.683 18.683 0 0 1 12 22.5c-2.786 0-5.433-.608-7.812-1.7a.75.75 0 0 1-.437-.695Z" clip-rule="evenodd"/>
</svg>'''

@app.route('/favicon.ico')
def favicon():
    static_dir = os.path.join(app.root_path, 'static')
    if os.path.exists(os.path.join(static_dir, 'favicon.ico')):
        resp = make_response(send_from_directory(static_dir, 'favicon.ico', mimetype='image/x-icon'))
        resp.headers['Cache-Control'] = 'no-cache, must-revalidate'
        return resp
    icon_dir = os.path.join(static_dir, 'images')
    if os.path.exists(os.path.join(icon_dir, 'favicon.ico')):
        resp = make_response(send_from_directory(icon_dir, 'favicon.ico', mimetype='image/x-icon'))
        resp.headers['Cache-Control'] = 'no-cache, must-revalidate'
        return resp
    if os.path.exists(os.path.join(icon_dir, 'argus_triangle_logo.png')):
        return send_from_directory(icon_dir, 'argus_triangle_logo.png', mimetype='image/png')
    return Response(status=204)

@app.errorhandler(500)
def handle_500(e):
    err = traceback.format_exc()
    print("500 Internal Error:", err)
    req_path = getattr(request, 'path', '') or ''
    if req_path.startswith('/api/face/recognize') or getattr(request, 'endpoint', None) == 'api_face_recognize':
        return jsonify({
            'matched': False,
            'confidence': 0.0,
            'message': 'Face recognition processing temporarily busy. Please wait a moment and try again.'
        }), 200
    if req_path.startswith('/api/'):
        return jsonify({
            'error': 'Internal server error',
            'message': 'A temporary server processing error occurred. Please try again.',
            'details': str(e)
        }), 500
    return f"<h1>Internal Server Error (500)</h1><pre>{err}</pre>", 500

@app.errorhandler(Exception)
def handle_exception(e):
    req_path = getattr(request, 'path', '') or ''
    if req_path.startswith('/api/face/recognize') or getattr(request, 'endpoint', None) == 'api_face_recognize':
        return jsonify({
            'matched': False,
            'confidence': 0.0,
            'message': 'Face recognition processing temporarily busy. Please wait a moment and try again.'
        }), 200
    if isinstance(e, HTTPException):
        if req_path.startswith('/api/'):
            return jsonify({'error': e.description, 'message': e.description}), e.code
        return e
    err = traceback.format_exc()
    print("Unhandled Exception:", err)
    if req_path.startswith('/api/'):
        return jsonify({
            'error': 'Server error',
            'message': 'A temporary server processing error occurred. Please try again.',
            'details': str(e)
        }), 500
    return f"<h1>Server Error</h1><pre>{err}</pre>", 500

@app.after_request
def add_performance_headers(response):
    if request.path.startswith('/static/'):
        if request.args.get('v'):
            response.headers['Cache-Control'] = 'public, max-age=31536000, immutable'
        else:
            response.headers['Cache-Control'] = 'no-cache, must-revalidate'
    elif response.status_code == 200 and request.method == 'GET' and not request.path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-cache, must-revalidate'

    accept_encoding = request.headers.get('Accept-Encoding', '')
    if (
        'gzip' in accept_encoding.lower()
        and response.status_code == 200
        and not response.direct_passthrough
        and response.mimetype in ['text/html', 'text/css', 'application/javascript', 'application/json', 'image/svg+xml']
    ):
        try:
            data = response.get_data()
            if len(data) > 1024:
                gzip_buffer = io.BytesIO()
                with gzip.GzipFile(mode='wb', fileobj=gzip_buffer, compresslevel=6) as gz:
                    gz.write(data)
                response.set_data(gzip_buffer.getvalue())
                response.headers['Content-Encoding'] = 'gzip'
                response.headers['Content-Length'] = len(response.get_data())
                response.headers['Vary'] = 'Accept-Encoding'
        except Exception:
            pass
    return response

@app.before_request
def ensure_persistent_session():
    """Keeps the session alive for 365 days across browser restarts, and evicts deactivated tenants."""
    # 1. Loopback Host Normalization: Force 127.0.0.1 -> localhost so cookies are 100% shared across tabs
    if request.host.startswith('127.0.0.1'):
        port = f":{request.host.split(':')[1]}" if ':' in request.host else ''
        new_url = request.url.replace(f"127.0.0.1{port}", f"localhost{port}", 1)
        return redirect(new_url, code=307)

    # 2. Session Auto-Rehydration from Persistent DB Remember Token
    if not (session.get('admin_logged_in') or session.get('employee_logged_in')):
        rem_token = request.cookies.get('argus_remember_token')
        if rem_token:
            s_doc = database.get_user_session_by_token(rem_token)
            if s_doc:
                session.permanent = True
                session.modified = True
                role = s_doc.get('role')
                session['role'] = role
                session['company_id'] = s_doc.get('company_id')
                session['company_name'] = s_doc.get('company_name')
                session['remember_token'] = rem_token
                if role in ['super_admin', 'company_admin', 'admin']:
                    session['admin_logged_in'] = True
                    session['logged_in'] = True
                    session['admin_username'] = s_doc.get('username') or ('Company Admin' if role == 'company_admin' else 'Admin')
                    session['admin_email'] = s_doc.get('email', '')
                    if role == 'company_admin':
                        session['company_email'] = s_doc.get('email', '')
                elif role == 'employee':
                    session['employee_logged_in'] = True
                    session['employee_id'] = s_doc.get('employee_id')
                    session['employee_name'] = s_doc.get('employee_name')
                    session['employee_email'] = s_doc.get('email', '')

    if session.get('admin_logged_in') or session.get('employee_logged_in'):
        session.permanent = True
        session.modified = True

        role = session.get('role')

        # Self-healing company_id for company_admin
        if role == 'company_admin':
            comp_id = session.get('company_id')
            if not comp_id or not database.get_company_by_id(comp_id):
                c_email = session.get('company_email') or session.get('admin_email')
                comp_rec = database.get_company_by_email(c_email) if c_email else None
                if not comp_rec:
                    db = database.get_db()
                    raw_c = db.company_admin.find_one({'id': {'$nin': ['ALL', 'ARGUS_MASTER']}, 'status': {'$nin': ['Deactive', 'Inactive']}})
                    if raw_c:
                        comp_rec = database.clean_doc(raw_c)
                if comp_rec and comp_rec.get('id'):
                    session['company_id'] = comp_rec['id']
                    session['company_name'] = comp_rec.get('company_name', session.get('company_name', ''))
                    token = session.get('remember_token')
                    if token:
                        db = database.get_db()
                        db.user_sessions.update_one({'token': token}, {'$set': {'company_id': comp_rec['id'], 'company_name': session['company_name']}})

        if role in ['company_admin', 'employee']:
            comp_id = session.get('company_id')
            if comp_id and str(comp_id).strip() not in ['ALL', 'ARGUS_MASTER']:
                # Allow static assets and explicit logout without interception
                if not request.path.startswith('/static') and request.endpoint not in ['static', 'logout', 'favicon']:
                    if not database.is_company_active(comp_id):
                        token = session.get('remember_token')
                        if token:
                            database.delete_user_session(token)
                        session.clear()
                        if request.path.startswith('/api/'):
                            return jsonify({
                                'success': False,
                                'error': 'COMPANY_DEACTIVATED',
                                'message': 'This company account has been deactivated. Please contact System Administrator.'
                            }), 403
                        flash('Your company account has been deactivated. Please contact System Administrator.', 'danger')
                        return redirect(url_for('login'))

# ----------------- UNIVERSAL EXCEL & CSV EXPORT HELPERS ----------------- #

def format_export_cell(val):
    """Formats cell values cleanly: dates to readable IST, numbers preserved, ISO timestamps parsed."""
    if val is None:
        return ''
    if isinstance(val, (int, float)):
        return val
    if isinstance(val, (datetime, date)):
        if isinstance(val, datetime):
            return val.strftime('%d/%m/%Y %I:%M %p')
        return val.strftime('%d/%m/%Y')
    val_str = str(val).strip()
    # Check if ISO timestamp like 2026-10-05T06:49:21.866000
    if re.match(r'^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}', val_str):
        try:
            clean_iso = val_str.replace('Z', '').split('+')[0]
            dt = datetime.fromisoformat(clean_iso)
            return dt.strftime('%d/%m/%Y %I:%M %p')
        except Exception:
            return val_str
    # Check if date like 2026-10-05
    if re.match(r'^\d{4}-\d{2}-\d{2}$', val_str):
        try:
            dt = datetime.strptime(val_str, '%Y-%m-%d')
            return dt.strftime('%d/%m/%Y')
        except Exception:
            return val_str
    return val_str

def export_table_data(headers, rows, filename_base, export_format='excel', sheet_name='Data'):
    """
    Universal table export handler:
    - If export_format is 'excel': Generates an openpyxl .xlsx workbook with dynamic column widths
      (auto-expanded so dates/times NEVER show '########'), professional header styling, and cell borders.
    - If export_format is 'csv': Generates UTF-8 with BOM (.csv) for Excel compatibility.
    - Bulletproof fallback: If openpyxl encounters any error or missing dependency, gracefully
      returns a clean CSV response so the client NEVER receives an HTTP 500 error.
    """
    clean_rows = []
    for row in rows:
        clean_rows.append([format_export_cell(v) for v in row])

    is_csv = str(export_format).strip().lower() == 'csv'

    if is_csv:
        output = io.StringIO()
        # UTF-8 BOM (\ufeff) ensures Excel opens special symbols (e.g. ₹, dashes) cleanly
        output.write('\ufeff')
        writer = csv.writer(output)
        writer.writerow(headers)
        for r in clean_rows:
            writer.writerow(r)
        output.seek(0)
        filename = f"{filename_base}.csv" if not filename_base.endswith('.csv') else filename_base
        return Response(
            output.getvalue(),
            mimetype="text/csv; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

    # Generate real styled .xlsx
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = str(sheet_name)[:31]
        try:
            ws.views.sheetView[0].showGridLines = True
        except Exception:
            try:
                ws.sheet_view.showGridLines = True
            except Exception:
                pass

        # Header styling (professional dark navy/slate fill, bold white font)
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        header_border = Border(
            left=Side(style="thin", color="0F172A"),
            right=Side(style="thin", color="0F172A"),
            top=Side(style="thin", color="0F172A"),
            bottom=Side(style="medium", color="0F172A")
        )
        cell_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0")
        )
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        ws.append(headers)
        ws.row_dimensions[1].height = 28

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = header_border

        data_font = Font(name="Calibri", size=10)
        for row_num, row_data in enumerate(clean_rows, 2):
            ws.row_dimensions[row_num].height = 20
            is_even = (row_num % 2 == 0)
            for col_num, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.font = data_font
                cell.border = cell_border
                if not is_even:
                    cell.fill = zebra_fill
                # Align numbers to right, dates/status/short codes center, strings left
                val_str = str(val if val is not None else '')
                if isinstance(val, (int, float)) or (re.match(r'^-?[\d,]+(\.\d+)?$', val_str) and not val_str.startswith('0') and len(val_str) < 15):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif any(k in (headers[col_num - 1] if col_num - 1 < len(headers) else '').lower() for k in ['date', 'time', 'status', 'sl no', 'id']):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths with extra padding (min 16, max 55) so dates/times NEVER show '#####'
        for col_idx, col in enumerate(ws.columns, 1):
            col_letter = get_column_letter(col_idx)
            header_text = headers[col_idx - 1] if col_idx - 1 < len(headers) else ''
            max_len = len(str(header_text))
            for cell in col:
                v_str = str(cell.value if cell.value is not None else '')
                if '\n' in v_str:
                    v_str = max(v_str.split('\n'), key=len)
                max_len = max(max_len, len(v_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 5, 16), 55)

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        filename = f"{filename_base}.xlsx" if not filename_base.endswith('.xlsx') else filename_base
        return Response(
            out.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except Exception as exc:
        print(f"Warning: openpyxl export failed ({exc}), falling back to CSV export: {traceback.format_exc()}")
        output = io.StringIO()
        output.write('\ufeff')
        writer = csv.writer(output)
        writer.writerow(headers)
        for r in clean_rows:
            writer.writerow(r)
        output.seek(0)
        filename = f"{filename_base}.csv" if not filename_base.endswith('.csv') else filename_base
        return Response(
            output.getvalue(),
            mimetype="text/csv; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )


# ----------------- AUTHENTICATION & PAGE ROUTES ----------------- #

@app.route('/')
def index():
    if not request.args.get('kiosk'):
        if session.get('admin_logged_in'):
            return redirect(url_for('dashboard'))
        if session.get('employee_logged_in'):
            return redirect(url_for('employee_portal'))
    return render_template('attendance.html')

@app.route('/kiosk')
def kiosk():
    """Direct alias to Attendance Kiosk interface."""
    return render_template('attendance.html')

@app.route('/login', methods=['GET', 'POST'])
def login():
    """Super Admin Login (Registered Email: technologiesargus@gmail.com)."""
    if request.method == 'GET' and session.get('admin_logged_in'):
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        email = (request.form.get('email') or request.form.get('username') or '').strip()
        password = request.form.get('password', '').strip()
        if not email:
            return render_template('login.html', error='Please enter your authorized email address.')
        if not password:
            return render_template(
                'login.html',
                error='Password is required for manual sign-in. First time signing in? Please use Continue with Google below.'
            )
        
        result = database.validate_admin_login(email, password=password)
        if isinstance(result, dict) and result.get('success'):
            admin_user = result['admin']
            session.clear()
            session.permanent = True
            session['admin_logged_in'] = True
            session['logged_in'] = True
            session['admin_username'] = admin_user.get('username', 'Admin')
            session['admin_email'] = admin_user.get('email', 'technologiesargus@gmail.com')
            session['role'] = 'super_admin'
            session['company_id'] = 'ARGUS_MASTER'
            session['company_name'] = 'ARGUS TECHNOLOGIES'
            resp = redirect(url_for('dashboard'))
            return finalize_login_session(resp, 'super_admin', admin_user)
        else:
            err_code = result.get('error') if isinstance(result, dict) else None
            if err_code == 'PASSWORD_NOT_SET':
                error_msg = "Password has not been set yet. First-time access? Please click 'Continue with Google' below to verify your account and set your password."
            elif err_code == 'NOT_REGISTERED':
                error_msg = 'Access Denied: Email is not authorized as System Admin. Please verify with Google below.'
            else:
                error_msg = 'Invalid password. Please verify your password or sign in with Google.'
            return render_template('login.html', error=error_msg)
    return render_template('login.html')

@app.route('/auth/google/login')
def google_login():
    """Initiates Google OAuth 2.0 Authorization Flow."""
    login_type = request.args.get('type', 'admin').strip().lower()
    if login_type not in ['admin', 'company', 'employee']:
        login_type = 'admin'

    next_url = (request.args.get('next') or '').strip()
    if next_url.startswith('/'):
        session['oauth_next_url'] = next_url
    else:
        session.pop('oauth_next_url', None)

    state = secrets.token_urlsafe(32)
    session['oauth_state'] = state
    session['oauth_login_type'] = login_type

    redirect_uri = url_for('google_callback', _external=True)
    if request.headers.get('X-Forwarded-Proto') == 'https' and redirect_uri.startswith('http://'):
        redirect_uri = redirect_uri.replace('http://', 'https://', 1)

    session['oauth_redirect_uri'] = redirect_uri

    params = {
        'client_id': GOOGLE_CLIENT_ID,
        'redirect_uri': redirect_uri,
        'response_type': 'code',
        'scope': 'openid email profile',
        'state': state,
        'access_type': 'online',
        'prompt': 'select_account'
    }
    auth_url = 'https://accounts.google.com/o/oauth2/v2/auth?' + urllib.parse.urlencode(params)
    return redirect(auth_url)

@app.route('/auth/google/callback')
def google_callback():
    """Handles Google OAuth 2.0 Callback and role-based authorization."""
    error = request.args.get('error')
    login_type = session.get('oauth_login_type', 'admin')
    oauth_next = session.pop('oauth_next_url', None)

    def render_oauth_error(msg):
        if oauth_next == '/field-worker':
            return render_template('field_worker_login.html', error=msg, selected_role='field_worker' if login_type == 'employee' else 'company_admin')
        if login_type == 'company':
            return render_template('company_login.html', error=msg)
        elif login_type == 'employee':
            return render_template('employee_login.html', error=msg)
        else:
            return render_template('login.html', error=msg)

    if error:
        return render_oauth_error(f"Google sign-in was cancelled or failed ({error}).")

    code = request.args.get('code')
    state = request.args.get('state')
    expected_state = session.get('oauth_state')

    if not code or not state or state != expected_state:
        return render_oauth_error("Invalid or expired OAuth state. Please try logging in again.")

    redirect_uri = session.get('oauth_redirect_uri') or url_for('google_callback', _external=True)
    if request.headers.get('X-Forwarded-Proto') == 'https' and redirect_uri.startswith('http://'):
        redirect_uri = redirect_uri.replace('http://', 'https://', 1)

    token_url = 'https://oauth2.googleapis.com/token'
    token_payload = {
        'code': code,
        'client_id': GOOGLE_CLIENT_ID,
        'client_secret': GOOGLE_CLIENT_SECRET,
        'redirect_uri': redirect_uri,
        'grant_type': 'authorization_code'
    }

    try:
        token_resp = requests.post(token_url, data=token_payload, timeout=12)
        token_data = token_resp.json()
        if 'error' in token_data:
            err_desc = token_data.get('error_description') or token_data.get('error')
            return render_oauth_error(f"Google authorization error: {err_desc}")

        access_token = token_data.get('access_token')
        userinfo_resp = requests.get(
            'https://www.googleapis.com/oauth2/v3/userinfo',
            headers={'Authorization': f"Bearer {access_token}"},
            timeout=12
        )
        user_info = userinfo_resp.json()
        google_email = (user_info.get('email') or '').strip().lower()
        if not google_email:
            return render_oauth_error("Unable to obtain verified email address from Google.")
    except Exception as exc:
        return render_oauth_error(f"Communication error with Google services: {str(exc)}")

    # Clear OAuth temporary session tokens
    session.pop('oauth_state', None)
    session.pop('oauth_redirect_uri', None)
    session.pop('oauth_login_type', None)

    # Helper login session finalizers
    def _do_admin_login(adm_user):
        session.clear()
        session.permanent = True
        session['admin_logged_in'] = True
        session['admin_username'] = adm_user.get('username', 'Admin')
        session['admin_email'] = adm_user.get('email', 'technologiesargus@gmail.com')
        session['role'] = 'super_admin'
        session['company_id'] = 'ARGUS_MASTER'
        session['company_name'] = 'ARGUS TECHNOLOGIES'
        resp = redirect(oauth_next if oauth_next == '/field-worker' else url_for('dashboard'))
        return finalize_login_session(resp, 'super_admin', adm_user)

    def _do_company_login(comp):
        session.clear()
        session.permanent = True
        session['admin_logged_in'] = True
        session['admin_username'] = comp.get('company_name', 'Company Admin')
        session['role'] = 'company_admin'
        session['company_id'] = str(comp.get('id') or comp.get('_id', ''))
        session['company_name'] = comp.get('company_name', '')
        session['company_email'] = comp.get('email', '')
        resp = redirect(oauth_next if oauth_next == '/field-worker' else url_for('dashboard'))
        return finalize_login_session(resp, 'company_admin', comp)

    def _do_employee_login(emp):
        session.clear()
        session.permanent = True
        session['employee_logged_in'] = True
        session['role'] = 'employee'
        session['employee_id'] = str(emp.get('id', emp.get('_id', '')))
        session['employee_name'] = emp.get('employee_name', '')
        session['employee_email'] = emp.get('email_id', '')
        c_id = emp.get('company_id', 'ARGUS_MASTER')
        session['company_id'] = c_id
        c_rec = database.get_company_by_id(c_id) if c_id != 'ARGUS_MASTER' else None
        session['company_name'] = c_rec.get('company_name', 'ARGUS TECHNOLOGIES') if c_rec else 'ARGUS TECHNOLOGIES'
        resp = redirect(oauth_next if oauth_next == '/field-worker' else url_for('employee_portal'))
        return finalize_login_session(resp, 'employee', emp)

    # Authorized Role Verification with cross-portal smart role detection
    if login_type == 'admin':
        admin_user = database.validate_admin_login(google_email)
        if admin_user:
            return _do_admin_login(admin_user)

        # Fallback 1: Check if registered as Company Administrator
        company = database.validate_company_login(google_email)
        if isinstance(company, dict) and company.get('error') == 'COMPANY_DEACTIVATED':
            return render_oauth_error("Access Denied: This company account has been deactivated. Please contact System Administrator.")
        elif company and isinstance(company, dict) and (company.get('id') or company.get('_id')):
            return _do_company_login(company)

        # Fallback 2: Check if registered as Employee
        employee = database.validate_employee_login(google_email)
        if isinstance(employee, dict) and employee.get('error') == 'COMPANY_DEACTIVATED':
            return render_oauth_error("Access Denied: Your company account has been deactivated. Please contact System Administrator.")
        elif isinstance(employee, dict) and employee.get('error') == 'EMPLOYEE_INACTIVE':
            return render_oauth_error("Access Denied: Your employee account is currently set to Inactive. Please contact your company administrator.")
        elif employee and isinstance(employee, dict) and (employee.get('id') or employee.get('_id')):
            return _do_employee_login(employee)

        return render_oauth_error(
            f"Access Denied: The Google account '{google_email}' is not authorized as System Admin or registered company."
        )

    elif login_type == 'company':
        company = database.validate_company_login(google_email)
        if isinstance(company, dict) and company.get('error') == 'COMPANY_DEACTIVATED':
            return render_oauth_error("Access Denied: This company account has been deactivated. Please contact System Administrator.")
        elif company and isinstance(company, dict) and (company.get('id') or company.get('_id')):
            return _do_company_login(company)

        # Fallback 1: Check if registered as Super Admin
        admin_user = database.validate_admin_login(google_email)
        if admin_user:
            return _do_admin_login(admin_user)

        # Fallback 2: Check if registered as Employee
        employee = database.validate_employee_login(google_email)
        if isinstance(employee, dict) and employee.get('error') == 'COMPANY_DEACTIVATED':
            return render_oauth_error("Access Denied: Your company account has been deactivated. Please contact System Administrator.")
        elif isinstance(employee, dict) and employee.get('error') == 'EMPLOYEE_INACTIVE':
            return render_oauth_error("Access Denied: Your employee account is currently set to Inactive. Please contact your company administrator.")
        elif employee and isinstance(employee, dict) and (employee.get('id') or employee.get('_id')):
            return _do_employee_login(employee)

        return render_oauth_error(
            f"Access Denied: The Google account '{google_email}' is not registered as a company administrator. Please contact Argus Support."
        )

    elif login_type == 'employee':
        employee = database.validate_employee_login(google_email)
        if isinstance(employee, dict) and employee.get('error') == 'COMPANY_DEACTIVATED':
            return render_oauth_error("Access Denied: Your company account has been deactivated. Please contact System Administrator.")
        elif isinstance(employee, dict) and employee.get('error') == 'EMPLOYEE_INACTIVE':
            return render_oauth_error("Access Denied: Your employee account is currently set to Inactive. Please contact your company administrator.")
        elif employee and isinstance(employee, dict) and (employee.get('id') or employee.get('_id')):
            return _do_employee_login(employee)

        # Fallback 1: Check if registered as Company Administrator
        company = database.validate_company_login(google_email)
        if isinstance(company, dict) and company.get('error') == 'COMPANY_DEACTIVATED':
            return render_oauth_error("Access Denied: This company account has been deactivated. Please contact System Administrator.")
        elif company and isinstance(company, dict) and (company.get('id') or company.get('_id')):
            return _do_company_login(company)

        # Fallback 2: Check if registered as Super Admin
        admin_user = database.validate_admin_login(google_email)
        if admin_user:
            return _do_admin_login(admin_user)

        return render_oauth_error(
            f"Access Denied: The Google account '{google_email}' is not registered with any organisation. Please contact your company HR."
        )

    return redirect(url_for('login'))

@app.route('/company-login', methods=['GET', 'POST'])
def company_login():
    """Dedicated Company Admin Portal Login via Registered Corporate Email."""
    if request.method == 'GET' and session.get('admin_logged_in'):
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        email = request.form.get('email', '').strip()
        password = request.form.get('password', '').strip()
        if not email:
            return render_template('company_login.html', error='Please enter your registered corporate email.')
        if not password:
            return render_template(
                'company_login.html',
                error='Password is required for manual sign-in. First time signing in? Please use Continue with Google below.'
            )

        result = database.validate_company_login(email, password=password)
        if isinstance(result, dict) and result.get('success'):
            company = result['company']
            session.clear()
            session.permanent = True
            session['admin_logged_in'] = True
            session['logged_in'] = True
            session['admin_username'] = company.get('company_name', 'Company Admin')
            session['role'] = 'company_admin'
            session['company_id'] = str(company.get('id') or company.get('_id', ''))
            session['company_name'] = company.get('company_name', '')
            session['company_email'] = company.get('email', '')
            resp = redirect(url_for('dashboard'))
            return finalize_login_session(resp, 'company_admin', company)
        else:
            err_code = result.get('error') if isinstance(result, dict) else None
            if err_code == 'COMPANY_DEACTIVATED':
                error_msg = result.get('message') or 'Access Denied: This company account has been deactivated. Please contact System Administrator.'
            elif err_code == 'PASSWORD_NOT_SET':
                error_msg = "Password has not been set yet. First-time access? Please click 'Continue with Google' below to verify your account and set your password in Company Profile."
            elif err_code == 'NOT_REGISTERED':
                error_msg = 'Access Denied: This email is not registered as a company administrator. Please contact Argus Support.'
            else:
                error_msg = 'Invalid password. Please check your credentials or continue with Google.'
            return render_template('company_login.html', error=error_msg)
    return render_template('company_login.html')

@app.route('/employee-login', methods=['GET', 'POST'])
def employee_login():
    """Dedicated Employee Portal Login via Registered Employee Email."""
    if request.method == 'GET' and session.get('employee_logged_in'):
        return redirect(url_for('employee_portal'))
    if request.method == 'POST':
        email = request.form.get('email', '').strip()
        password = request.form.get('password', '').strip()
        if not email:
            return render_template('employee_login.html', error='Please enter your registered employee email.')
        if not password:
            return render_template(
                'employee_login.html',
                error='Password is required for manual sign-in. First time signing in? Please use Continue with Google below.'
            )

        result = database.validate_employee_login(email, password=password)
        if isinstance(result, dict) and result.get('success'):
            employee = result['employee']
            is_fw = bool(
                employee.get('is_field_worker') is True
                or str(employee.get('employee_type') or '').lower() in ['field_worker', 'field', 'fw']
            )
            if is_fw:
                return render_template(
                    'employee_login.html',
                    error='Access Denied: This account is registered as a Field Worker. Please sign in via the Field Worker Module login.',
                    form_data={'email': email}
                )
            session.clear()
            session.permanent = True
            session['employee_logged_in'] = True
            session['role'] = 'employee'
            session['employee_id'] = str(employee.get('id', employee.get('_id', '')))
            session['employee_name'] = employee.get('employee_name', '')
            session['employee_email'] = employee.get('email_id', '')
            comp_id = employee.get('company_id', 'ARGUS_MASTER')
            session['company_id'] = comp_id
            comp_record = database.get_company_by_id(comp_id) if comp_id != 'ARGUS_MASTER' else None
            session['company_name'] = comp_record.get('company_name', 'ARGUS TECHNOLOGIES') if comp_record else 'ARGUS TECHNOLOGIES'
            resp = redirect(url_for('employee_portal'))
            return finalize_login_session(resp, 'employee', employee)
        else:
            err_code = result.get('error') if isinstance(result, dict) else None
            if err_code == 'COMPANY_DEACTIVATED':
                error_msg = result.get('message') or 'Access Denied: Your company account has been deactivated. Please contact System Administrator.'
            elif err_code == 'EMPLOYEE_INACTIVE':
                error_msg = result.get('message') or 'Access Denied: Your employee account is currently set to Inactive. Please contact your company administrator.'
            elif err_code == 'PASSWORD_NOT_SET':
                error_msg = "Password has not been set yet. First-time access? Please click 'Continue with Google' below to verify your account and set your password in My Credentials."
            elif err_code == 'NOT_REGISTERED':
                error_msg = 'Access Denied: Email not registered with any organisation. Please contact your company HR.'
            else:
                error_msg = 'Invalid password. Please check your credentials or continue with Google.'
            return render_template('employee_login.html', error=error_msg)
    return render_template('employee_login.html')

@app.route('/logout')
def logout():
    role = session.get('role')
    token = session.get('remember_token') or request.cookies.get('argus_remember_token')
    if token:
        database.delete_user_session(token)
    session.clear()
    cookie_name = app.config.get('SESSION_COOKIE_NAME', 'argus_session')
    target = url_for('company_login') if role == 'company_admin' else (url_for('employee_login') if role == 'employee' else url_for('login'))
    resp = redirect(target)
    resp.delete_cookie(cookie_name, path='/')
    resp.delete_cookie('session', path='/')
    resp.delete_cookie('argus_remember_token', path='/')
    return resp

@app.route('/api/auth/restore-session', methods=['POST'])
def api_restore_session():
    """Client-side fallback: Re-hydrates an authenticated session from localStorage remember token."""
    data = request.get_json(silent=True) or {}
    token = (data.get('token') or request.cookies.get('argus_remember_token') or '').strip()
    if not token:
        return jsonify({'success': False, 'error': 'NO_TOKEN'}), 401

    s_doc = database.get_user_session_by_token(token)
    if not s_doc:
        return jsonify({'success': False, 'error': 'INVALID_TOKEN'}), 401

    session.permanent = True
    session.modified = True
    role = s_doc.get('role', 'company_admin')
    session['role'] = role
    session['company_id'] = s_doc.get('company_id', '')
    session['company_name'] = s_doc.get('company_name', '')
    session['remember_token'] = token

    if role in ['super_admin', 'company_admin']:
        session['admin_logged_in'] = True
        session['admin_username'] = s_doc.get('username') or ('Company Admin' if role == 'company_admin' else 'Admin')
        session['admin_email'] = s_doc.get('email', '')
        target = url_for('dashboard')
    elif role == 'employee':
        session['employee_logged_in'] = True
        session['employee_id'] = s_doc.get('employee_id', '')
        session['employee_name'] = s_doc.get('employee_name', '')
        session['employee_email'] = s_doc.get('email', '')
    else:
        target = url_for('dashboard')

    resp = jsonify({'success': True, 'role': role, 'redirect': target, 'redirect_url': target})
    exp = datetime.now(timezone.utc) + timedelta(days=365)
    resp.set_cookie(
        'argus_remember_token',
        token,
        max_age=31536000,
        expires=exp,
        path='/',
        httponly=False,
        samesite='Lax'
    )
    return resp

@app.route('/manage-companies')
@super_admin_required
def manage_companies():
    return render_template('manage_companies.html', active_tab='MANAGE COMPANIES')

@app.route('/company-profile')
@login_required
def company_profile():
    role = session.get('role', 'company_admin')
    return render_template('company_profile.html', active_tab='COMPANY PROFILE', role=role)

@app.route('/leave-permission')
@login_required
def leave_permission():
    comp_id = get_current_company_id()
    branch_id = get_current_branch_id()
    policy = database.get_company_leave_policy(comp_id)
    stats = database.get_admin_leave_stats(company_id=comp_id, branch_id=branch_id)
    return render_template('leave_permission.html', active_tab='LEAVE & PERMISSION', policy=policy, stats=stats)

@app.route('/leave-permission/review/<request_id>')
@login_required
def leave_permission_review(request_id):
    req = database.get_leave_request_by_id(request_id)
    if not req:
        flash('Leave request not found', 'danger')
        return redirect(url_for('leave_permission'))
    emp_id = req.get('employee_id')
    comp_id = req.get('company_id')
    emp = database.get_employee_by_id(emp_id)
    bal = database.get_employee_leave_balance(emp_id, company_id=comp_id)
    return render_template('leave_review.html', active_tab='LEAVE & PERMISSION', req=req, employee=emp or {}, balance=bal or {})

@app.route('/employee/portal')
@employee_required
def employee_portal():
    emp_id = session.get('employee_id')
    comp_id = session.get('company_id')
    employee = database.get_employee_by_id(emp_id)
    company = database.get_company_by_id(comp_id) if comp_id != 'ARGUS_MASTER' else {'company_name': 'Argus Technologies'}
    if company and company.get('company_name'):
        session['company_name'] = company['company_name']
    return render_template('employee_portal.html', employee=employee, company=company)

@app.route('/dashboard')
@login_required
def dashboard():
    comp_id = get_current_company_id()
    if comp_id and comp_id != 'ARGUS_MASTER':
        comp = database.get_company_by_id(comp_id)
        if comp and comp.get('company_name'):
            session['company_name'] = comp['company_name']
    # Resolve multi-branch context if company has branches
    branches = []
    active_bid = request.args.get('branch_id')
    if active_bid is None:
        active_bid = session.get('active_branch_id')
    if not active_bid or str(active_bid).strip().upper() in ['ALL', 'NONE', '']:
        active_bid = None

    active_branch = None
    if comp_id:
        try:
            branches = database.get_company_branches(comp_id)
            if active_bid:
                if str(active_bid).strip().upper() in ['HOME', 'HQ', 'MAIN', 'UNASSIGNED']:
                    active_bid = 'HOME'
                    active_branch = {'id': 'HOME', 'branch_name': 'Home Branch'}
                else:
                    active_branch = database.get_branch_by_id(active_bid, comp_id)
                    if not active_branch:
                        active_bid = None
        except Exception as e:
            print(f"Error resolving branches for dashboard: {e}")
            branches = []
            active_bid = None

    session['active_branch_id'] = active_bid
    if active_branch:
        session['active_branch_name'] = active_branch.get('branch_name')
    else:
        session.pop('active_branch_name', None)

    stats = database.get_dashboard_stats(company_id=comp_id, branch_id=active_bid)

    try:
        from report_scheduler import get_yesterday_ist, get_previous_month_ist
        yesterday_str = get_yesterday_ist()
        prev_month_str = get_previous_month_ist()

        try:
            y_dt = datetime.strptime(yesterday_str, '%Y-%m-%d')
            yesterday_display = y_dt.strftime('%d %b %Y')
        except Exception:
            yesterday_display = yesterday_str

        try:
            m_dt = datetime.strptime(prev_month_str, '%Y-%m')
            prev_month_display = m_dt.strftime('%B %Y')
        except Exception:
            prev_month_display = prev_month_str

        stats['yesterday_date'] = yesterday_str
        stats['yesterday_display'] = yesterday_display
        stats['previous_month'] = prev_month_str
        stats['previous_month_display'] = prev_month_display
    except Exception as e:
        print(f"Error computing dashboard date badges: {e}")
        stats['yesterday_date'] = ''
        stats['yesterday_display'] = 'Yesterday'
        stats['previous_month'] = ''
        stats['previous_month_display'] = 'Previous Month'

    return render_template(
        'dashboard.html',
        active_tab='DASHBOARD',
        stats=stats,
        branches=branches,
        active_branch=active_branch,
        active_branch_id=active_bid
    )

@app.route('/api/dashboard/yesterday-activity/pdf', methods=['GET'])
@login_required
def api_dashboard_yesterday_activity_pdf():
    try:
        from report_scheduler import get_yesterday_ist, build_yesterdays_activity_full_data
        role = session.get('role')
        req_comp_id = request.args.get('company_id')
        current_comp_id = get_current_company_id()

        # Support company_id override for super_admin or company_admin accessing own company
        if req_comp_id and (role == 'super_admin' or req_comp_id == current_comp_id):
            comp_id = req_comp_id
            company_info = database.get_company_by_id(comp_id) or get_current_company_info()
        else:
            comp_id = current_comp_id
            company_info = get_current_company_info()

        target_date = request.args.get('date') or get_yesterday_ist()
        branch_id = get_current_branch_id()
        activity_data = build_yesterdays_activity_full_data(comp_id, target_date=target_date, branch_id=branch_id)
        pdf_buffer = pdf_generator.generate_yesterdays_activity_report_pdf(
            activity_data,
            company_info=company_info
        )

        disposition_type = 'attachment' if request.args.get('download') == '1' else 'inline'
        filename = f"yesterday_activity_{target_date}.pdf"

        resp = make_response(pdf_buffer.getvalue())
        resp.headers['Content-Type'] = 'application/pdf'
        resp.headers['Content-Disposition'] = f'{disposition_type}; filename="{filename}"'
        resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        return resp
    except Exception as e:
        print(f"Error generating yesterday activity PDF: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/dashboard/previous-month-salary/pdf', methods=['GET'])
@login_required
def api_dashboard_previous_month_salary_pdf():
    try:
        from report_scheduler import get_previous_month_ist, build_monthly_salary_report
        comp_id = get_current_company_id()
        prev_month = get_previous_month_ist()
        company_info = get_current_company_info()

        result = database.get_salary_reports(
            start_month=prev_month,
            end_month=prev_month,
            limit=1000,
            company_id=comp_id,
            branch_id=get_current_branch_id()
        )
        data = result.get('data', []) if isinstance(result, dict) else []
        if not data:
            build_monthly_salary_report(comp_id, target_month=prev_month)
            result = database.get_salary_reports(
                start_month=prev_month,
                end_month=prev_month,
                limit=1000,
                company_id=comp_id,
                branch_id=get_current_branch_id()
            )
            data = result.get('data', []) if isinstance(result, dict) else []

        pdf_buffer = pdf_generator.generate_salary_report_pdf(data, company_info=company_info)

        disposition_type = 'attachment' if request.args.get('download') == '1' else 'inline'
        filename = f"salary_report_{prev_month}.pdf"

        resp = make_response(pdf_buffer.getvalue())
        resp.headers['Content-Type'] = 'application/pdf'
        resp.headers['Content-Disposition'] = f'{disposition_type}; filename="{filename}"'
        resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        return resp
    except Exception as e:
        print(f"Error generating previous month salary PDF: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/employee-details')
@login_required
def employee_details():
    comp_id = get_current_company_id()
    branches = []
    if comp_id:
        try:
            branches = database.get_company_branches(comp_id)
        except Exception:
            branches = []
    return render_template('employees.html', active_tab='EMPLOYEE DETAILS', branches=branches, active_branch_id=get_current_branch_id())

@app.route('/live-report')
@login_required
def live_report():
    return render_template('live_report.html', active_tab='LIVE REPORT')

@app.route('/attendance-report')
@login_required
def attendance_report():
    employees = database.get_employee_names_cached(company_id=get_current_company_id(), branch_id=get_current_branch_id())
    return render_template('attendance_report.html', active_tab='ATTENDANCE REPORT', employees=employees)

@app.route('/manual-entry')
@login_required
def manual_entry():
    employees = database.get_employee_names_cached(company_id=get_current_company_id(), branch_id=get_current_branch_id())
    return render_template('manual_entry.html', active_tab='MANUAL ENTRY', employees=employees)

@app.route('/payment-entry')
@login_required
def payment_entry():
    employees = database.get_employee_names_cached(company_id=get_current_company_id(), branch_id=get_current_branch_id())
    return render_template('payment_entry.html', active_tab='PAYMENT ENTRY', employees=employees)

@app.route('/advance-management')
@login_required
def advance_management():
    employees = database.get_employee_names_cached(company_id=get_current_company_id(), branch_id=get_current_branch_id())
    return render_template('advance_management.html', active_tab='PAYMENT ENTRY', employees=employees)

@app.route('/advance-summary')
@app.route('/balance-report')
@login_required
def balance_report():
    employees = database.get_employee_names_cached(company_id=get_current_company_id(), branch_id=get_current_branch_id())
    return render_template('balance_report.html', active_tab='PAYMENT ENTRY', employees=employees)

@app.route('/monthly-payslip')
@app.route('/payslip-preview')
@login_required
def monthly_payslip():
    comp_id = get_current_company_id()
    employees_res = database.get_all_employees(company_id=comp_id, branch_id=get_current_branch_id(), limit=1000)
    employees = [e['employee_name'] for e in employees_res['data']]
    return render_template('monthly_payslip.html', active_tab='PAYSLIP PREVIEW', employees=employees)

@app.route('/salary-report')
@login_required
def salary_report():
    employees_res = database.get_all_employees(company_id=get_current_company_id(), branch_id=get_current_branch_id(), limit=1000)
    employees = [e['employee_name'] for e in employees_res['data']]
    return render_template('salary_report.html', active_tab='SALARY REPORT', employees=employees)

# ----------------- COMPANY MANAGEMENT API (SUPER ADMIN ONLY) ----------------- #

@app.route('/api/companies', methods=['GET'])
@super_admin_required
def api_get_companies():
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    result = database.get_all_companies(search=search, page=page, limit=limit)
    result['companies'] = result.get('data', [])
    import math
    result['pages'] = math.ceil(result['total'] / limit) if limit else 1
    return jsonify(result)

@app.route('/api/companies', methods=['POST'])
@super_admin_required
def api_create_company():
    data = request.get_json() if request.is_json else request.form.to_dict()
    try:
        comp_id = database.create_company(data)
        return jsonify({'success': True, 'id': comp_id, 'message': 'Company registered successfully'})
    except ValueError as ve:
        return jsonify({'success': False, 'error': str(ve)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': f'Failed to create company: {e}'}), 500

@app.route('/api/companies/<comp_id>', methods=['GET'])
@super_admin_required
def api_get_company(comp_id):
    comp = database.get_company_by_id(comp_id)
    if not comp:
        return jsonify({'error': 'Company not found'}), 404
    return jsonify(comp)

@app.route('/api/companies/<comp_id>', methods=['PUT', 'POST'])
@super_admin_required
def api_update_company(comp_id):
    data = request.get_json() if request.is_json else request.form.to_dict()
    try:
        database.update_company(comp_id, data)
        return jsonify({'success': True, 'message': 'Company updated successfully'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400

@app.route('/api/companies/<comp_id>', methods=['DELETE'])
@super_admin_required
def api_delete_company(comp_id):
    try:
        res = database.delete_company(comp_id)
        return jsonify({
            'success': True,
            'message': res.get('message', 'Company and all associated records deleted successfully.'),
            'details': res
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400

@app.route('/api/companies/<comp_id>/toggle-support', methods=['POST'])
@super_admin_required
def api_toggle_company_support(comp_id):
    data = request.get_json(silent=True) or {}
    enabled = data.get('enabled')
    new_state = database.toggle_company_support(comp_id, enabled)
    if new_state is None:
        return jsonify({'success': False, 'error': 'Company not found'}), 404
    return jsonify({
        'success': True,
        'company_id': comp_id,
        'support_enabled': new_state,
        'message': f"Support feature {'enabled' if new_state else 'disabled'} successfully."
    })

@app.route('/api/companies/<comp_id>/toggle-auto-reports', methods=['POST'])
@super_admin_required
def api_toggle_company_auto_reports(comp_id):
    data = request.get_json(silent=True) or {}
    enabled = data.get('enabled')
    new_state = database.toggle_company_auto_reports(comp_id, enabled)
    if new_state is None:
        return jsonify({'success': False, 'error': 'Company not found'}), 404
    return jsonify({
        'success': True,
        'company_id': comp_id,
        'auto_email_reports': new_state,
        'message': f"Automatic email reports {'enabled' if new_state else 'disabled'} successfully."
    })

@app.route('/api/companies/<comp_id>/toggle-delete-entries', methods=['POST'])
@super_admin_required
def api_toggle_company_delete_entries(comp_id):
    data = request.get_json(silent=True) or {}
    enabled = data.get('enabled')
    new_state = database.toggle_company_delete_entries(comp_id, enabled)
    if new_state is None:
        return jsonify({'success': False, 'error': 'Company not found'}), 404
    return jsonify({
        'success': True,
        'company_id': comp_id,
        'can_delete_entries': new_state,
        'message': f"Delete entries feature {'enabled' if new_state else 'disabled'} successfully."
    })

@app.route('/api/companies/reports', methods=['GET'])
@super_admin_required
def api_company_reports():
    summary = database.get_company_reports_summary()
    return jsonify({'success': True, 'summary': summary})

@app.route('/api/companies/export/excel', methods=['GET'])
@app.route('/api/companies/export/csv', methods=['GET'])
@super_admin_required
def api_companies_export_excel():
    search = request.args.get('search', '').strip()
    result = database.get_all_companies(search=search, limit=10000)
    
    headers = ['SL NO', 'COMPANY NAME', 'REGISTERED DATE', 'DAYS ACTIVE', 'GSTIN', 'EMAIL', 'PHONE', 'ADDRESS', 'LATITUDE', 'LONGITUDE', 'SHIFT HOURS', 'EMPLOYEES', 'EMPLOYEE LIMIT', 'STATUS', 'SUPPORT']
    rows = []
    for idx, c in enumerate(result['data'], 1):
        rows.append([
            idx,
            c.get('company_name', ''),
            c.get('registered_date', '-'),
            c.get('days_active_label', f"{c.get('days_active', 0)} Days"),
            c.get('gstin', ''),
            c.get('email', ''),
            c.get('phone', ''),
            c.get('address', ''),
            c.get('latitude', ''),
            c.get('longitude', ''),
            c.get('shift_hours', '08:00'),
            c.get('employee_count', 0),
            c.get('employee_limit', 50),
            c.get('status', 'Active'),
            'Enabled' if c.get('support_enabled') else 'Disabled'
        ])
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    return export_table_data(headers, rows, "registered_companies", export_format=fmt, sheet_name="Companies")

@app.route('/api/companies/export/pdf', methods=['GET'])
@super_admin_required
def api_companies_export_pdf():
    search = request.args.get('search', '').strip()
    result = database.get_all_companies(search=search, limit=10000)
    company_info = get_current_company_info()
    pdf_buffer = pdf_generator.generate_companies_pdf(result['data'], company_info=company_info)
    filename = "registered_companies.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- SYSTEM ADMIN MULTI-TENANT SWITCHER API ----------------- #

@app.route('/api/system-admin/set-active-company', methods=['POST'])
@login_required
def api_system_admin_set_active_company():
    role = session.get('role')
    if role != 'super_admin':
        return jsonify({'success': False, 'error': 'Unauthorized. System Admin only.'}), 403
    
    data = request.get_json(silent=True) or {}
    company_id = str(data.get('company_id', '')).strip()
    if not company_id or company_id == 'ARGUS_MASTER':
        session['company_id'] = 'ARGUS_MASTER'
        session['company_name'] = 'ARGUS TECHNOLOGIES'
        session['active_branch_id'] = 'HOME'
        session['active_branch_name'] = 'Home Branch'
        session.modified = True
        return jsonify({
            'success': True,
            'company_id': 'ARGUS_MASTER',
            'company_name': 'System Administration (ARGUS TECHNOLOGIES)'
        })
    
    comp = database.get_company_by_id(company_id)
    if not comp:
        return jsonify({'success': False, 'error': f'Company {company_id} not found'}), 404
    
    c_name = comp.get('company_name', company_id)
    session['company_id'] = company_id
    session['company_name'] = c_name
    session['active_branch_id'] = 'HOME'
    session['active_branch_name'] = 'Home Branch'
    session.modified = True
    
    return jsonify({
        'success': True,
        'company_id': company_id,
        'company_name': c_name
    })

# ----------------- COMPANY PROFILE API ROUTES ----------------- #

@app.route('/api/company-profile', methods=['GET'])
@login_required
def api_get_company_profile():
    role = session.get('role', 'company_admin')
    if role == 'super_admin':
        req_id = request.args.get('company_id', '').strip()
        all_comps = database.get_all_companies(limit=1000)['data']
        
        # Build dropdown list with System Administration prominently at top
        comps_dropdown = [{'id': 'ARGUS_MASTER', 'company_name': 'System Administration (ARGUS TECHNOLOGIES)'}]
        for c in all_comps:
            if c.get('id') != 'ARGUS_MASTER':
                comps_dropdown.append({'id': c['id'], 'company_name': c.get('company_name', c['id'])})

        target_id = req_id or 'ARGUS_MASTER'
        company = database.get_company_by_id(target_id)
        if not company and target_id == 'ARGUS_MASTER':
            company = database.get_company_by_id('ARGUS_MASTER') or {
                'id': 'ARGUS_MASTER',
                'company_name': 'ARGUS TECHNOLOGIES',
                'email': 'technologiesargus@gmail.com'
            }

        is_master = (target_id == 'ARGUS_MASTER')
        if company:
            is_locked = company.get('coordinates_locked')
            if is_locked is None:
                lat = company.get('latitude')
                lng = company.get('longitude')
                is_locked = bool(lat is not None and lng is not None and str(lat).strip() != '' and str(lng).strip() != '')
            company['coordinates_locked'] = bool(is_locked)

            raw_pass = company.get('password_raw') or ''
            if not raw_pass and is_master:
                db = database.get_db()
                adm = db.admin_users.find_one({'role': 'super_admin'}) or db.admin_users.find_one({'username': 'Admin'})
                if adm:
                    raw_pass = adm.get('password_raw') or adm.get('password', '')
                    if raw_pass and (raw_pass.startswith('pbkdf2:') or raw_pass.startswith('scrypt:')):
                        raw_pass = ''
            elif not raw_pass:
                p = company.get('password', '')
                if p and not p.startswith('pbkdf2:') and not p.startswith('scrypt:'):
                    raw_pass = p
            company['saved_password'] = raw_pass
            company['has_password'] = bool(raw_pass or company.get('password') or company.get('password_hash'))
            company['logo'] = company.get('logo') or ''
            company['logo_url'] = company.get('logo_data') or (f"/uploads/{company['logo']}" if company.get('logo') else '')
            company.pop('password', None)
            company.pop('password_hash', None)

        return jsonify({
            'success': True,
            'role': 'super_admin',
            'company': company,
            'is_master': is_master,
            'companies_list': comps_dropdown
        })
    else:
        comp_id = session.get('company_id')
        company = database.get_company_by_id(comp_id) if comp_id else None

        # Multi-tier self-healing company lookup
        if not company:
            # Tier 1: Look up by company_email or admin_email in session
            c_email = session.get('company_email') or session.get('admin_email')
            if c_email:
                comp_rec = database.get_company_by_email(c_email)
                if comp_rec and comp_rec.get('id'):
                    comp_id = comp_rec['id']
                    session['company_id'] = comp_id
                    session['company_name'] = comp_rec.get('company_name', session.get('company_name', ''))
                    session.modified = True
                    company = database.get_company_by_id(comp_id)

            # Tier 2: Check query parameter ?company_id=
            if not company and request.args.get('company_id'):
                arg_id = request.args.get('company_id').strip()
                comp_arg = database.get_company_by_id(arg_id)
                if comp_arg:
                    comp_id = comp_arg['id']
                    session['company_id'] = comp_id
                    session['company_name'] = comp_arg.get('company_name', session.get('company_name', ''))
                    session.modified = True
                    company = comp_arg

            # Tier 3: Super Admin fallback to ARGUS_MASTER
            if not company and (c_email in ['technologiesargus@gmail.com', 'productionargus@gmail.com'] or comp_id == 'ARGUS_MASTER'):
                company = database.get_company_by_id('ARGUS_MASTER')
                if company:
                    comp_id = 'ARGUS_MASTER'
                    session['company_id'] = comp_id
                    session['role'] = 'super_admin'
                    session.modified = True

            # Tier 4: Fallback to first active company in database
            if not company:
                db = database.get_db()
                first_comp = db.company_admin.find_one({'id': {'$nin': ['ALL', 'ARGUS_MASTER']}, 'status': {'$nin': ['Deactive', 'Inactive']}})
                if first_comp:
                    first_id = first_comp.get('id')
                    company = database.get_company_by_id(first_id)
                    if company:
                        session['company_id'] = first_id
                        session['company_name'] = company.get('company_name', '')
                        session.modified = True

            # Tier 5: Fallback to ARGUS_MASTER
            if not company:
                company = database.get_company_by_id('ARGUS_MASTER')
                if company:
                    session['company_id'] = 'ARGUS_MASTER'
                    session.modified = True

            # Tier 6: Safe company dictionary fallback if none found in DB
            if not company:
                company = {
                    'id': comp_id or 'ARGUS_MASTER',
                    'company_name': session.get('company_name') or 'ARGUS TECHNOLOGIES',
                    'email': c_email or 'technologiesargus@gmail.com',
                    'phone': '',
                    'address': '',
                    'latitude': None,
                    'longitude': None,
                    'radius_meters': 150,
                    'coordinates_locked': False,
                    'status': 'Active',
                    'shift_hours': '08:00',
                    'employee_limit': 50,
                    'employee_count': 0
                }
                session['company_id'] = company['id']
                session.modified = True

        is_locked = company.get('coordinates_locked')
        if is_locked is None:
            lat = company.get('latitude')
            lng = company.get('longitude')
            is_locked = bool(lat is not None and lng is not None and str(lat).strip() != '' and str(lng).strip() != '')
        company['coordinates_locked'] = bool(is_locked)

        raw_pass = company.get('password_raw') or ''
        if not raw_pass:
            p = company.get('password', '')
            if p and not p.startswith('pbkdf2:') and not p.startswith('scrypt:'):
                raw_pass = p
        company['saved_password'] = raw_pass
        company['has_password'] = bool(raw_pass or company.get('password') or company.get('password_hash'))
        company['logo'] = company.get('logo') or ''
        company['logo_url'] = company.get('logo_data') or (f"/uploads/{company['logo']}" if company.get('logo') else '')
        company.pop('password', None)
        company.pop('password_hash', None)
        return jsonify({
            'success': True,
            'role': 'company_admin',
            'company': company,
            'has_password': company['has_password']
        })

@app.route('/api/company-profile', methods=['PUT'])
@login_required
def api_update_company_profile():
    role = session.get('role')
    if not role and session.get('admin_logged_in'):
        role = 'super_admin' if session.get('company_id') == 'ARGUS_MASTER' or not session.get('company_id') else 'company_admin'
    role = role or 'company_admin'
    data = request.get_json(silent=True) or {}
    
    if role == 'company_admin':
        comp_id = session.get('company_id')
        comp = database.get_company_by_id(comp_id) if comp_id else None

        if not comp:
            c_email = session.get('company_email') or session.get('admin_email')
            if c_email:
                comp_rec = database.get_company_by_email(c_email)
                if comp_rec and comp_rec.get('id'):
                    comp_id = comp_rec['id']
                    session['company_id'] = comp_id
                    session.modified = True
                    comp = database.get_company_by_id(comp_id)

            if not comp and data.get('id'):
                comp_arg = database.get_company_by_id(data.get('id'))
                if comp_arg:
                    comp_id = comp_arg['id']
                    session['company_id'] = comp_id
                    session.modified = True
                    comp = comp_arg

            if not comp:
                db = database.get_db()
                first_comp = db.company_admin.find_one({'id': {'$nin': ['ALL', 'ARGUS_MASTER']}, 'status': {'$nin': ['Deactive', 'Inactive']}})
                if first_comp:
                    first_id = first_comp.get('id')
                    comp = database.get_company_by_id(first_id)
                    if comp:
                        comp_id = first_id
                        session['company_id'] = comp_id
                        session.modified = True

        if not comp:
            return jsonify({'success': False, 'error': 'Company profile not found'}), 404

        is_locked = comp.get('coordinates_locked')
        if is_locked is None:
            lat = comp.get('latitude')
            lng = comp.get('longitude')
            is_locked = bool(lat is not None and lng is not None and str(lat).strip() != '' and str(lng).strip() != '')

        # Coordinates handling: only once allowed for company_admin
        if 'latitude' in data or 'longitude' in data:
            lat_val = data.get('latitude')
            lng_val = data.get('longitude')

            # Check if user typed or pasted combined coordinates (e.g. "11.0298, 76.9740")
            if lat_val and (',' in str(lat_val) or ';' in str(lat_val)) and (not lng_val or str(lng_val).strip() == ''):
                import re
                parts = re.split(r'[,;]+', str(lat_val))
                if len(parts) >= 2:
                    lat_val = parts[0].strip()
                    lng_val = parts[1].strip()

            # If not locked yet, company admin can lock them for the first time
            if not is_locked and lat_val is not None and lng_val is not None and str(lat_val).strip() != '' and str(lng_val).strip() != '':
                new_lat = database.parse_coordinate_value(lat_val)
                new_lng = database.parse_coordinate_value(lng_val)
                database.update_company(comp_id, {
                    'latitude': new_lat,
                    'longitude': new_lng,
                    'coordinates_locked': True,
                    'coordinates_locked_at': database.get_ist_now()
                })

        if 'address' in data:
            address = str(data.get('address', '')).strip()
            if address:
                database.update_company(comp_id, {'address': address})

        new_pass = str(data.get('password', '')).strip()
        if new_pass:
            if len(new_pass) < 4:
                return jsonify({'success': False, 'error': 'Password must be at least 4 characters.'}), 400
            database.set_company_password(comp_id, new_pass)
            
        return jsonify({'success': True, 'message': 'Company profile updated successfully.'})
    
    elif role == 'super_admin':
        target_id = data.get('id') or data.get('company_id') or session.get('company_id') or 'ARGUS_MASTER'
        
        new_pass = str(data.get('password', '')).strip()
        if new_pass:
            if len(new_pass) < 4:
                return jsonify({'success': False, 'error': 'Password must be at least 4 characters.'}), 400
            if target_id == 'ARGUS_MASTER':
                database.set_admin_password(new_pass)
            else:
                database.set_company_password(target_id, new_pass)
            
        try:
            data_to_save = dict(data)
            data_to_save.pop('password', None)
            data_to_save.pop('password_hash', None)

            lat_val = data_to_save.get('latitude')
            lng_val = data_to_save.get('longitude')
            # Check if user typed or pasted combined coordinates
            if lat_val and (',' in str(lat_val) or ';' in str(lat_val)) and (not lng_val or str(lng_val).strip() == ''):
                import re
                parts = re.split(r'[,;]+', str(lat_val))
                if len(parts) >= 2:
                    lat_val = parts[0].strip()
                    lng_val = parts[1].strip()
                    data_to_save['longitude'] = lng_val

            if 'latitude' in data_to_save:
                val = lat_val
                if val is not None and str(val).strip() != '':
                    data_to_save['latitude'] = database.parse_coordinate_value(val)
                else:
                    data_to_save['latitude'] = None

            if 'longitude' in data_to_save:
                val = data_to_save.get('longitude')
                if val is not None and str(val).strip() != '':
                    data_to_save['longitude'] = database.parse_coordinate_value(val)
                else:
                    data_to_save['longitude'] = None

            if data_to_save.get('latitude') is not None and data_to_save.get('longitude') is not None and str(data_to_save.get('latitude')).strip() != '' and str(data_to_save.get('longitude')).strip() != '':
                data_to_save['coordinates_locked'] = True
                data_to_save['coordinates_locked_at'] = database.get_ist_now()

            database.update_company(target_id, data_to_save)
            return jsonify({'success': True, 'message': 'Company profile updated successfully.'})
        except Exception as e:
            print("Error in super_admin update_company_profile:", traceback.format_exc())
            return jsonify({'success': False, 'error': str(e)}), 400
            
    return jsonify({'success': False, 'error': 'Unauthorized: Admin role not recognized.'}), 403

@app.route('/api/company-profile/logo', methods=['POST'])
@login_required
def api_upload_company_logo():
    role = session.get('role', 'company_admin')
    if role not in ['super_admin', 'company_admin']:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 403

    if 'logo' not in request.files:
        return jsonify({'success': False, 'error': 'No logo file provided'}), 400

    file = request.files['logo']
    if not file or not file.filename:
        return jsonify({'success': False, 'error': 'No file selected'}), 400

    allowed_exts = {'png', 'jpg', 'jpeg', 'webp', 'svg'}
    ext = file.filename.rsplit('.', 1)[-1].lower() if '.' in file.filename else ''
    if ext not in allowed_exts:
        return jsonify({'success': False, 'error': f'Unsupported format .{ext}. Please upload PNG, JPG, JPEG, or WEBP.'}), 400

    if role == 'super_admin':
        comp_id = request.form.get('company_id') or session.get('company_id') or 'ARGUS_MASTER'
    else:
        comp_id = session.get('company_id')
        if not comp_id or not database.get_company_by_id(comp_id):
            c_email = session.get('company_email') or session.get('admin_email')
            if c_email:
                comp_rec = database.get_company_by_email(c_email)
                if comp_rec and comp_rec.get('id'):
                    comp_id = comp_rec['id']
                    session['company_id'] = comp_id
                    session.modified = True
            if not comp_id and request.form.get('company_id'):
                comp_id = request.form.get('company_id')
                session['company_id'] = comp_id
                session.modified = True
            if not comp_id:
                db = database.get_db()
                first_comp = db.company_admin.find_one({'id': {'$nin': ['ALL', 'ARGUS_MASTER']}, 'status': {'$nin': ['Deactive', 'Inactive']}})
                if first_comp:
                    comp_id = first_comp.get('id')
                    session['company_id'] = comp_id
                    session.modified = True
        if not comp_id:
            return jsonify({'success': False, 'error': 'No company ID associated with session'}), 403

    comp_id_clean = re.sub(r'[^a-zA-Z0-9_-]', '_', str(comp_id))
    timestamp = int(datetime.now().timestamp())
    filename = f"company_logo_{comp_id_clean}_{timestamp}.{ext}"
    upload_folder = app.config['UPLOAD_FOLDER']
    os.makedirs(upload_folder, exist_ok=True)
    file_path = os.path.join(upload_folder, filename)

    file_bytes = file.read()
    with open(file_path, 'wb') as f:
        f.write(file_bytes)

    # Generate base64 data URI for database storage and recovery
    mime = f"image/{'svg+xml' if ext == 'svg' else ext}"
    if ext == 'jpg': mime = 'image/jpeg'
    b64_str = base64.b64encode(file_bytes).decode('utf-8')
    logo_data = f"data:{mime};base64,{b64_str}"

    db = database.get_db()
    if comp_id == 'ARGUS_MASTER':
        db.company_admin.update_one(
            {'id': 'ARGUS_MASTER'},
            {'$set': {'id': 'ARGUS_MASTER', 'company_name': 'ARGUS TECHNOLOGIES', 'logo': filename, 'logo_data': logo_data, 'updated_at': database.get_ist_now()}},
            upsert=True
        )
    else:
        database.update_company(comp_id, {'logo': filename, 'logo_data': logo_data})

    return jsonify({
        'success': True,
        'logo': filename,
        'logo_url': f"/uploads/{filename}",
        'message': 'Company logo uploaded and saved successfully.'
    })

@app.route('/api/company-profile/logo', methods=['DELETE'])
@login_required
def api_delete_company_logo():
    role = session.get('role', 'company_admin')
    if role not in ['super_admin', 'company_admin']:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 403

    if role == 'super_admin':
        comp_id = request.args.get('company_id') or session.get('company_id') or 'ARGUS_MASTER'
    else:
        comp_id = session.get('company_id')
        if not comp_id or not database.get_company_by_id(comp_id):
            c_email = session.get('company_email') or session.get('admin_email')
            if c_email:
                comp_rec = database.get_company_by_email(c_email)
                if comp_rec and comp_rec.get('id'):
                    comp_id = comp_rec['id']
                    session['company_id'] = comp_id
                    session.modified = True
            if not comp_id and request.args.get('company_id'):
                comp_id = request.args.get('company_id')
                session['company_id'] = comp_id
                session.modified = True
            if not comp_id:
                db = database.get_db()
                first_comp = db.company_admin.find_one({'id': {'$nin': ['ALL', 'ARGUS_MASTER']}, 'status': {'$nin': ['Deactive', 'Inactive']}})
                if first_comp:
                    comp_id = first_comp.get('id')
                    session['company_id'] = comp_id
                    session.modified = True

    db = database.get_db()
    db.company_admin.update_one(
        {'id': str(comp_id)},
        {'$set': {'logo': '', 'logo_data': '', 'updated_at': database.get_ist_now()}}
    )
    return jsonify({'success': True, 'message': 'Company logo removed successfully.'})

# ----------------- EMPLOYEE PORTAL API ROUTES ----------------- #

@app.route('/api/employee/my-attendance', methods=['GET'])
@employee_required
def api_employee_my_attendance():
    emp_id = session.get('employee_id')
    curr_emp = database.get_employee_by_id(emp_id)
    if curr_emp and curr_emp.get('permissions') and curr_emp['permissions'].get('attendance_history') is False:
        return jsonify({'error': 'Attendance history view is disabled for your account.'}), 403

    emp_name = session.get('employee_name')
    comp_id = session.get('company_id')
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()
    result = database.get_attendance_reports(
        report_type='all',
        start_date=from_date,
        end_date=to_date,
        employee=emp_name,
        limit=500,
        company_id=comp_id
    )
    return jsonify(result)

@app.route('/api/employee/my-payslip', methods=['GET'])
@employee_required
def api_employee_my_payslip():
    emp_id = session.get('employee_id')
    curr_emp = database.get_employee_by_id(emp_id)
    if curr_emp and curr_emp.get('permissions') and curr_emp['permissions'].get('monthly_payslip') is False:
        return jsonify({'error': 'Monthly payslip view is disabled for your account.'}), 403

    emp_name = session.get('employee_name')
    comp_id = session.get('company_id')
    month = request.args.get('month', '').strip()
    if not month:
        month = datetime.now().strftime("%Y-%m")
    payslip = database.get_payslip_data(emp_name, month, company_id=comp_id)
    return jsonify(payslip)

@app.route('/api/employee/credentials', methods=['POST'])
@employee_required
def api_employee_set_credentials():
    """Allows authenticated employee to create or update their portal password."""
    emp_id = session.get('employee_id')
    curr_emp = database.get_employee_by_id(emp_id)
    if curr_emp and curr_emp.get('permissions') and curr_emp['permissions'].get('employee_credentials') is False:
        return jsonify({'success': False, 'error': 'Credential updates are disabled for your account.'}), 403

    data = request.get_json(silent=True) or request.form.to_dict()
    new_password = str(data.get('password', '')).strip()
    if not new_password or len(new_password) < 4:
        return jsonify({'success': False, 'error': 'Password must be at least 4 characters long.'}), 400
    
    success = database.set_employee_password(emp_id, new_password)
    if success:
        return jsonify({
            'success': True,
            'message': 'Your password has been saved successfully! You can now sign in using your email and password.'
        })
    return jsonify({'success': False, 'error': 'Failed to save password. Please try again.'}), 500

# ----------------- LEAVE & PERMISSION API ROUTES ----------------- #

LEAVES_FOLDER = os.path.join(app.config['UPLOAD_FOLDER'], 'leaves')
try:
    os.makedirs(LEAVES_FOLDER, exist_ok=True)
except Exception:
    pass

@app.route('/uploads/leaves/<path:filename>')
def serve_leave_attachment(filename):
    return send_from_directory(LEAVES_FOLDER, filename)

@app.route('/api/employee/leave-balance', methods=['GET'])
@employee_required
def api_employee_leave_balance():
    emp_id = session.get('employee_id')
    comp_id = session.get('company_id')
    bal = database.get_employee_leave_balance(emp_id, company_id=comp_id)
    return jsonify({'success': True, 'balance': bal, 'balances': bal})

@app.route('/api/employee/leave-requests', methods=['GET'])
@employee_required
def api_employee_leave_requests():
    emp_id = session.get('employee_id')
    limit = int(request.args.get('limit', 20))
    page = int(request.args.get('page', 1))
    status = request.args.get('status', 'ALL')
    res = database.get_leave_requests(employee_id=emp_id, status=status, limit=limit, page=page)
    return jsonify({'success': True, 'requests': res['requests'], 'total': res['total']})

@app.route('/api/employee/leave-requests/apply', methods=['POST'])
@employee_required
def api_employee_apply_leave():
    emp_id = session.get('employee_id')
    emp = database.get_employee_by_id(emp_id)
    if not emp:
        return jsonify({'success': False, 'error': 'Employee profile not found'}), 404

    data = request.form.to_dict()
    leave_type = data.get('leave_type', '').strip()
    from_date = data.get('from_date', '').strip()
    to_date = data.get('to_date', '').strip() or from_date
    session_type = data.get('session', 'Full Day').strip()
    reason = data.get('reason', '').strip()

    if not leave_type or not from_date or not reason:
        return jsonify({'success': False, 'error': 'Please fill in all required fields (Leave Type, Dates, Reason).'}), 400

    attachments = []
    if 'attachment' in request.files:
        file = request.files['attachment']
        if file and file.filename:
            fname = werkzeug.utils.secure_filename(file.filename) or 'leave_doc'
            base_name, ext = os.path.splitext(fname)
            unique_fname = f"{base_name}_{int(time.time())}{ext}"
            os.makedirs(LEAVES_FOLDER, exist_ok=True)
            dest_path = os.path.join(LEAVES_FOLDER, unique_fname)
            file.save(dest_path)
            size_kb = round(os.path.getsize(dest_path) / 1024, 1)
            attachments.append({
                'filename': unique_fname,
                'original_name': fname,
                'file_url': f"/uploads/leaves/{unique_fname}",
                'size_str': f"{size_kb} KB",
                'mime_type': file.mimetype or 'application/octet-stream',
                'uploaded_at': database.get_ist_now().strftime('%d %b %Y %I:%M %p')
            })

    req_doc = {
        'company_id': emp.get('company_id') or session.get('company_id'),
        'employee_id': emp_id,
        'employee_code': emp.get('employee_id') or emp.get('emp_id') or str(emp_id),
        'employee_name': emp.get('employee_name', ''),
        'department': emp.get('department', 'General'),
        'designation': emp.get('designation', ''),
        'mobile_number': emp.get('mobile_number', ''),
        'leave_type': leave_type,
        'from_date': from_date,
        'to_date': to_date,
        'session': session_type,
        'reason': reason,
        'attachments': attachments
    }
    # For hourly employees, store time-based leave details
    if session_type == 'Hourly':
        req_doc['leave_from_time'] = data.get('leave_from_time', '').strip()
        req_doc['leave_to_time'] = data.get('leave_to_time', '').strip()
        req_doc['leave_duration'] = data.get('leave_duration', '').strip()
        try:
            req_doc['duration_hours'] = float(data.get('duration_hours', 0))
        except (ValueError, TypeError):
            req_doc['duration_hours'] = 0.0
    new_req = database.submit_leave_request(req_doc)
    return jsonify({
        'success': True,
        'message': 'Your leave request has been submitted successfully!',
        'request': new_req
    })

@app.route('/api/employee/permission-requests/apply', methods=['POST'])
@employee_required
def api_employee_apply_permission():
    emp_id = session.get('employee_id')
    emp = database.get_employee_by_id(emp_id)
    if not emp:
        return jsonify({'success': False, 'error': 'Employee profile not found'}), 404

    data = request.form.to_dict()
    perm_type = data.get('permission_type', 'Short Leave').strip()
    perm_date = data.get('date', '').strip()
    from_time = data.get('from_time', '').strip()
    to_time = data.get('to_time', '').strip()
    duration = data.get('duration', '02:00').strip()
    reason = data.get('reason', '').strip()

    if not perm_type or not perm_date or not from_time or not to_time:
        return jsonify({'success': False, 'error': 'Please fill in all required fields (Permission Type, Date, Times).'}), 400

    attachments = []
    if 'attachment' in request.files:
        file = request.files['attachment']
        if file and file.filename:
            fname = werkzeug.utils.secure_filename(file.filename) or 'perm_doc'
            base_name, ext = os.path.splitext(fname)
            unique_fname = f"{base_name}_{int(time.time())}{ext}"
            os.makedirs(LEAVES_FOLDER, exist_ok=True)
            dest_path = os.path.join(LEAVES_FOLDER, unique_fname)
            file.save(dest_path)
            size_kb = round(os.path.getsize(dest_path) / 1024, 1)
            attachments.append({
                'filename': unique_fname,
                'original_name': fname,
                'file_url': f"/uploads/leaves/{unique_fname}",
                'size_str': f"{size_kb} KB",
                'mime_type': file.mimetype or 'application/octet-stream',
                'uploaded_at': database.get_ist_now().strftime('%d %b %Y %I:%M %p')
            })

    req_doc = {
        'company_id': emp.get('company_id') or session.get('company_id'),
        'employee_id': emp_id,
        'employee_code': emp.get('employee_id') or emp.get('emp_id') or str(emp_id),
        'employee_name': emp.get('employee_name', ''),
        'department': emp.get('department', 'General'),
        'designation': emp.get('designation', ''),
        'mobile_number': emp.get('mobile_number', ''),
        'permission_type': perm_type,
        'date': perm_date,
        'from_time': from_time,
        'to_time': to_time,
        'duration': duration,
        'reason': reason,
        'attachments': attachments
    }
    new_req = database.submit_permission_request(req_doc)
    return jsonify({
        'success': True,
        'message': 'Your permission request has been submitted successfully!',
        'request': new_req
    })

@app.route('/api/employee/leave-requests/<request_id>', methods=['DELETE'])
@employee_required
def api_employee_delete_leave_request(request_id):
    emp_id = session.get('employee_id')
    res = database.delete_employee_leave_request(request_id, emp_id=emp_id)
    if not res.get('success'):
        return jsonify(res), 400
    return jsonify(res)

@app.route('/api/employee/leave-requests/<request_id>/edit', methods=['POST'])
@employee_required
def api_employee_edit_leave_request(request_id):
    emp_id = session.get('employee_id')
    data = request.form.to_dict()
    
    req = database.get_leave_request_by_id(request_id)
    if not req:
        return jsonify({'success': False, 'error': 'Request not found'}), 404

    if str(req.get('status', 'Pending')).strip().lower() != 'pending':
        return jsonify({'success': False, 'error': f"Cannot edit request that has already been {req.get('status')}"}), 400

    updated_fields = {}
    if req.get('request_type') == 'Permission':
        perm_type = data.get('permission_type', 'Short Leave').strip()
        perm_date = data.get('date', '').strip()
        from_time = data.get('from_time', '').strip()
        to_time = data.get('to_time', '').strip()
        duration_str = data.get('duration', '02:00').strip()
        reason = data.get('reason', '').strip()

        duration_hours = 2.0
        try:
            clean_dur = duration_str.split(' ')[0]
            if ':' in clean_dur:
                hh, mm = clean_dur.split(':')
                duration_hours = round(int(hh) + int(mm) / 60.0, 2)
            else:
                duration_hours = float(clean_dur)
        except Exception:
            duration_hours = 2.0

        updated_fields = {
            'permission_type': perm_type,
            'date': perm_date,
            'from_date': perm_date,
            'to_date': perm_date,
            'from_time': from_time,
            'to_time': to_time,
            'duration': duration_str,
            'duration_str': duration_str,
            'duration_hours': duration_hours,
            'hours_count': duration_hours,
            'reason': reason
        }
    else:
        leave_type = data.get('leave_type', 'Casual Leave (CL)').strip()
        session_type = data.get('session', 'Full Day').strip()
        from_date = data.get('from_date', '').strip()
        to_date = data.get('to_date', '').strip() or from_date
        reason = data.get('reason', '').strip()

        total_days = 1.0
        if session_type == 'Half Day':
            total_days = 0.5
            to_date = from_date
        elif session_type == 'Hourly':
            total_days = 0.0
            to_date = from_date
        else:
            try:
                d1 = datetime.strptime(from_date, '%Y-%m-%d').date()
                d2 = datetime.strptime(to_date, '%Y-%m-%d').date()
                total_days = max(1.0, float((d2 - d1).days + 1))
            except Exception:
                total_days = 1.0

        updated_fields = {
            'leave_type': leave_type,
            'session': session_type,
            'from_date': from_date,
            'to_date': to_date,
            'total_days': total_days,
            'days_count': total_days,
            'reason': reason
        }
        if session_type == 'Hourly':
            updated_fields['from_time'] = data.get('leave_from_time', '').strip()
            updated_fields['to_time'] = data.get('leave_to_time', '').strip()
            updated_fields['leave_from_time'] = data.get('leave_from_time', '').strip()
            updated_fields['leave_to_time'] = data.get('leave_to_time', '').strip()
            updated_fields['leave_duration'] = data.get('leave_duration', '').strip()
            try:
                updated_fields['duration_hours'] = float(data.get('duration_hours', 0))
            except Exception:
                updated_fields['duration_hours'] = 0.0

    if 'attachment' in request.files:
        file = request.files['attachment']
        if file and file.filename:
            fname = werkzeug.utils.secure_filename(file.filename) or 'doc'
            base_name, ext = os.path.splitext(fname)
            unique_fname = f"{base_name}_{int(time.time())}{ext}"
            os.makedirs(LEAVES_FOLDER, exist_ok=True)
            dest_path = os.path.join(LEAVES_FOLDER, unique_fname)
            file.save(dest_path)
            size_kb = round(os.path.getsize(dest_path) / 1024, 1)
            updated_fields['attachments'] = [{
                'filename': unique_fname,
                'original_name': fname,
                'file_url': f"/uploads/leaves/{unique_fname}",
                'size_str': f"{size_kb} KB",
                'mime_type': file.mimetype or 'application/octet-stream',
                'uploaded_at': database.get_ist_now().strftime('%d %b %Y %I:%M %p')
            }]

    res = database.edit_employee_leave_request(request_id, emp_id, updated_fields)
    if not res.get('success'):
        return jsonify(res), 400
    return jsonify(res)

@app.route('/api/admin/leave-permission/stats', methods=['GET'])
@login_required
def api_admin_leave_stats():
    comp_id = get_current_company_id()
    req_comp = request.args.get('company_id')
    if req_comp and session.get('role') == 'super_admin':
        comp_id = req_comp
    stats = database.get_admin_leave_stats(company_id=comp_id, branch_id=get_current_branch_id())
    return jsonify({'success': True, 'stats': stats})

@app.route('/api/admin/leave-permission/policy', methods=['GET', 'POST'])
@login_required
def api_admin_leave_policy():
    comp_id = get_current_company_id()
    if request.method == 'POST':
        data = request.get_json(silent=True) or request.form.to_dict() or {}
        req_comp = data.get('company_id') or request.args.get('company_id')
        if req_comp and session.get('role') == 'super_admin':
            comp_id = req_comp
        saved = database.save_company_leave_policy(comp_id, data)
        return jsonify({
            'success': saved,
            'message': 'Annual leave policy saved and persisted successfully!',
            'policy': database.get_company_leave_policy(comp_id)
        })
    req_comp = request.args.get('company_id')
    if req_comp and session.get('role') == 'super_admin':
        comp_id = req_comp
    pol = database.get_company_leave_policy(comp_id)
    return jsonify({'success': True, 'policy': pol})

@app.route('/api/admin/leave-permission/balances', methods=['GET'])
@login_required
def api_admin_leave_balances():
    comp_id = get_current_company_id()
    req_comp = request.args.get('company_id')
    if req_comp and session.get('role') == 'super_admin':
        comp_id = req_comp
    search = request.args.get('search', '').strip()
    bals = database.get_all_employees_leave_balances(company_id=comp_id, search=search, branch_id=get_current_branch_id())
    return jsonify({'success': True, 'balances': bals})

@app.route('/api/admin/leave-permission/employee-quota', methods=['POST'])
@login_required
def api_admin_save_employee_leave_quota():
    data = request.get_json(silent=True) or request.form.to_dict() or {}
    emp_id = data.get('employee_id') or data.get('id')
    if not emp_id:
        return jsonify({'success': False, 'error': 'Employee ID is required'}), 400
    comp_id = get_current_company_id()
    req_comp = data.get('company_id') or request.args.get('company_id')
    if req_comp and session.get('role') == 'super_admin':
        comp_id = req_comp
    if 'custom_leaves' in data:
        res = database.save_employee_custom_leaves(emp_id, data.get('custom_leaves', []), company_id=comp_id)
    else:
        res = database.save_employee_leave_quota(emp_id, data, company_id=comp_id)
    if not res.get('success'):
        return jsonify(res), 400
    return jsonify(res)

@app.route('/api/admin/leave-permission/employee-custom-leaves', methods=['POST'])
@login_required
def api_admin_save_employee_custom_leaves():
    data = request.get_json(silent=True) or request.form.to_dict() or {}
    emp_id = data.get('employee_id') or data.get('id')
    if not emp_id:
        return jsonify({'success': False, 'error': 'Employee ID is required'}), 400
    comp_id = get_current_company_id()
    req_comp = data.get('company_id') or request.args.get('company_id')
    if req_comp and session.get('role') == 'super_admin':
        comp_id = req_comp
    custom_leaves = data.get('custom_leaves', [])
    res = database.save_employee_custom_leaves(emp_id, custom_leaves, company_id=comp_id)
    if not res.get('success'):
        return jsonify(res), 400
    return jsonify(res)

@app.route('/api/admin/leave-permission/requests', methods=['GET'])
@login_required
def api_admin_leave_requests():
    comp_id = get_current_company_id()
    req_comp = request.args.get('company_id')
    if req_comp and session.get('role') == 'super_admin':
        comp_id = req_comp
    status = request.args.get('status', 'ALL')
    limit = int(request.args.get('limit', 50))
    page = int(request.args.get('page', 1))
    emp_id = request.args.get('employee_id', None)
    res = database.get_leave_requests(company_id=comp_id, employee_id=emp_id, status=status, limit=limit, page=page, branch_id=get_current_branch_id())
    return jsonify({
        'success': True,
        'requests': res['requests'],
        'total': res['total'],
        'page': res['page'],
        'limit': res['limit'],
        'status_counts': res.get('status_counts', {'all': 0, 'pending': 0, 'approved': 0, 'rejected': 0})
    })

@app.route('/api/admin/leave-permission/requests/<request_id>', methods=['GET'])
@login_required
def api_admin_leave_request_detail(request_id):
    req = database.get_leave_request_by_id(request_id)
    if not req:
        return jsonify({'success': False, 'error': 'Leave request not found'}), 404
    emp_id = req.get('employee_id')
    comp_id = req.get('company_id')
    emp = database.get_employee_by_id(emp_id)
    bal = database.get_employee_leave_balance(emp_id, company_id=comp_id)
    return jsonify({
        'success': True,
        'request': req,
        'employee': emp or {},
        'balance': bal or {}
    })

@app.route('/api/admin/leave-permission/requests/<request_id>/action', methods=['POST'])
@login_required
def api_admin_leave_request_action(request_id):
    data = request.get_json(silent=True) or request.form.to_dict()
    action = str(data.get('action') or data.get('status') or '').strip()
    if action.lower() in ['approve', 'approved']:
        action = 'Approved'
    elif action.lower() in ['reject', 'rejected']:
        action = 'Rejected'
    else:
        return jsonify({'success': False, 'error': "Action must be 'Approved' or 'Rejected'"}), 400
    remark = data.get('admin_remark', '').strip()
    reviewer = session.get('user') or 'Administrator'
    res = database.update_leave_request_status(request_id, action, admin_remark=remark, reviewer=reviewer)
    return jsonify(res)

# ----------------- EMPLOYEE MANAGEMENT API ROUTES ----------------- #

@app.route('/api/dashboard/stats')
@login_required
def api_dashboard_stats():
    stats = database.get_dashboard_stats(company_id=get_current_company_id(), branch_id=get_current_branch_id())
    return jsonify(stats)

@app.route('/api/employees', methods=['GET'])
@login_required
def api_get_employees():
    search = request.args.get('search', '').strip()
    sort_col = request.args.get('sort_col', 'id')
    sort_dir = request.args.get('sort_dir', 'desc')
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    status = request.args.get('status', '').strip()
    
    branch_id = request.args.get('branch_id') or get_current_branch_id()
    result = database.get_all_employees(
        company_id=get_current_company_id(),
        search=search,
        sort_col=sort_col,
        sort_dir=sort_dir,
        page=page,
        limit=limit,
        status=status,
        branch_id=branch_id
    )
    return jsonify(result)

@app.route('/api/employees/<emp_id>', methods=['GET'])
@login_required
def api_get_employee(emp_id):
    emp = database.get_employee_by_id(emp_id)
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404
    return jsonify(emp)

def extract_employee_permissions(data, default_val=True):
    raw_perms = data.get('permissions')
    if isinstance(raw_perms, str):
        try:
            raw_perms = database.json.loads(raw_perms)
        except Exception:
            raw_perms = {}
    if not isinstance(raw_perms, dict):
        raw_perms = {}

    def to_bool(val, fallback):
        if val is None:
            return fallback
        if isinstance(val, bool):
            return val
        s = str(val).strip().lower()
        if s in ('true', '1', 'yes', 'on'):
            return True
        if s in ('false', '0', 'no', 'off'):
            return False
        return fallback

    return {
        'punch_attendance': to_bool(data.get('perm_punch_attendance', raw_perms.get('punch_attendance')), default_val),
        'attendance_history': to_bool(data.get('perm_attendance_history', raw_perms.get('attendance_history')), default_val),
        'monthly_payslip': to_bool(data.get('perm_monthly_payslip', raw_perms.get('monthly_payslip')), default_val),
        'employee_credentials': to_bool(data.get('perm_employee_credentials', raw_perms.get('employee_credentials')), default_val),
        'leave_permission': to_bool(data.get('perm_leave_permission', raw_perms.get('leave_permission')), default_val)
    }

@app.route('/api/employees', methods=['POST'])
@login_required
def api_create_employee():
    try:
        data = {}
        if request.is_json:
            data = request.get_json() or {}
        else:
            data = request.form.to_dict() or {}
            
        photo_filename = str(data.get('photo_filename') or data.get('photo') or '')
        photo_data_url = str(data.get('photo_data') or '')
        face_registered = False

        # 0. Check if pre-extracted face embedding was passed (from instant camera capture)
        if data.get('face_embedding'):
            try:
                raw_emb = data['face_embedding']
                if isinstance(raw_emb, str):
                    parsed = json.loads(raw_emb)
                    if isinstance(parsed, list) and len(parsed) > 0:
                        data['face_embedding'] = json.dumps(parsed)
                        face_registered = True
                elif isinstance(raw_emb, list) and len(raw_emb) > 0:
                    data['face_embedding'] = json.dumps(raw_emb)
                    face_registered = True
            except Exception as emb_e:
                print("Warning: parsing pre-generated face_embedding:", emb_e)

        if 'photo' in request.files:
            file = request.files['photo']
            if file and file.filename and allowed_file(file.filename):
                try:
                    file_bytes = file.read()
                    if file_bytes:
                        # 1. Optimize photo first (auto-orient, resize to max 600x600, convert to RGB JPEG)
                        unique_filename, photo_data_url, optimized_bytes = process_uploaded_photo(file_bytes, file.filename)
                        photo_filename = unique_filename
                        
                        # 2. Extract 128-d face embedding if not already supplied
                        if not face_registered:
                            fe = get_face_engine()
                            if fe and hasattr(fe, 'extract_face_embedding_from_image'):
                                try:
                                    embedding = fe.extract_face_embedding_from_image(optimized_bytes)
                                    if embedding:
                                        data['face_embedding'] = json.dumps(embedding)
                                        face_registered = True
                                except Exception as fe_err:
                                    print(f"Warning: face embedding extraction error: {fe_err}")
                except Exception as pe:
                    print(f"Warning: photo processing error: {pe}")
                
        # Fallback: if photo was sent as base64 data URI in form body
        if not photo_data_url and data.get('photo_data') and str(data['photo_data']).startswith('data:image'):
            try:
                raw_b64 = str(data['photo_data'])
                header, b64_part = raw_b64.split(',', 1)
                img_bytes = base64.b64decode(b64_part)
                
                # 1. Optimize photo first
                unique_filename, photo_data_url, optimized_bytes = process_uploaded_photo(img_bytes, 'face_capture.jpg')
                photo_filename = unique_filename
                
                # 2. Extract 128-d face embedding if not already supplied
                if not face_registered:
                    fe = get_face_engine()
                    if fe and hasattr(fe, 'extract_face_embedding_from_image'):
                        try:
                            embedding = fe.extract_face_embedding_from_image(optimized_bytes)
                            if embedding:
                                data['face_embedding'] = json.dumps(embedding)
                                face_registered = True
                        except Exception as fe_err:
                            print(f"Warning: face embedding extraction error: {fe_err}")
            except Exception as b64_err:
                print(f"Warning: processing base64 photo_data failed: {b64_err}")

        if photo_filename:
            data['photo_filename'] = photo_filename
            data['photo'] = photo_filename
        if photo_data_url:
            data['photo_data'] = photo_data_url

        if not data.get('employee_name'):
            return jsonify({'error': 'Employee name is required'}), 400
        if not (data.get('email_id') or '').strip():
            return jsonify({'error': 'Email ID is required'}), 400

        # Normalize employee type / field worker flag
        emp_type = str(data.get('employee_type') or '').strip().lower()
        if emp_type:
            data['is_field_worker'] = (emp_type == 'field_worker')
        elif 'is_field_worker' in data:
            data['is_field_worker'] = str(data['is_field_worker']).strip().lower() in ['true', '1', 'yes']
            
        data['permissions'] = extract_employee_permissions(data, default_val=True)
        
        # Process uploaded employee documents
        try:
            data['documents'] = process_uploaded_employee_documents(request.form, request.files, emp_id_prefix='emp')
        except Exception as doc_err:
            print(f"Warning: processing employee documents failed: {doc_err}")
            data['documents'] = []

        if not data.get('branch_id') and get_current_branch_id():
            data['branch_id'] = get_current_branch_id()

        emp_id = database.create_employee(data, company_id=get_current_company_id())
        success_msg = 'Employee created successfully'
        if face_registered:
            success_msg += ' (AI Face Biometrics Registered)'
        return jsonify({
            'success': True,
            'id': emp_id,
            'face_registered': face_registered,
            'message': success_msg
        })
    except ValueError as ve:
        return jsonify({'error': str(ve)}), 400
    except Exception as e:
        err_str = str(e)
        if 'duplicate key' in err_str.lower() or 'e11000' in err_str.lower():
            return jsonify({'error': 'An employee with this ID or email is already registered.'}), 400
        print("Error in api_create_employee:", traceback.format_exc())
        return jsonify({'error': f'Failed to create employee: {err_str}'}), 500

@app.route('/api/employees/<emp_id>', methods=['PUT', 'POST'])
@login_required
def api_update_employee(emp_id):
    try:
        emp = database.get_employee_by_id(emp_id)
        if not emp:
            return jsonify({'error': 'Employee not found'}), 404
            
        data = {}
        if request.is_json:
            data = request.get_json() or {}
        else:
            data = request.form.to_dict() or {}
            
        if 'email_id' in data and not (data.get('email_id') or '').strip():
            return jsonify({'error': 'Email ID is required'}), 400

        # Normalize employee type / field worker flag
        emp_type = str(data.get('employee_type') or '').strip().lower()
        if emp_type:
            data['is_field_worker'] = (emp_type == 'field_worker')
        elif 'is_field_worker' in data:
            data['is_field_worker'] = str(data['is_field_worker']).strip().lower() in ['true', '1', 'yes']
            
        face_registered = False

        # 0. Check if pre-extracted face embedding was passed
        if data.get('face_embedding'):
            try:
                raw_emb = data['face_embedding']
                if isinstance(raw_emb, str):
                    parsed = json.loads(raw_emb)
                    if isinstance(parsed, list) and len(parsed) > 0:
                        data['face_embedding'] = json.dumps(parsed)
                        face_registered = True
                elif isinstance(raw_emb, list) and len(raw_emb) > 0:
                    data['face_embedding'] = json.dumps(raw_emb)
                    face_registered = True
            except Exception as emb_e:
                print("Warning: parsing pre-generated face_embedding in update:", emb_e)

        if 'photo' in request.files:
            file = request.files['photo']
            if file and file.filename and allowed_file(file.filename):
                try:
                    file_bytes = file.read()
                    if file_bytes:
                        # 1. Optimize photo first (auto-orient, resize to max 600x600, convert to RGB JPEG)
                        unique_filename, photo_data_url, optimized_bytes = process_uploaded_photo(file_bytes, file.filename)
                        data['photo_filename'] = unique_filename
                        data['photo'] = unique_filename
                        data['photo_data'] = photo_data_url
                        
                        # 2. Extract 128-d face embedding if not already provided
                        if not face_registered:
                            fe = get_face_engine()
                            if fe and hasattr(fe, 'extract_face_embedding_from_image'):
                                try:
                                    embedding = fe.extract_face_embedding_from_image(optimized_bytes)
                                    if embedding:
                                        data['face_embedding'] = json.dumps(embedding)
                                        face_registered = True
                                except Exception as fe_err:
                                    print(f"Warning: face embedding extraction error: {fe_err}")
                except Exception as pe:
                    print(f"Warning: photo processing error: {pe}")
                
        # Fallback: if photo was sent as base64 data URI in update form body
        if not face_registered and 'photo_data' in data and data['photo_data'] and str(data['photo_data']).startswith('data:image'):
            try:
                raw_b64 = str(data['photo_data'])
                header, b64_part = raw_b64.split(',', 1)
                img_bytes = base64.b64decode(b64_part)
                
                # 1. Optimize photo first
                unique_filename, photo_data_url, optimized_bytes = process_uploaded_photo(img_bytes, 'face_capture.jpg')
                data['photo_filename'] = unique_filename
                data['photo'] = unique_filename
                data['photo_data'] = photo_data_url
                
                # 2. Extract face embedding from optimized image
                fe = get_face_engine()
                if fe and hasattr(fe, 'extract_face_embedding_from_image'):
                    try:
                        embedding = fe.extract_face_embedding_from_image(optimized_bytes)
                        if embedding:
                            data['face_embedding'] = json.dumps(embedding)
                            face_registered = True
                    except Exception as fe_err:
                        print(f"Warning: face embedding extraction from photo_data error: {fe_err}")
            except Exception as b64_err:
                print(f"Warning: processing base64 photo_data in update failed: {b64_err}")

                
        # Update permissions if permission fields or permissions dict was passed
        if any(k in data for k in ['perm_punch_attendance', 'perm_attendance_history', 'perm_monthly_payslip', 'perm_employee_credentials', 'permissions']):
            data['permissions'] = extract_employee_permissions(data, default_val=True)

        # Process existing and new employee documents
        existing_docs = []
        if 'existing_documents' in request.form:
            try:
                raw_docs = request.form.get('existing_documents')
                if raw_docs:
                    parsed = json.loads(raw_docs)
                    if isinstance(parsed, list):
                        existing_docs = parsed
            except Exception as doc_ex:
                print(f"Warning: parsing existing_documents failed: {doc_ex}")
        elif emp and isinstance(emp.get('documents'), list) and not any(k.startswith('doc_') for k in request.files.keys()):
            existing_docs = emp.get('documents', [])

        try:
            new_docs = process_uploaded_employee_documents(request.form, request.files, emp_id_prefix=emp_id)
        except Exception as doc_err:
            print(f"Warning: processing new documents failed: {doc_err}")
            new_docs = []

        if existing_docs or new_docs or 'existing_documents' in request.form:
            data['documents'] = existing_docs + new_docs

        database.update_employee(emp_id, data)
        success_msg = 'Employee updated successfully'
        if face_registered:
            success_msg += ' (AI Face Biometrics Registered)'
        return jsonify({
            'success': True,
            'face_registered': face_registered,
            'message': success_msg
        })
    except ValueError as ve:
        return jsonify({'error': str(ve)}), 400
    except Exception as e:
        err_str = str(e)
        if 'duplicate key' in err_str.lower() or 'e11000' in err_str.lower():
            return jsonify({'error': 'An employee with this ID or email is already registered.'}), 400
        print("Error in api_update_employee:", traceback.format_exc())
        return jsonify({'error': f'Failed to update employee: {err_str}'}), 500

@app.route('/api/employees/<emp_id>', methods=['DELETE'])
@login_required
def api_delete_employee(emp_id):
    if not (session.get('role') in ['super_admin', 'company_admin'] or check_user_can_delete_entries()):
        return jsonify({'error': 'Delete permission is disabled for your company. Please contact System Administrator.'}), 403
    emp = database.get_employee_by_id(emp_id)
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404
    database.delete_employee(emp_id)
    return jsonify({'success': True, 'message': 'Employee deleted successfully'})

@app.route('/api/employees/<emp_id>/pdf', methods=['GET'])
@login_required
def api_employee_pdf(emp_id):
    emp = database.get_employee_by_id(emp_id)
    if not emp:
        return jsonify({'error': 'Employee not found'}), 404
        
    if emp.get('company_id') and emp.get('company_id') != 'ARGUS_MASTER':
        comp = database.get_company_by_id(emp['company_id'])
        if comp:
            emp['company_name'] = comp.get('company_name', 'ARGUS TECHNOLOGIES')
            
    pdf_buffer = pdf_generator.generate_employee_pdf(emp)
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=f"employee_{emp_id}_{emp.get('employee_name', 'details')}.pdf",
        mimetype='application/pdf'
    )

@app.route('/api/employees/export/excel', methods=['GET'])
@app.route('/api/employees/export/csv', methods=['GET'])
@login_required
def api_employees_export_excel():
    search = request.args.get('search', '').strip()
    status = request.args.get('status', '').strip()
    comp_id = get_current_company_id()
    branch_id = request.args.get('branch_id') or get_current_branch_id()
    result = database.get_all_employees(search=search, limit=10000, company_id=comp_id, status=status, branch_id=branch_id)
    
    headers = ['SL NO', 'EMPLOYEE ID', 'EMPLOYEE NAME', 'STATUS', 'SALARY BASIS', 'HOURLY SALARY', 'DAY SALARY', 'HALF DAY SALARY', 'MOBILE NUMBER', 'SHIFT HOURS', 'BANK NAME', 'ACCOUNT NUMBER', 'DATE OF JOINING']
    rows = []
    for idx, emp in enumerate(result['data'], 1):
        st = str(emp.get('salary_type', 'hourly')).lower()
        st_label = 'Hourly' if st == 'hourly' else ('Day-Based' if st == 'daily' else 'Half-Day')
        emp_status = 'Inactive' if str(emp.get('status', 'Active')).strip().capitalize() == 'Inactive' else 'Active'
        rows.append([
            idx,
            emp.get('id', ''),
            emp.get('employee_name', ''),
            emp_status,
            st_label,
            emp.get('hourly_salary', 0),
            emp.get('day_salary', 0),
            emp.get('half_day_salary', 0),
            emp.get('mobile_number', ''),
            emp.get('shift_hours', '08:00'),
            emp.get('bank_name', ''),
            emp.get('account_number', ''),
            emp.get('joining_date', '') or emp.get('date_of_joining', '')
        ])
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    return export_table_data(headers, rows, "employee_directory", export_format=fmt, sheet_name="Employees")

@app.route('/api/employees/export/pdf', methods=['GET'])
@login_required
def api_employees_export_pdf():
    search = request.args.get('search', '').strip()
    status = request.args.get('status', '').strip()
    comp_id = get_current_company_id()
    branch_id = request.args.get('branch_id') or get_current_branch_id()
    result = database.get_all_employees(search=search, limit=10000, company_id=comp_id, status=status, branch_id=branch_id)
    company_info = get_current_company_info()
    pdf_buffer = pdf_generator.generate_employees_pdf(result['data'], company_info=company_info)
    filename = "employee_directory.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

@app.route('/uploads/<filename>')
def uploaded_file(filename):
    upload_folder = app.config['UPLOAD_FOLDER']
    file_path = os.path.join(upload_folder, filename)

    # 1. If file exists on disk, serve it directly
    if os.path.isfile(file_path):
        return send_from_directory(upload_folder, filename)

    # 2. If file missing from disk (e.g. server restart, ephemeral container, hard refresh),
    # recover it permanently from MongoDB Atlas!
    try:
        db = database.get_db()
        emp = db.employees.find_one({
            '$or': [
                {'photo': filename},
                {'photo_filename': filename}
            ]
        })
        if not emp:
            clean_name = os.path.basename(filename)
            import re
            emp = db.employees.find_one({
                '$or': [
                    {'photo': {'$regex': re.escape(clean_name)}},
                    {'photo_filename': {'$regex': re.escape(clean_name)}}
                ]
            })

        if emp and emp.get('photo_data'):
            raw_data = str(emp['photo_data'])
            if ',' in raw_data:
                header, b64_data = raw_data.split(',', 1)
                mime = header.split(';')[0].replace('data:', '') if 'data:' in header else 'image/jpeg'
            else:
                b64_data = raw_data
                mime = 'image/jpeg'
            img_bytes = base64.b64decode(b64_data)

            # Re-cache to disk for fast subsequent requests
            try:
                os.makedirs(upload_folder, exist_ok=True)
                with open(file_path, 'wb') as f:
                    f.write(img_bytes)
            except Exception:
                pass

            return Response(img_bytes, mimetype=mime)

        # Check company_admin for company logos
        comp = db.company_admin.find_one({'logo': filename})
        if not comp:
            clean_name = os.path.basename(filename)
            import re
            comp = db.company_admin.find_one({'logo': {'$regex': re.escape(clean_name)}})
        if comp and comp.get('logo_data'):
            raw_data = str(comp['logo_data'])
            if ',' in raw_data:
                header, b64_data = raw_data.split(',', 1)
                mime = header.split(';')[0].replace('data:', '') if 'data:' in header else 'image/png'
            else:
                b64_data = raw_data
                mime = 'image/png'
            img_bytes = base64.b64decode(b64_data)
            try:
                os.makedirs(upload_folder, exist_ok=True)
                with open(file_path, 'wb') as f:
                    f.write(img_bytes)
            except Exception:
                pass
            return Response(img_bytes, mimetype=mime)
    except Exception as ex:
        print(f"Notice: Image recovery from DB error for {filename}: {ex}")

    return Response(DEFAULT_AVATAR_SVG, mimetype='image/svg+xml')

# ----------------- LIVE & TIMEOUT REPORT API ROUTES ----------------- #

@app.route('/api/live-entries', methods=['GET'])
@login_required
def api_get_live_entries():
    entry_type = request.args.get('type', 'live')
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    
    result = database.get_live_entries(
        tab=entry_type,
        start_date=start_date,
        end_date=end_date,
        search=search,
        page=page,
        limit=limit,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    return jsonify(result)

@app.route('/api/live-entries', methods=['POST'])
@login_required
def api_add_live_entry():
    data = request.get_json() if request.is_json else request.form.to_dict()
    employee_id = data.get('employee_id')
    employee_name = data.get('employee_name')
    if not employee_id or not employee_name:
        return jsonify({'error': 'Employee ID and name are required'}), 400
        
    user_lat = data.get('latitude') or data.get('user_lat')
    user_lng = data.get('longitude') or data.get('user_lng')
    entry_loc = data.get('entry_location') or data.get('live_address')

    inserted_id = database.add_live_entry(
        employee_id=employee_id,
        employee_name=employee_name,
        entry_time=data.get('entry_time'),
        site_name=data.get('site_name', 'OFFICE'),
        entry_location=entry_loc,
        entry_distance=float(data.get('entry_distance', 0.0)),
        is_timeout=int(data.get('is_timeout', 0)),
        user_lat=user_lat,
        user_lng=user_lng,
        company_id=get_current_company_id(),
        live_address=data.get('live_address')
    )
    return jsonify({'success': True, 'id': inserted_id})

@app.route('/api/live-entries/<entry_id>', methods=['DELETE'])
@login_required
def api_delete_live_entry(entry_id):
    if not check_user_can_delete_entries():
        return jsonify({'error': 'Delete permission is disabled for your company. Please contact System Administrator.'}), 403
    entry_type = request.args.get('type', 'live')
    success = database.delete_live_report_entry(entry_id, entry_type=entry_type, company_id=get_current_company_id())
    if not success:
        return jsonify({'error': 'Entry not found or could not be deleted'}), 404
    return jsonify({'success': True, 'message': 'Live report entry deleted successfully'})

@app.route('/api/live-entries/export/excel', methods=['GET'])
@app.route('/api/live-entries/export/csv', methods=['GET'])
@login_required
def api_export_excel():
    entry_type = request.args.get('type', 'live')
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    result = database.get_live_entries(
        tab=entry_type,
        start_date=start_date,
        end_date=end_date,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    
    time_header = 'EXIT TIME' if entry_type == 'timeout' else 'ENTRY TIME'
    loc_header = 'EXIT LOCATION' if entry_type == 'timeout' else 'ENTRY LOCATION'
    dist_header = 'EXIT DISTANCE' if entry_type == 'timeout' else 'ENTRY DISTANCE'
    status_header = 'EXIT STATUS' if entry_type == 'timeout' else 'ENTRY STATUS'
    headers = ['EMPLOYEE NAME', time_header, loc_header, dist_header, status_header]
    rows = []
    for r in result.get('data', []):
        time_val = (r.get('exit_time') if entry_type == 'timeout' and r.get('exit_time') else r.get('entry_time', '')) or '----'
        loc_val = (r.get('exit_location') if entry_type == 'timeout' and r.get('exit_location') else r.get('entry_location', '')) or '----'
        dist_val = (r.get('exit_distance') if entry_type == 'timeout' and r.get('exit_distance') and r.get('exit_distance') != '----' else (r.get('formatted_distance') or (f"OFFICE DISTANCE {r.get('entry_distance', 0)}M" if r.get('entry_distance') is not None else '----')))
        status_val = (r.get('exit_status') if entry_type == 'timeout' else r.get('entry_status')) or '-'
        rows.append([
            r.get('employee_name', ''),
            time_val,
            loc_val,
            dist_val,
            status_val
        ])
        
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    sheet_title = "Timeout Entries" if entry_type == 'timeout' else "Live Entries"
    return export_table_data(headers, rows, f"{entry_type}_entries_report", export_format=fmt, sheet_name=sheet_title)

@app.route('/api/live-entries/export/pdf', methods=['GET'])
@login_required
def api_export_pdf():
    entry_type = request.args.get('type', 'live')
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    result = database.get_live_entries(
        tab=entry_type,
        start_date=start_date,
        end_date=end_date,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    
    title = "Live Entries Report" if entry_type == 'live' else "Timeout Entries Report"
    company_info = get_current_company_info()
    pdf_buffer = pdf_generator.generate_live_report_pdf(title, result['data'], company_info=company_info)
    filename = f"{entry_type}_entries_report.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- UNIVERSAL TAMIL TTS VOICE API ----------------- #

TTS_CACHE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'static', 'cache', 'tts')
os.makedirs(TTS_CACHE_DIR, exist_ok=True)

@app.route('/api/voice/tamil-tts', methods=['GET'])
def api_voice_tamil_tts():
    """
    Synthesizes and streams authentic Tamil voice audio (MP3) using gTTS.
    Implements persistent disk caching so repeated voice messages return in < 10ms.
    Used as an automatic fallback for desktop devices without native Tamil OS voices.
    """
    text = (request.args.get('text') or '').strip()
    lang = request.args.get('lang', 'ta').strip()
    if not text:
        return jsonify({'error': 'No text provided'}), 400

    if len(text) > 300:
        text = text[:300]

    cache_hash = hashlib.md5(f"{text}_{lang}".encode('utf-8')).hexdigest()
    cache_path = os.path.join(TTS_CACHE_DIR, f"{cache_hash}.mp3")

    if not os.path.exists(cache_path):
        try:
            from gtts import gTTS
            tts = gTTS(text=text, lang=lang, slow=False)
            tts.save(cache_path)
        except Exception as e:
            print("gTTS generation error:", e)
            return jsonify({'error': 'TTS synthesis failed', 'details': str(e)}), 500

    return send_file(
        cache_path,
        mimetype='audio/mpeg',
        as_attachment=False,
        max_age=86400 * 7
    )

# ----------------- ZERO-IMAGE FACE RECOGNITION API ----------------- #

@app.route('/api/face/extract-biometrics', methods=['POST'])
@login_required
def api_face_extract_biometrics():
    """
    Extracts 128-dimensional AI face biometrics in real-time during photo capture or file upload.
    Optimizes photo, runs deep neural face detection (YuNet + SFace), and returns the 128-d embedding
    and persistent photo data URL immediately.
    """
    try:
        file_bytes = None
        orig_filename = 'face_capture.jpg'

        if 'photo' in request.files:
            file = request.files['photo']
            if file and file.filename:
                file_bytes = file.read()
                orig_filename = file.filename

        if not file_bytes:
            data = request.get_json(silent=True) or request.form.to_dict()
            photo_data = data.get('photo_data') or data.get('photo') or ''
            if photo_data and 'base64,' in str(photo_data):
                try:
                    _, b64_str = str(photo_data).split('base64,', 1)
                    file_bytes = base64.b64decode(b64_str)
                except Exception as b_err:
                    print("Error decoding base64 photo in extract-biometrics:", b_err)

        if not file_bytes:
            return jsonify({
                'success': False,
                'error': 'No image data received. Please look directly at the camera and take a clear snapshot.',
                'spoken_error': 'No image data received. Please take a snapshot.',
                'spoken_tamil_error': 'படம் பெறப்படவில்லை. கேமராவை நேராகப் பார்த்து மீண்டும் புகைப்படம் எடுக்கவும்.'
            }), 400

        # 1. Optimize photo (600x600 LANCZOS, EXIF transpose, RGB JPEG, persistent base64)
        unique_filename, photo_data_url, optimized_bytes = process_uploaded_photo(file_bytes, orig_filename)

        # 2. Validate face quality & extract 128-d embedding via deep learning engine
        fe = get_face_engine()
        embedding = None
        quality_res = None

        if fe and hasattr(fe, 'validate_registration_face_quality'):
            try:
                # Run quality validation on the original uploaded image to preserve natural aspect ratio & resolution
                quality_res = fe.validate_registration_face_quality(file_bytes)
                if not quality_res.get('valid'):
                    return jsonify({
                        'success': False,
                        'error': quality_res.get('error') or 'Photo did not meet face registration requirements.',
                        'spoken_error': quality_res.get('spoken_error') or 'This image is not accepted. Please upload a clear photo.',
                        'spoken_tamil_error': quality_res.get('spoken_tamil_error') or 'படம் ஏற்கப்படவில்லை. தெளிவான புகைப்படத்தை பதிவேற்றவும்.',
                        'advice': quality_res.get('advice') or '',
                        'code': quality_res.get('code', 'REJECTED'),
                        'metrics': quality_res.get('metrics', {}),
                        'photo_filename': unique_filename,
                        'photo_data': photo_data_url
                    }), 400
                embedding = quality_res.get('embedding')
            except Exception as v_err:
                print("Warning: face quality validation exception:", v_err)

        if not embedding and fe and hasattr(fe, 'extract_face_embedding_from_image'):
            try:
                embedding = fe.extract_face_embedding_from_image(optimized_bytes)
            except Exception as fe_err:
                print("Warning: face embedding extraction error:", fe_err)

        if not embedding:
            return jsonify({
                'success': False,
                'error': 'No clear face detected in the photo. Please ensure your face is well-lit, clearly visible, and facing the camera directly, then try again.',
                'spoken_error': 'This image is not accepted. No face detected. Please face the camera directly in good lighting.',
                'spoken_tamil_error': 'படம் ஏற்கப்படவில்லை. முகம் தெளிவாகத் தெரியவில்லை. வெளிச்சமான இடத்தில் கேமராவை நேராகப் பார்க்கவும்.',
                'code': 'NO_FACE',
                'photo_filename': unique_filename,
                'photo_data': photo_data_url
            }), 400

        return jsonify({
            'success': True,
            'embedding': embedding,
            'photo_filename': unique_filename,
            'photo_data': photo_data_url,
            'vector_length': len(embedding),
            'message': 'AI Face Biometrics verified and registered successfully!',
            'spoken_message': 'Face photo accepted.',
            'spoken_tamil_message': (quality_res and quality_res.get('spoken_tamil_message')) or 'முகம் வெற்றிகரமாக ஏற்றுக்கொள்ளப்பட்டது!'
        })
    except Exception as e:
        print("Error in api_face_extract_biometrics:", traceback.format_exc())
        return jsonify({
            'success': False,
            'error': f'Face biometrics generation failed: {str(e)}',
            'spoken_tamil_error': 'படம் ஏற்கப்படவில்லை. தயவுசெய்து மீண்டும் முயற்சிக்கவும்.'
        }), 500

@app.route('/api/face/enroll', methods=['POST'])
@login_required
def api_face_enroll():
    """
    Enrolls employee face:
    Computes 128-d embedding and discards the image immediately.
    """
    emp_id = request.form.get('employee_id')
    if not emp_id:
        return jsonify({'error': 'Employee ID is required'}), 400
        
    if 'photo' not in request.files:
        return jsonify({'error': 'Face photo frame is required'}), 400
        
    file = request.files['photo']
    file_bytes = file.read()
    
    embedding = get_face_engine().extract_face_embedding_from_image(file_bytes)
    if not embedding:
        return jsonify({'error': 'No face detected in photo. Please ensure your face is clearly visible in good lighting.'}), 400

    database.save_face_embedding(emp_id, embedding)
    
    return jsonify({
        'success': True,
        'message': 'Face biometrics enrolled successfully. Image discarded for privacy.',
        'vector_dimensions': len(embedding)
    })

@app.route('/api/face/recognize', methods=['POST'])
@app.route('/api/recognize_face', methods=['POST'])
def api_face_recognize():
    """
    Recognizes employee face from live camera frame:
    Compares 128-d deep vector with stored embeddings using Cosine Similarity.
    Marks attendance & creates live entry with company-specific geofencing.
    """
    try:
        file_bytes = None
        if 'photo' in request.files:
            file = request.files['photo']
            if file and file.filename:
                file_bytes = file.read()

        burst_file_bytes = None
        data_dict = {}
        if not file_bytes:
            data_dict = request.get_json(silent=True) or request.form.to_dict()
            photo_data = data_dict.get('photo_data') or data_dict.get('photo') or data_dict.get('image') or ''
            if photo_data and 'base64,' in str(photo_data):
                try:
                    _, b64_str = str(photo_data).split('base64,', 1)
                    file_bytes = base64.b64decode(b64_str)
                except Exception:
                    pass
            elif photo_data and isinstance(photo_data, str) and len(photo_data) > 100:
                try:
                    file_bytes = base64.b64decode(photo_data)
                except Exception:
                    pass

        # Optional burst frame for temporal liveness / micro-movement verification
        if 'burst_photo' in request.files:
            b_file = request.files['burst_photo']
            if b_file and b_file.filename:
                burst_file_bytes = b_file.read()
        if not burst_file_bytes:
            if not data_dict:
                data_dict = request.get_json(silent=True) or request.form.to_dict()
            burst_data = data_dict.get('burst_photo') or data_dict.get('burst_data') or ''
            if burst_data and 'base64,' in str(burst_data):
                try:
                    _, b64_b = str(burst_data).split('base64,', 1)
                    burst_file_bytes = base64.b64decode(b64_b)
                except Exception:
                    pass

        if not file_bytes:
            return jsonify({
                'matched': False,
                'confidence': 0.0,
                'message': 'Camera frame is required. Please face the camera and try again.'
            }), 200

        engine = get_face_engine()
        if engine is None:
            return jsonify({
                'matched': False,
                'confidence': 0.0,
                'message': 'Face recognition engine is initializing or unavailable. Please try again in a few moments.'
            }), 200
        
        # Presentation Attack Detection (Anti-Spoofing: detect phone screens, 2D photos, tablets, static attacks)
        try:
            if hasattr(engine, 'verify_liveness_from_bytes'):
                is_live, liveness_conf, spoof_reason = engine.verify_liveness_from_bytes(file_bytes, burst_bytes=burst_file_bytes)
            elif hasattr(engine, 'verify_liveness_anti_spoofing'):
                is_live, liveness_conf, spoof_reason = engine.verify_liveness_anti_spoofing(file_bytes, burst_img=burst_file_bytes)
            else:
                is_live, liveness_conf, spoof_reason = True, 1.0, "OK"

            # Strictly trigger anti-spoofing alert ONLY for confirmed presentation attacks (screens, glass glare, moire, static 2D prints)
            if is_live is False and spoof_reason and "No face detected" not in spoof_reason:
                return jsonify({
                    'matched': False,
                    'confidence': 0.0,
                    'spoof_detected': True,
                    'message': f"Anti-Spoofing Alert: {spoof_reason}. Please face the camera live."
                }), 200
        except Exception as pad_err:
            print(f"Warning: PAD verification notice: {pad_err}")
        
        try:
            query_embedding = engine.extract_face_embedding_from_image(file_bytes)
        except Exception as fe_err:
            traceback.print_exc()
            query_embedding = None

        if not query_embedding:
            return jsonify({
                'matched': False,
                'confidence': 0.0,
                'message': 'No face detected in camera frame. Please face the camera directly in good lighting.'
            }), 200

        # If marked from employee portal, scope search to employee's company
        portal_company_id = session.get('company_id') if session.get('employee_logged_in') else None
        try:
            result = engine.recognize_face(query_embedding, company_id=portal_company_id)
        except Exception as rf_err:
            traceback.print_exc()
            result = {
                'matched': False,
                'confidence': 0.0,
                'message': 'Face recognition matching error. Please try again.'
            }
        
        if result.get('matched'):
            emp_id = result['employee_id']
            emp_name = result['employee_name']
            
            user_lat = request.form.get('latitude')
            user_lng = request.form.get('longitude')
            if str(user_lat).strip().lower() in ['undefined', 'null', '', 'none']:
                user_lat = None
            if str(user_lng).strip().lower() in ['undefined', 'null', '', 'none']:
                user_lng = None
            client_time = request.form.get('client_time')
            live_address = request.form.get('live_address') or request.form.get('location')
            
            # Enrich with employee details for dynamic punch card
            db = database.get_db()
            emp_doc = db.employees.find_one(database.build_id_filter(emp_id)) or db.employees.find_one({'employee_name': emp_name})
            
            # Check company deactivation
            emp_comp_id = (emp_doc.get('company_id') if emp_doc else None) or portal_company_id
            if emp_comp_id and str(emp_comp_id).strip() not in ['ALL', 'ARGUS_MASTER']:
                if not database.is_company_active(emp_comp_id):
                    return jsonify({
                        'matched': False,
                        'message': 'This company account has been deactivated. Biometric attendance cannot be recorded.'
                    }), 200

            # Check if employee status is Inactive
            if emp_doc and str(emp_doc.get('status', 'Active')).strip().capitalize() == 'Inactive':
                return jsonify({
                    'matched': False,
                    'message': f"Attendance punch disabled: Employee profile '{emp_name}' is currently set to Inactive. Contact your company administrator."
                }), 200

            # Check if Punch Attendance is permitted for employee when punching from portal
            if session.get('employee_logged_in') and emp_doc and emp_doc.get('permissions') and emp_doc['permissions'].get('punch_attendance') is False:
                return jsonify({
                    'matched': False,
                    'message': 'Punch attendance is disabled for your account by your administrator.'
                }), 403

            if emp_doc:
                # Dynamic company name lookup
                c_name = emp_doc.get('company_name')
                c_id = emp_doc.get('company_id')
                if not c_name or c_name == 'Argus Technologies':
                    if c_id and c_id != 'ARGUS_MASTER':
                        c_doc = db.company_admin.find_one({'id': str(c_id)})
                        if c_doc and c_doc.get('company_name'):
                            c_name = c_doc['company_name']
                result['employee_code'] = emp_doc.get('employee_id') or emp_doc.get('id') or emp_id
                result['designation'] = emp_doc.get('designation', '') or 'Staff'
                result['department'] = emp_doc.get('department', '') or 'General'
                result['company_name'] = c_name or 'Argus Technologies'
                result['shift_hours'] = emp_doc.get('shift_hours', '08:00')
            else:
                result['employee_code'] = emp_id
                result['designation'] = 'Staff'
                result['department'] = 'General'
                result['company_name'] = 'Argus Technologies'

            # Mark Attendance Punch In / Punch Out Lifecycle & Live Entry (Company-aware, 12-hour format)
            req_branch_id = (request.form.get('branch_id') or data_dict.get('branch_id') or request.args.get('branch_id') or '').strip() or None
            try:
                punch_res = database.record_face_attendance(
                    employee_id=emp_id,
                    employee_name=emp_name,
                    user_lat=user_lat,
                    user_lng=user_lng,
                    company_id=portal_company_id,
                    client_time=client_time,
                    live_address=live_address,
                    branch_id=req_branch_id
                )
                result['live_entry_id'] = str(punch_res.get('live_id', ''))
                result['punch_status'] = str(punch_res.get('status', ''))
                result['already_punched_in'] = bool(punch_res.get('already_punched_in', False))
                result['cooldown_remaining'] = punch_res.get('cooldown_remaining', 0)
                if punch_res.get('message'):
                    result['message'] = str(punch_res.get('message'))
                if punch_res.get('entry_time'):
                    result['entry_time'] = str(punch_res.get('entry_time'))
                result['formatted_distance'] = str(punch_res.get('formatted_dist', ''))
                result['live_location'] = str(punch_res.get('live_location', ''))
                result['entry_location'] = str(punch_res.get('live_location', ''))
                result['punch_time'] = client_time or datetime.now().strftime('%d %b %Y, %I:%M:%S %p')
                if punch_res.get('working_hours'):
                    result['working_hours'] = str(punch_res.get('working_hours'))

                # If this is a Field Worker, record an automatic GPS check-in in Field Worker Module
                if emp_doc and emp_doc.get('is_field_worker'):
                    result['is_field_worker'] = True
                    try:
                        c_office = database.get_company_by_id(emp_comp_id) or {}
                        f_lat = parse_coordinate_to_float(user_lat, c_office.get('latitude', 11.0298))
                        f_lng = parse_coordinate_to_float(user_lng, c_office.get('longitude', 76.9740))
                        f_comp = emp_comp_id or 'ARGUS_MASTER'
                        database.record_field_location_ping({
                            'employee_name': emp_name,
                            'latitude': f_lat,
                            'longitude': f_lng,
                            'accuracy': 15,
                            'activity_note': f"Main Gate AI Camera Punch ({result.get('punch_status', 'Verified')})",
                            'company_id': f_comp
                        }, company_id=f_comp)
                    except Exception as f_ping_err:
                        print(f"Warning: Field worker auto GPS check-in on main camera: {f_ping_err}")
            except Exception as punch_err:
                traceback.print_exc()
                result['live_entry_id'] = ''
                result['punch_status'] = 'punch_in'
                result['formatted_distance'] = 'OFFICE DISTANCE 0.0M'
                result['live_location'] = 'Company Office'
                result['entry_location'] = 'Company Office'
                result['punch_time'] = client_time or datetime.now().strftime('%d %b %Y, %I:%M:%S %p')
            
        clean_result = database.clean_doc(result)
        return jsonify(clean_result)

    except Exception as e:
        traceback.print_exc()
        return jsonify({
            'matched': False,
            'message': f'Face recognition processing error: {str(e)}. Please face the camera directly in good lighting and try again.'
        }), 200

# ----------------- ATTENDANCE REPORTS API ROUTES ----------------- #

@app.route('/api/attendance-reports', methods=['GET'])
@login_required
def api_get_attendance_reports():
    report_type = request.args.get('type', 'all')
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    employee = request.args.get('employee', 'All').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    
    result = database.get_attendance_reports(
        report_type=report_type,
        start_date=start_date,
        end_date=end_date,
        employee=employee,
        search=search,
        page=page,
        limit=limit,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    return jsonify(result)

@app.route('/api/attendance-reports/compact', methods=['GET'])
@app.route('/api/attendance-reports/simple', methods=['GET'])
@login_required
def api_get_attendance_simple():
    employee = request.args.get('employee', 'All Employees').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    
    result = database.get_attendance_simple_table(
        employee=employee,
        start_date=start_date,
        end_date=end_date,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    return jsonify(result)

@app.route('/api/attendance-reports/update', methods=['POST'])
@login_required
def api_update_attendance_records():
    return jsonify({'success': True, 'message': 'Attendance records updated successfully'})

@app.route('/api/attendance-reports/<report_id>', methods=['DELETE'])
@login_required
def api_delete_attendance_report(report_id):
    if not check_user_can_delete_entries():
        return jsonify({'error': 'Delete permission is disabled for your company. Please contact System Administrator.'}), 403
    is_manual = request.args.get('is_manual', 'false').lower() in ['true', '1', 'yes']
    success = database.delete_attendance_report(report_id, is_manual=is_manual, company_id=get_current_company_id())
    if not success:
        return jsonify({'error': 'Attendance record not found or could not be deleted'}), 404
    return jsonify({'success': True, 'message': 'Attendance record deleted successfully'})

@app.route('/api/attendance-reports/export/excel', methods=['GET'])
@app.route('/api/attendance-reports/export/csv', methods=['GET'])
@login_required
def api_attendance_export_excel():
    report_type = request.args.get('type', 'all')
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    employee = request.args.get('employee', 'All').strip()
    comp_id = get_current_company_id()
    branch_id = get_current_branch_id()
    
    rows = []
    if report_type in ['simple', 'compact']:
        result = database.get_attendance_simple_table(employee=employee, start_date=start_date, end_date=end_date, company_id=comp_id, branch_id=branch_id)
        headers = ['EMPLOYEE NAME', 'ENTRY TIME', 'EXIT TIME', 'WORKING HOURS', 'SHIFT VARIANCE', 'WORKING SALARY', 'STATUS']
        for r in result.get('data', []):
            rows.append([
                r.get('employee_name', ''),
                r.get('entry_time', ''),
                r.get('exit_time', '') or '',
                r.get('working_hours', '00:00') or '00:00',
                r.get('shift_variance', '') or '',
                r.get('working_salary', 0),
                r.get('status', 'Manual')
            ])
        rows.append(['TOTAL', '', '', result.get('total_working_hours', '00:00'), '----', result.get('total_working_salary', '0.00'), ''])
    else:
        result = database.get_attendance_reports(report_type=report_type, start_date=start_date, end_date=end_date, employee=employee, limit=10000, company_id=comp_id, branch_id=branch_id)
        headers = ['EMPLOYEE NAME', 'ENTRY TIME', 'ENTRY DISTANCE', 'ENTRY LOCATION', 'ENTRY STATUS', 'EXIT TIME', 'EXIT DISTANCE', 'EXIT LOCATION', 'EXIT STATUS', 'WORKING HOURS', 'SHIFT VARIANCE', 'WORKING SALARY']
        for r in result.get('data', []):
            rows.append([
                r.get('employee_name', ''),
                r.get('entry_time', '') or '',
                r.get('entry_distance', '----') or '----',
                r.get('entry_location', '----') or '----',
                r.get('entry_status', '-') or '-',
                r.get('exit_time', '----') or '----',
                r.get('exit_distance', '----') or '----',
                r.get('exit_location', '----') or '----',
                r.get('exit_status', '-') or '-',
                r.get('working_hours', '00:00') or '00:00',
                r.get('shift_variance', '----') or '----',
                r.get('working_salary', 0) if r.get('working_salary') is not None else 0
            ])
            
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    return export_table_data(headers, rows, f"attendance_{report_type}_report", export_format=fmt, sheet_name="Attendance")

@app.route('/api/attendance-reports/export/pdf', methods=['GET'])
@login_required
def api_attendance_export_pdf():
    report_type = request.args.get('type', 'all')
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    employee = request.args.get('employee', 'All').strip()
    comp_id = get_current_company_id()
    branch_id = get_current_branch_id()
    
    company_info = get_current_company_info()
    if report_type in ['simple', 'compact']:
        result = database.get_attendance_simple_table(employee=employee, start_date=start_date, end_date=end_date, company_id=comp_id, branch_id=branch_id)
        emp_title = f" - {employee}" if (employee and employee not in ['All', 'All Employees']) else ""
        date_title = f" ({start_date} to {end_date})" if (start_date and end_date) else (f" (From {start_date})" if start_date else (f" (Up to {end_date})" if end_date else ""))
        title = f"Attendance Report (Compact View){emp_title}{date_title}"
        totals = {
            'total_working_hours': result['total_working_hours'],
            'total_working_salary': result['total_working_salary']
        }
        pdf_buffer = pdf_generator.generate_attendance_report_pdf(title, result['data'], is_simple=True, totals=totals, company_info=company_info)
    else:
        result = database.get_attendance_reports(report_type=report_type, start_date=start_date, end_date=end_date, employee=employee, limit=10000, company_id=comp_id, branch_id=branch_id)
        titles = {
            'all': 'All Attendance Entries',
            'proper': 'Proper Attendance Entries',
            'improper': 'Improper Attendance Entries',
            'manual': 'Manual Attendance Entries'
        }
        title = titles.get(report_type, 'Attendance Report')
        pdf_buffer = pdf_generator.generate_attendance_report_pdf(title, result['data'], is_simple=False, company_info=company_info)
        
    filename = f"attendance_{report_type}_report.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- MANUAL ENTRIES API ROUTES ----------------- #

@app.route('/api/manual-entries', methods=['GET'])
@login_required
def api_get_manual_entries():
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()
    status = request.args.get('status', 'All').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    
    result = database.get_manual_entries(
        from_date=from_date,
        to_date=to_date,
        status=status,
        search=search,
        page=page,
        limit=limit,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    return jsonify(result)

@app.route('/api/manual-entries/<entry_id>', methods=['GET'])
@login_required
def api_get_manual_entry(entry_id):
    entry = database.get_manual_entry_by_id(entry_id)
    if not entry:
        return jsonify({'error': 'Manual entry not found'}), 404
    return jsonify(entry)

@app.route('/api/manual-entries', methods=['POST'])
@login_required
def api_create_manual_entry():
    data = request.get_json() if request.is_json else request.form.to_dict()
    if not data.get('employee_name'):
        return jsonify({'error': 'Employee name is required'}), 400
        
    bid = get_current_branch_id()
    if bid and 'branch_id' not in data:
        data['branch_id'] = bid

    inserted_id = database.create_manual_entry(data, company_id=get_current_company_id())
    return jsonify({'success': True, 'id': inserted_id, 'message': 'Manual entry created successfully'})

@app.route('/api/manual-entries/<entry_id>', methods=['PUT', 'POST'])
@login_required
def api_update_manual_entry(entry_id):
    entry = database.get_manual_entry_by_id(entry_id)
    if not entry:
        return jsonify({'error': 'Manual entry not found'}), 404
        
    data = request.get_json() if request.is_json else request.form.to_dict()
    database.update_manual_entry(entry_id, data)
    return jsonify({'success': True, 'message': 'Manual entry updated successfully'})

@app.route('/api/manual-entries/<entry_id>', methods=['DELETE'])
@login_required
def api_delete_manual_entry(entry_id):
    if not check_user_can_delete_entries():
        return jsonify({'error': 'Delete permission is disabled for your company. Please contact System Administrator.'}), 403
    entry = database.get_manual_entry_by_id(entry_id)
    if not entry:
        return jsonify({'error': 'Manual entry not found'}), 404
        
    database.delete_manual_entry(entry_id)
    return jsonify({'success': True, 'message': 'Manual entry deleted successfully'})

@app.route('/api/manual-entries/export/excel', methods=['GET'])
@app.route('/api/manual-entries/export/csv', methods=['GET'])
@login_required
def api_manual_entries_export_excel():
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()
    status = request.args.get('status', 'All').strip()
    
    result = database.get_manual_entries(
        from_date=from_date,
        to_date=to_date,
        status=status,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    
    headers = ['NAME', 'DATE', 'HOURS', 'STATUS', 'SUBMITTED', 'HOURLY', 'DAY', 'HALF', 'WORKING SALARY']
    rows = []
    for r in result['data']:
        sal = f"+{int(float(r.get('working_salary', 0)))}" if float(r.get('working_salary', 0)) > 0 else "0"
        st_val = r.get('status', '')
        rs_val = (r.get('reason') or '').strip()
        status_display = f"{st_val} ({rs_val})" if rs_val else st_val
        rows.append([
            r.get('employee_name', ''), r.get('entry_date', ''), r.get('hours', ''), status_display,
            r.get('submitted_at', ''), int(float(r.get('hourly_rate', 0))), int(float(r.get('day_rate', 0))),
            int(float(r.get('half_rate', 0))), sal
        ])
        
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    return export_table_data(headers, rows, "manual_entries_report", export_format=fmt, sheet_name="Manual Entries")

@app.route('/api/manual-entries/export/pdf', methods=['GET'])
@login_required
def api_manual_entries_export_pdf():
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()
    status = request.args.get('status', 'All').strip()
    
    result = database.get_manual_entries(
        from_date=from_date,
        to_date=to_date,
        status=status,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    company_info = get_current_company_info()
    pdf_buffer = pdf_generator.generate_manual_entries_pdf(result['data'], company_info=company_info)
    filename = "manual_entries_report.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- PAYMENT MANAGEMENT API ROUTES ----------------- #

RECEIPTS_FOLDER = os.path.join(app.config['UPLOAD_FOLDER'], 'receipts')
try:
    os.makedirs(RECEIPTS_FOLDER, exist_ok=True)
except Exception:
    pass

@app.route('/api/payments', methods=['GET'])
@login_required
def api_get_payments():
    employee = request.args.get('employee', 'All').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    bank = request.args.get('bank', 'All').strip()
    payment_type = request.args.get('payment_type', 'All').strip()
    status = (request.args.get('status') or request.args.get('reason') or 'All').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    
    result = database.get_payments(
        employee=employee,
        start_date=start_date,
        end_date=end_date,
        bank=bank,
        payment_type=payment_type,
        status=status,
        search=search,
        page=page,
        limit=limit,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    return jsonify(result)

@app.route('/api/payments/<payment_id>', methods=['GET'])
@login_required
def api_get_payment(payment_id):
    p = database.get_payment_by_id(payment_id)
    if not p:
        return jsonify({'error': 'Payment not found'}), 404
    return jsonify(p)

@app.route('/api/payments', methods=['POST'])
@login_required
def api_create_payment():
    data = {}
    if request.is_json:
        data = request.get_json()
    else:
        data = request.form.to_dict()
        
    if 'receipt' in request.files:
        file = request.files['receipt']
        if file and file.filename:
            fname = werkzeug.utils.secure_filename(file.filename)
            uniq = f"{int(database.time.time())}_{fname}"
            file.save(os.path.join(RECEIPTS_FOLDER, uniq))
            data['receipt_filename'] = uniq
            
    if not data.get('employee_name'):
        return jsonify({'error': 'Employee name is required'}), 400
        
    bid = get_current_branch_id()
    if bid and 'branch_id' not in data:
        data['branch_id'] = bid

    inserted_id = database.create_payment(data, company_id=get_current_company_id())
    return jsonify({'success': True, 'id': inserted_id, 'message': 'Payment created successfully'})

@app.route('/api/payments/<payment_id>', methods=['PUT', 'POST'])
@login_required
def api_update_payment(payment_id):
    p = database.get_payment_by_id(payment_id)
    if not p:
        return jsonify({'error': 'Payment not found'}), 404
        
    data = {}
    if request.is_json:
        data = request.get_json()
    else:
        data = request.form.to_dict()
        
    if 'receipt' in request.files:
        file = request.files['receipt']
        if file and file.filename:
            fname = werkzeug.utils.secure_filename(file.filename)
            uniq = f"{int(database.time.time())}_{fname}"
            file.save(os.path.join(RECEIPTS_FOLDER, uniq))
            data['receipt_filename'] = uniq
            
    database.update_payment(payment_id, data)
    return jsonify({'success': True, 'message': 'Payment updated successfully'})

@app.route('/api/payments/<payment_id>', methods=['DELETE'])
@login_required
def api_delete_payment(payment_id):
    if not check_user_can_delete_entries():
        return jsonify({'error': 'Delete permission is disabled for your company. Please contact System Administrator.'}), 403
    p = database.get_payment_by_id(payment_id)
    if not p:
        return jsonify({'error': 'Payment not found'}), 404
        
    database.delete_payment(payment_id)
    return jsonify({'success': True, 'message': 'Payment deleted successfully'})

@app.route('/api/payments/export/excel', methods=['GET'])
@app.route('/api/payments/export/csv', methods=['GET'])
@login_required
def api_payments_export_excel():
    employee = request.args.get('employee', 'All').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    bank = request.args.get('bank', 'All').strip()
    payment_type = request.args.get('payment_type', 'All').strip()
    status = (request.args.get('status') or request.args.get('reason') or 'All').strip()
    
    result = database.get_payments(
        employee=employee,
        start_date=start_date,
        end_date=end_date,
        bank=bank,
        payment_type=payment_type,
        status=status,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    
    headers = ['TIMESTAMP', 'EMPLOYEE', 'DATE', 'AMOUNT', 'BANK', 'PAYMENT TYPE', 'STATUS']
    rows = []
    for r in result['data']:
        st = r.get('status') or r.get('reason') or 'Payment'
        rs = r.get('reason') if r.get('status') else ''
        disp_st = f"{st} - {rs}" if rs else st
        rows.append([
            r.get('timestamp', ''), r.get('employee_name', ''), r.get('payment_date', ''),
            float(r.get('amount', 0)), r.get('bank') or '-',
            r.get('payment_type', ''), disp_st
        ])
        
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    return export_table_data(headers, rows, "payment_management_report", export_format=fmt, sheet_name="Payments")

@app.route('/api/payments/export/pdf', methods=['GET'])
@login_required
def api_payments_export_pdf():
    employee = request.args.get('employee', 'All').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    bank = request.args.get('bank', 'All').strip()
    payment_type = request.args.get('payment_type', 'All').strip()
    status = (request.args.get('status') or request.args.get('reason') or 'All').strip()
    
    result = database.get_payments(
        employee=employee,
        start_date=start_date,
        end_date=end_date,
        bank=bank,
        payment_type=payment_type,
        status=status,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    company_info = get_current_company_info()
    pdf_buffer = pdf_generator.generate_payments_pdf(result['data'], company_info=company_info)
    filename = "payment_management_report.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- ADVANCE MANAGEMENT API ROUTES ----------------- #

@app.route('/api/advances', methods=['GET'])
@login_required
def api_get_advances():
    employee = request.args.get('employee', 'All').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    
    result = database.get_advances(
        employee=employee,
        start_date=start_date,
        end_date=end_date,
        search=search,
        page=page,
        limit=limit,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    return jsonify(result)

@app.route('/api/advances/<advance_id>', methods=['GET'])
@login_required
def api_get_advance(advance_id):
    a = database.get_advance_by_id(advance_id)
    if not a:
        return jsonify({'error': 'Advance not found'}), 404
    return jsonify(a)

@app.route('/api/advances', methods=['POST'])
@login_required
def api_create_advance():
    data = request.get_json() if request.is_json else request.form.to_dict()
    if not data.get('employee_name'):
        return jsonify({'error': 'Employee name is required'}), 400
        
    bid = get_current_branch_id()
    if bid and 'branch_id' not in data:
        data['branch_id'] = bid

    inserted_id = database.create_advance(data, company_id=get_current_company_id())
    return jsonify({'success': True, 'id': inserted_id, 'message': 'Advance created successfully'})

@app.route('/api/advances/<advance_id>', methods=['PUT', 'POST'])
@login_required
def api_update_advance(advance_id):
    a = database.get_advance_by_id(advance_id)
    if not a:
        return jsonify({'error': 'Advance not found'}), 404
        
    data = request.get_json() if request.is_json else request.form.to_dict()
    database.update_advance(advance_id, data)
    return jsonify({'success': True, 'message': 'Advance updated successfully'})

@app.route('/api/advances/<advance_id>', methods=['DELETE'])
@login_required
def api_delete_advance(advance_id):
    if not check_user_can_delete_entries():
        return jsonify({'error': 'Delete permission is disabled for your company. Please contact System Administrator.'}), 403
    a = database.get_advance_by_id(advance_id)
    if not a:
        return jsonify({'error': 'Advance not found'}), 404
        
    database.delete_advance(advance_id)
    return jsonify({'success': True, 'message': 'Advance deleted successfully'})

@app.route('/api/advances/export/excel', methods=['GET'])
@app.route('/api/advances/export/csv', methods=['GET'])
@login_required
def api_advances_export_excel():
    employee = request.args.get('employee', 'All').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    
    result = database.get_advances(
        employee=employee,
        start_date=start_date,
        end_date=end_date,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    
    headers = ['TIMESTAMP', 'EMPLOYEE', 'DATE', 'AMOUNT']
    rows = []
    for r in result['data']:
        rows.append([
            r.get('timestamp', ''), r.get('employee_name', ''), r.get('advance_date', ''),
            float(r.get('amount', 0))
        ])
        
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    return export_table_data(headers, rows, "advance_management_report", export_format=fmt, sheet_name="Advances")

@app.route('/api/advances/export/pdf', methods=['GET'])
@login_required
def api_advances_export_pdf():
    employee = request.args.get('employee', 'All').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    
    result = database.get_advances(
        employee=employee,
        start_date=start_date,
        end_date=end_date,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    company_info = get_current_company_info()
    pdf_buffer = pdf_generator.generate_advances_pdf(result['data'], company_info=company_info)
    filename = "advance_management_report.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- BALANCE REPORT API ROUTES ----------------- #

@app.route('/api/advance-summary', methods=['GET'])
@app.route('/api/balance-report', methods=['GET'])
@login_required
def api_get_balance_report():
    employee = request.args.get('employee', 'All').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    result = database.get_balance_report(
        employee=employee,
        search=search,
        page=page,
        limit=limit,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    return jsonify(result)

@app.route('/api/advance-summary/export/excel', methods=['GET'])
@app.route('/api/advance-summary/export/csv', methods=['GET'])
@app.route('/api/balance-report/export/excel', methods=['GET'])
@app.route('/api/balance-report/export/csv', methods=['GET'])
@login_required
def api_balance_report_export_excel():
    employee = request.args.get('employee', 'All').strip()
    result = database.get_balance_report(employee=employee, limit=10000, company_id=get_current_company_id(), branch_id=get_current_branch_id())
    
    headers = ['DATE&TIME', 'NAME', 'DATE', 'ADVANCE AMOUNT', 'REPAYMENT AMOUNT', 'BALANCE AMOUNT']
    rows = []
    for r in result['data']:
        pay_amt = r.get('payment_amount', r.get('advance_repayment_amount', 0.0))
        rows.append([
            r.get('timestamp', ''), r.get('name', ''), r.get('date', ''),
            float(r.get('advance_amount', 0)),
            float(pay_amt),
            float(r.get('balance_amount', 0))
        ])
    rows.append([])
    rows.append(['TOTAL ADVANCE', float(result.get('total_advance', 0)), 'TOTAL REPAYMENT', float(result.get('total_payment', result.get('total_repayment', 0.0))), 'BALANCE AMOUNT', float(result.get('balance_amount', 0))])
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    return export_table_data(headers, rows, "advance_summary", export_format=fmt, sheet_name="Advance Summary")

@app.route('/api/advance-summary/export/pdf', methods=['GET'])
@app.route('/api/balance-report/export/pdf', methods=['GET'])
@login_required
def api_balance_report_export_pdf():
    employee = request.args.get('employee', 'All').strip()
    result = database.get_balance_report(employee=employee, limit=10000, company_id=get_current_company_id(), branch_id=get_current_branch_id())
    totals = {
        'total_advance': result['total_advance'],
        'total_repayment': result.get('total_payment', result.get('total_repayment', 0.0)),
        'total_payment': result.get('total_payment', result.get('total_repayment', 0.0)),
        'balance_amount': result['balance_amount']
    }
    company_info = get_current_company_info()
    pdf_buffer = pdf_generator.generate_balance_report_pdf(result['data'], totals, company_info=company_info)
    filename = "advance_summary.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- MONTHLY PAYSLIP API ROUTES ----------------- #

@app.route('/api/payslip/generate', methods=['GET', 'POST'])
@login_required
def api_generate_payslip():
    if request.method == 'POST':
        data = request.get_json() if request.is_json else request.form.to_dict()
        employee = (data.get('employee_name') or data.get('employee') or '').strip()
        month = data.get('month', '').strip()
        branch_id = data.get('branch_id') or request.args.get('branch_id') or get_current_branch_id()
    else:
        employee = (request.args.get('employee_name') or request.args.get('employee') or '').strip()
        month = request.args.get('month', '').strip()
        branch_id = request.args.get('branch_id') or get_current_branch_id()
        
    if not employee:
        return jsonify({'error': 'Employee name is required'}), 400
    if not month:
        month = datetime.now().strftime("%Y-%m")
        
    comp_id = get_current_company_id()
    if session.get('role') == 'super_admin':
        comp_id = None
    payslip = database.get_payslip_data(employee, month, company_id=comp_id, branch_id=branch_id)
    database.save_generated_salary_report(payslip, company_id=payslip.get('company_id') or comp_id)
    return jsonify(payslip)

@app.route('/api/payslip/export/pdf', methods=['GET'])
def api_export_payslip_pdf():
    """Generates and downloads or views payslip PDF. Available to Admins and Employees."""
    employee = (request.args.get('employee_name') or request.args.get('employee') or '').strip()
    month = request.args.get('month', '').strip()
    branch_id = request.args.get('branch_id') or get_current_branch_id()
    
    if not month:
        month = datetime.now().strftime("%Y-%m")
    
    comp_id = get_current_company_id()
    if session.get('role') == 'super_admin':
        comp_id = None
    if not employee:
        emp_res = database.get_all_employees(company_id=comp_id or 'ALL', branch_id=branch_id, limit=1)
        if emp_res['data']:
            employee = emp_res['data'][0]['employee_name']
        else:
            employee = 'Employee'
    
    template = request.args.get('template', 'template_1').strip()
    theme = request.args.get('theme', 'navy').strip()
    payslip = database.get_payslip_data(employee, month, company_id=comp_id, branch_id=branch_id)
    pdf_buffer = pdf_generator.generate_payslip_pdf(payslip, template_id=template, theme=theme)
    filename = f"payslip_{employee}_{month}.pdf"

    view_inline = (request.args.get('view') == '1' or request.args.get('inline') == '1')
    return send_file(
        pdf_buffer,
        as_attachment=not view_inline,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- SALARY REPORTS API ROUTES ----------------- #

@app.route('/api/salary-reports', methods=['GET'])
@login_required
def api_get_salary_reports():
    start_month = request.args.get('start_month', '').strip()
    end_month = request.args.get('end_month', '').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))
    
    result = database.get_salary_reports(
        start_month=start_month,
        end_month=end_month,
        search=search,
        page=page,
        limit=limit,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    return jsonify(result)

@app.route('/api/salary-reports/export/excel', methods=['GET'])
@app.route('/api/salary-reports/export/csv', methods=['GET'])
@app.route('/api/salary/export/excel', methods=['GET'])
@app.route('/api/salary/export/csv', methods=['GET'])
@login_required
def api_salary_reports_export_excel():
    start_month = request.args.get('start_month', '').strip()
    end_month = request.args.get('end_month', '').strip()
    
    result = database.get_salary_reports(
        start_month=start_month,
        end_month=end_month,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    
    headers = [
        'EMPLOYEE NAME', 'PAY PERIOD', 'WORKING DAYS', 'WORKING HOURS',
        'ALLOWANCE', 'INCENTIVE', 'OTHER EARNINGS', 'BASIC EARNINGS',
        'TOTAL EARNINGS', 'PAID SALARY', 'ADVANCE REPAYMENT', 'OTHER DEDUCTIONS',
        'TOTAL DEDUCTIONS', 'NET SALARY'
    ]
    rows = []
    for r in result['data']:
        rows.append([
            r.get('employee_name', ''), r.get('pay_period', ''), r.get('working_days', ''), r.get('working_hours', ''),
            int(float(r.get('allowance', 0))), int(float(r.get('incentive', 0))), int(float(r.get('other_earnings', 0))),
            int(float(r.get('basic_earnings', 0))), int(float(r.get('total_earnings', 0))), int(float(r.get('paid_salary', 0))),
            int(float(r.get('advance_repayment', 0))), int(float(r.get('other_deductions', 0))), int(float(r.get('total_deductions', 0))),
            int(float(r.get('net_salary', 0)))
        ])
        
    fmt = 'csv' if request.path.endswith('/csv') else 'excel'
    return export_table_data(headers, rows, "salary_report", export_format=fmt, sheet_name="Salary Report")

@app.route('/api/salary-reports/export/pdf', methods=['GET'])
@login_required
def api_salary_reports_export_pdf():
    start_month = request.args.get('start_month', '').strip()
    end_month = request.args.get('end_month', '').strip()
    
    result = database.get_salary_reports(
        start_month=start_month,
        end_month=end_month,
        limit=10000,
        company_id=get_current_company_id(),
        branch_id=get_current_branch_id()
    )
    company_info = get_current_company_info()
    pdf_buffer = pdf_generator.generate_salary_report_pdf(result['data'], company_info=company_info)
    filename = "salary_report.pdf"
    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/pdf'
    )

# ----------------- AUTOMATED COMPANY EMAIL REPORTING ROUTES ----------------- #

@app.route('/api/admin/reports/send-daily', methods=['POST'])
@super_admin_required
def api_admin_send_daily_reports():
    """Manually triggers or tests yesterday's daily activity reports to companies."""
    data = request.get_json() if request.is_json else request.form.to_dict()
    company_id = (data.get('company_id') or '').strip()
    target_date = (data.get('date') or '').strip() or None
    force = bool(data.get('force', False))

    import report_scheduler
    if company_id and company_id != 'ALL':
        ok, msg, report = report_scheduler.send_daily_activity_email(company_id, target_date=target_date, force=force)
        return jsonify({
            'success': ok,
            'message': msg,
            'company_id': company_id,
            'report': report
        })
    else:
        results = report_scheduler.dispatch_all_daily_reports(target_date=target_date, force=force)
        return jsonify({
            'success': True,
            'total_companies': len(results),
            'results': results
        })

@app.route('/api/admin/reports/send-monthly', methods=['POST'])
@super_admin_required
def api_admin_send_monthly_reports():
    """Manually triggers or tests monthly payroll reports to companies."""
    data = request.get_json() if request.is_json else request.form.to_dict()
    company_id = (data.get('company_id') or '').strip()
    target_month = (data.get('month') or '').strip() or None
    force = bool(data.get('force', False))

    import report_scheduler
    if company_id and company_id != 'ALL':
        ok, msg, report = report_scheduler.send_monthly_salary_email(company_id, target_month=target_month, force=force)
        return jsonify({
            'success': ok,
            'message': msg,
            'company_id': company_id,
            'report': report
        })
    else:
        results = report_scheduler.dispatch_all_monthly_reports(target_month=target_month, force=force)
        return jsonify({
            'success': True,
            'total_companies': len(results),
            'results': results
        })

@app.route('/api/admin/reports/email-logs', methods=['GET'])
@super_admin_required
def api_admin_email_logs():
    """Retrieves recent email dispatch audit logs from db.email_logs."""
    db = database.get_db()
    limit = int(request.args.get('limit', 50))
    company_id = request.args.get('company_id', '').strip()
    q = {}
    if company_id and company_id != 'ALL':
        q['company_id'] = company_id
    logs = list(db.email_logs.find(q).sort('_id', -1).limit(limit))
    return jsonify({'logs': database.clean_doc(logs)})

# ==============================================================================
# SUPPORT AND SUPPORT TICKETS ROUTES & APIS
# ==============================================================================

SUPPORT_UPLOADS_FOLDER = os.path.join(app.config['UPLOAD_FOLDER'], 'support_tickets')
try:
    os.makedirs(SUPPORT_UPLOADS_FOLDER, exist_ok=True)
except Exception as e:
    print(f"Warning: Could not create support_tickets upload folder: {e}")

ALLOWED_SUPPORT_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp', 'pdf', 'xls', 'xlsx', 'csv', 'doc', 'docx', 'txt'}

def allowed_support_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_SUPPORT_EXTENSIONS

def _format_file_size(bytes_size):
    if bytes_size < 1024:
        return f"{bytes_size} B"
    elif bytes_size < 1024 * 1024:
        return f"{round(bytes_size / 1024, 1)} KB"
    else:
        return f"{round(bytes_size / (1024 * 1024), 1)} MB"

@app.route('/support')
@login_required
def support_center():
    """Renders the Company Support Center page with 3 cards, raise ticket, my tickets, and demo videos."""
    comp_id = session.get('company_id') or 'ARGUS_MASTER'
    if session.get('role') == 'company_admin' and not database.is_company_support_enabled(comp_id):
        flash('Support & Help is currently disabled for your organization. Please contact System Administrator.', 'warning')
        return redirect(url_for('dashboard'))

    comp_name = session.get('company_name')
    if not comp_name:
        comp_doc = database.get_company_by_id(comp_id)
        comp_name = comp_doc.get('company_name', 'Company') if comp_doc else 'Company'
    
    current_ist = database.get_ist_now().strftime('%d/%m/%Y %I:%M %p')
    return render_template('support_company.html',
                           active_tab='SUPPORT & HELP',
                           company_id=comp_id,
                           company_name=comp_name,
                           current_time_str=current_ist)

@app.route('/admin/support-tickets')
@super_admin_required
def admin_support_tickets():
    """Renders the System Admin Support Tickets Dashboard."""
    companies = database.get_all_companies()
    return render_template('support_admin.html',
                           active_tab='SUPPORT TICKETS',
                           companies=companies)

@app.route('/api/support/tickets', methods=['GET'])
@login_required
def api_company_support_tickets():
    """Returns tickets for the logged-in company admin."""
    comp_id = session.get('company_id') or 'ARGUS_MASTER'
    status_filter = request.args.get('status', 'all')
    search = request.args.get('search', '')
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))

    result = database.get_company_support_tickets(
        company_id=comp_id,
        status=status_filter,
        search=search,
        page=page,
        limit=limit
    )
    return jsonify({
        'success': True,
        'tickets': result['tickets'],
        'counts': result['counts'],
        'total': result['total'],
        'page': result['page'],
        'limit': result['limit']
    })

@app.route('/api/support/ticket', methods=['POST'])
@login_required
def api_create_support_ticket():
    """Creates a new support ticket with attachments."""
    comp_id = session.get('company_id') or 'ARGUS_MASTER'
    comp_name = session.get('company_name')
    if not comp_name:
        comp_doc = database.get_company_by_id(comp_id)
        comp_name = comp_doc.get('company_name', 'Company') if comp_doc else 'Company'

    subject = request.form.get('subject', '').strip()
    description = request.form.get('description', '').strip()

    if not subject or not description:
        return jsonify({'success': False, 'error': 'Subject and Description are required.'}), 400

    raised_by = session.get('user_name') or f"Admin ({comp_name[:10]})"
    role = session.get('role', 'company_admin')

    # Process attachments
    attachments = []
    comp_upload_dir = os.path.join(SUPPORT_UPLOADS_FOLDER, str(comp_id))
    os.makedirs(comp_upload_dir, exist_ok=True)

    uploaded_files = request.files.getlist('attachments')
    for f in uploaded_files:
        if f and f.filename and allowed_support_file(f.filename):
            clean_filename = werkzeug.utils.secure_filename(f.filename) or 'file'
            uniq_filename = f"{int(time.time())}_{uuid.uuid4().hex[:6]}_{clean_filename}"
            target_path = os.path.join(comp_upload_dir, uniq_filename)
            f.save(target_path)

            file_size = os.path.getsize(target_path) if os.path.exists(target_path) else 0
            ext = clean_filename.rsplit('.', 1)[1].lower() if '.' in clean_filename else ''

            attachments.append({
                'id': f"att_{uuid.uuid4().hex[:8]}",
                'filename': uniq_filename,
                'original_name': f.filename,
                'file_path': target_path,
                'file_size_bytes': file_size,
                'file_size_display': _format_file_size(file_size),
                'extension': ext
            })

    ticket = database.create_support_ticket(
        company_id=comp_id,
        company_name=comp_name,
        raised_by=raised_by,
        subject=subject,
        description=description,
        attachments=attachments,
        role=role
    )

    return jsonify({
        'success': True,
        'message': f'Ticket {ticket["ticket_id"]} submitted successfully!',
        'ticket': ticket
    })

@app.route('/api/support/ticket/<ticket_id>', methods=['GET'])
@login_required
def api_get_support_ticket(ticket_id):
    """Fetches details of a specific ticket."""
    comp_id = None if session.get('role') == 'super_admin' else (session.get('company_id') or 'ARGUS_MASTER')
    ticket = database.get_support_ticket_by_id(ticket_id, company_id=comp_id)
    if not ticket:
        return jsonify({'success': False, 'error': 'Ticket not found or unauthorized access.'}), 404
    return jsonify({'success': True, 'ticket': ticket})

@app.route('/api/support/ticket/<ticket_id>/reply', methods=['POST'])
@login_required
def api_reply_support_ticket(ticket_id):
    """Posts a reply message to the ticket conversation thread."""
    role = session.get('role')
    comp_id = None if role == 'super_admin' else (session.get('company_id') or 'ARGUS_MASTER')
    ticket = database.get_support_ticket_by_id(ticket_id, company_id=comp_id)
    if not ticket:
        return jsonify({'success': False, 'error': 'Ticket not found or unauthorized.'}), 404

    message = request.form.get('message', '').strip()
    if not message:
        return jsonify({'success': False, 'error': 'Message content cannot be empty.'}), 400

    # Sender resolution
    if role == 'super_admin':
        sender_name = session.get('user_name', 'Super Admin')
        sender_role = 'super_admin'
        badge = None
    else:
        sender_name = session.get('user_name') or f"Admin ({ticket.get('company_name', 'Company')[:10]})"
        sender_role = 'company_admin'
        badge = None

    # Handle any uploaded attachments in reply
    attachments = []
    ticket_comp_id = ticket.get('company_id', 'ARGUS_MASTER')
    comp_upload_dir = os.path.join(SUPPORT_UPLOADS_FOLDER, str(ticket_comp_id))
    os.makedirs(comp_upload_dir, exist_ok=True)

    uploaded_files = request.files.getlist('attachments')
    for f in uploaded_files:
        if f and f.filename and allowed_support_file(f.filename):
            clean_filename = werkzeug.utils.secure_filename(f.filename) or 'file'
            uniq_filename = f"{int(time.time())}_{uuid.uuid4().hex[:6]}_{clean_filename}"
            target_path = os.path.join(comp_upload_dir, uniq_filename)
            f.save(target_path)

            file_size = os.path.getsize(target_path) if os.path.exists(target_path) else 0
            ext = clean_filename.rsplit('.', 1)[1].lower() if '.' in clean_filename else ''

            attachments.append({
                'id': f"att_{uuid.uuid4().hex[:8]}",
                'filename': uniq_filename,
                'original_name': f.filename,
                'file_path': target_path,
                'file_size_bytes': file_size,
                'file_size_display': _format_file_size(file_size),
                'extension': ext
            })

    # Status update if super admin reply
    new_status = None
    if role == 'super_admin' and ticket.get('status') == 'Open':
        new_status = 'In Progress'

    res = database.add_support_ticket_message(
        ticket_id=ticket['ticket_id'],
        sender_name=sender_name,
        sender_role=sender_role,
        message_text=message,
        attachments=attachments,
        new_status=new_status,
        badge=badge,
        company_id=comp_id
    )

    if not res.get('success'):
        return jsonify(res), 400

    return jsonify({
        'success': True,
        'message': 'Reply sent successfully.',
        'ticket': res['ticket']
    })

@app.route('/api/support/ticket/<ticket_id>/resolve', methods=['POST'])
@login_required
def api_resolve_support_ticket(ticket_id):
    """Marks ticket as resolved."""
    role = session.get('role')
    comp_id = None if role == 'super_admin' else (session.get('company_id') or 'ARGUS_MASTER')
    ticket = database.get_support_ticket_by_id(ticket_id, company_id=comp_id)
    if not ticket:
        return jsonify({'success': False, 'error': 'Ticket not found or unauthorized.'}), 404

    remark = request.json.get('remark') if request.is_json else request.form.get('remark')
    user_name = session.get('user_name', 'Admin')

    res = database.update_support_ticket_status(
        ticket_id=ticket['ticket_id'],
        new_status='Resolved',
        updated_by_name=user_name,
        updated_by_role=role,
        remark=remark or "Ticket marked as resolved.",
        company_id=comp_id
    )
    return jsonify(res)

@app.route('/api/support/ticket/<ticket_id>/reopen', methods=['POST'])
@login_required
def api_reopen_support_ticket(ticket_id):
    """Reopens a resolved or closed ticket."""
    role = session.get('role')
    comp_id = None if role == 'super_admin' else (session.get('company_id') or 'ARGUS_MASTER')
    ticket = database.get_support_ticket_by_id(ticket_id, company_id=comp_id)
    if not ticket:
        return jsonify({'success': False, 'error': 'Ticket not found or unauthorized.'}), 404

    remark = request.json.get('remark') if request.is_json else request.form.get('remark')
    user_name = session.get('user_name', 'Admin')

    res = database.update_support_ticket_status(
        ticket_id=ticket['ticket_id'],
        new_status='Open',
        updated_by_name=user_name,
        updated_by_role=role,
        remark=remark or "Ticket reopened for further assistance.",
        company_id=comp_id
    )
    return jsonify(res)

@app.route('/api/support/demo-videos', methods=['GET'])
@login_required
def api_get_demo_videos():
    """Returns the list of 8 demo video chapters with learning objectives."""
    videos = database.get_demo_videos()
    return jsonify({'success': True, 'videos': videos})

@app.route('/api/support/badge-count', methods=['GET'])
@login_required
def api_get_support_badge_count():
    """Returns live open ticket count for sidebar badges."""
    role = session.get('role')
    comp_id = None if role == 'super_admin' else (session.get('company_id') or 'ARGUS_MASTER')
    count = database.get_support_badge_count(comp_id)
    return jsonify({'success': True, 'count': count})

# Admin Support Ticket APIs
@app.route('/api/admin/support/tickets', methods=['GET'])
@super_admin_required
def api_admin_support_tickets():
    """Returns tickets across all companies for super admin dashboard."""
    comp_filter = request.args.get('company_id', '')
    status_filter = request.args.get('status', 'all')
    search = request.args.get('search', '')
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))

    result = database.get_all_support_tickets(
        company_filter=comp_filter,
        status_filter=status_filter,
        search=search,
        page=page,
        limit=limit
    )
    return jsonify({
        'success': True,
        'tickets': result['tickets'],
        'counts': result['counts'],
        'companies': result['companies'],
        'total': result['total'],
        'page': result['page'],
        'limit': result['limit']
    })

@app.route('/api/admin/support/ticket/<ticket_id>/status', methods=['POST'])
@super_admin_required
def api_admin_update_ticket_status(ticket_id):
    """Updates status of any ticket with admin remark."""
    data = request.get_json(silent=True) or request.form
    new_status = (data.get('status') or '').strip()
    remark = (data.get('remark') or '').strip()

    if not new_status:
        return jsonify({'success': False, 'error': 'Status is required.'}), 400

    admin_name = session.get('user_name', 'Super Admin')
    res = database.update_support_ticket_status(
        ticket_id=ticket_id,
        new_status=new_status,
        updated_by_name=admin_name,
        updated_by_role='super_admin',
        remark=remark or f"Status updated to {new_status} by Super Admin."
    )
    return jsonify(res)

@app.route('/api/admin/support/export', methods=['GET'])
@super_admin_required
def api_admin_support_export():
    """Exports tickets in CSV or Excel format."""
    export_format = request.args.get('format', 'csv').lower()
    comp_filter = request.args.get('company_id', '')
    status_filter = request.args.get('status', 'all')
    search = request.args.get('search', '')

    result = database.get_all_support_tickets(
        company_filter=comp_filter,
        status_filter=status_filter,
        search=search,
        page=1,
        limit=5000
    )
    tickets = result['tickets']

    if export_format not in ['csv', 'excel']:
        return jsonify({'success': False, 'error': 'Unsupported export format'}), 400

    headers = ['Ticket ID', 'Date & Time', 'Company Name', 'Raised By', 'Subject', 'Description', 'Status', 'Attachments Count', 'Created At']
    rows = []
    for t in tickets:
        rows.append([
            t.get('ticket_id', ''),
            t.get('date_time_str', ''),
            t.get('company_name', ''),
            t.get('raised_by', ''),
            t.get('subject', ''),
            t.get('description', ''),
            t.get('status', ''),
            len(t.get('attachments', [])),
            t.get('created_at', '')
        ])

    filename = f"Support_Tickets_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    return export_table_data(headers, rows, filename, export_format=export_format, sheet_name="Support Tickets")

# Attachment Download and Download-All (ZIP)
@app.route('/api/support/attachment/<ticket_id>/<filename>', methods=['GET'])
@login_required
def api_get_support_attachment(ticket_id, filename):
    """Serves individual ticket attachment."""
    role = session.get('role')
    comp_id = None if role == 'super_admin' else (session.get('company_id') or 'ARGUS_MASTER')
    ticket = database.get_support_ticket_by_id(ticket_id, company_id=comp_id)
    if not ticket:
        return jsonify({'error': 'Ticket not found or unauthorized.'}), 404

    target_comp_id = ticket.get('company_id', 'ARGUS_MASTER')
    folder = os.path.join(SUPPORT_UPLOADS_FOLDER, str(target_comp_id))
    file_path = os.path.join(folder, werkzeug.utils.secure_filename(filename))

    if not os.path.exists(file_path):
        return jsonify({'error': 'Attachment file not found.'}), 404

    as_attachment = request.args.get('download', '0') == '1'
    return send_file(file_path, as_attachment=as_attachment)

@app.route('/api/support/attachment/<ticket_id>/download-all', methods=['GET'])
@login_required
def api_download_all_attachments(ticket_id):
    """Zips all attachments belonging to ticket and downloads them in one go."""
    role = session.get('role')
    comp_id = None if role == 'super_admin' else (session.get('company_id') or 'ARGUS_MASTER')
    ticket = database.get_support_ticket_by_id(ticket_id, company_id=comp_id)
    if not ticket:
        return jsonify({'error': 'Ticket not found or unauthorized.'}), 404

    attachments = ticket.get('attachments', [])
    if not attachments:
        return jsonify({'error': 'No attachments found for this ticket.'}), 400

    target_comp_id = ticket.get('company_id', 'ARGUS_MASTER')
    folder = os.path.join(SUPPORT_UPLOADS_FOLDER, str(target_comp_id))

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for att in attachments:
            fname = att.get('filename')
            orig_name = att.get('original_name', fname)
            fpath = os.path.join(folder, fname) if fname else None
            if fpath and os.path.exists(fpath):
                zip_file.write(fpath, arcname=orig_name)

    zip_buffer.seek(0)
    zip_filename = f"Ticket_{ticket.get('ticket_id', ticket_id)}_Attachments.zip"
    return send_file(
        zip_buffer,
        mimetype='application/zip',
        as_attachment=True,
        download_name=zip_filename
    )


# ==============================================================================
# FIELD WORKER MODULE ROUTES (/field-worker/login, /field-worker & /api/field-worker/*)
# ==============================================================================

def _is_field_authenticated():
    """Checks if the current session is authenticated as Admin, Company Admin, or Field Worker."""
    return bool(session.get('admin_logged_in') or session.get('employee_logged_in'))


def _resolve_field_company_id(req_company_id=None):
    """
    Resolves the active company_id with strict multi-tenant isolation:
    - Company Admins & Employees (Field Workers) are strictly locked to their session['company_id'].
    - Only Super Admin can switch across companies via req_company_id.
    """
    role = session.get('role')
    sess_cid = session.get('company_id')
    if role in ['company_admin', 'admin', 'employee']:
        return str(sess_cid or 'ARGUS_MASTER')
    if role == 'super_admin' and req_company_id and str(req_company_id).strip() not in ['', 'ALL', 'None', 'null']:
        return str(req_company_id).strip()
    if sess_cid and sess_cid != 'ALL':
        return str(sess_cid)
    return 'ARGUS_MASTER'


@app.route('/field-worker/login', methods=['GET', 'POST'])
def field_worker_login():
    """
    Dedicated Multi-Tenant Login Gate for the Field Worker Module.
    Automatically detects which company a Field Worker or Company Admin belongs to
    and locks their session to prevent unauthorized cross-company access.
    """
    if request.method == 'GET' and _is_field_authenticated():
        return redirect(url_for('field_worker_module_page'))

    if request.method == 'POST':
        login_role = (request.form.get('login_role') or 'field_worker').strip().lower()
        email = (request.form.get('email') or '').strip()
        password = (request.form.get('password') or '').strip()

        if not email or not password:
            return render_template(
                'field_worker_login.html',
                error='Please enter your registered email and password.',
                selected_role=login_role
            )

        if login_role == 'field_worker':
            result = database.validate_employee_login(email, password=password)
            if isinstance(result, dict) and result.get('success'):
                employee = result['employee']
                is_fw = bool(
                    employee.get('is_field_worker') is True
                    or str(employee.get('employee_type') or '').lower() in ['field_worker', 'field', 'fw']
                )
                if not is_fw:
                    return render_template(
                        'field_worker_login.html',
                        error='Access Denied: This account is registered as a Standard Employee. Please login using the Standard Employee Portal.',
                        selected_role=login_role
                    )
                session.clear()
                session.permanent = True
                session['employee_logged_in'] = True
                session['role'] = 'employee'
                session['employee_id'] = str(employee.get('id', employee.get('_id', '')))
                session['employee_name'] = employee.get('employee_name', '')
                session['employee_email'] = employee.get('email_id', '')
                comp_id = employee.get('company_id', 'ARGUS_MASTER')
                session['company_id'] = comp_id
                comp_rec = database.get_company_by_id(comp_id) if comp_id != 'ARGUS_MASTER' else None
                session['company_name'] = comp_rec.get('company_name', 'ARGUS TECHNOLOGIES') if comp_rec else 'ARGUS TECHNOLOGIES'
                resp = redirect(url_for('field_worker_module_page'))
                return finalize_login_session(resp, 'employee', employee)

            err_code = result.get('error') if isinstance(result, dict) else None
            if err_code == 'COMPANY_DEACTIVATED':
                msg = result.get('message') or 'Access Denied: Your company account is deactivated.'
            elif err_code == 'EMPLOYEE_INACTIVE':
                msg = result.get('message') or 'Access Denied: Your employee profile is currently Inactive.'
            elif err_code == 'PASSWORD_NOT_SET':
                msg = "Password not set yet. Please click 'Continue with Google' below for first-time verification."
            elif err_code == 'NOT_REGISTERED':
                msg = 'Access Denied: This email is not registered as a Field Worker in any company.'
            else:
                msg = 'Invalid password. Please verify your credentials or continue with Google.'
            return render_template('field_worker_login.html', error=msg, selected_role='field_worker')

        else:
            # Company Admin (or Super Admin fallback) Sign-In
            comp_res = database.validate_company_login(email, password=password)
            if isinstance(comp_res, dict) and comp_res.get('success'):
                company = comp_res['company']
                session.clear()
                session.permanent = True
                session['admin_logged_in'] = True
                session['admin_username'] = company.get('company_name', 'Company Admin')
                session['role'] = 'company_admin'
                session['company_id'] = str(company.get('id') or company.get('_id', ''))
                session['company_name'] = company.get('company_name', '')
                session['company_email'] = company.get('email', '')
                resp = redirect(url_for('field_worker_module_page'))
                return finalize_login_session(resp, 'company_admin', company)

            # Fallback: Check if System Super Admin
            adm_res = database.validate_admin_login(email, password=password)
            if isinstance(adm_res, dict) and adm_res.get('success'):
                admin_user = adm_res['admin']
                session.clear()
                session.permanent = True
                session['admin_logged_in'] = True
                session['admin_username'] = admin_user.get('username', 'Admin')
                session['admin_email'] = admin_user.get('email', 'technologiesargus@gmail.com')
                session['role'] = 'super_admin'
                session['company_id'] = 'ARGUS_MASTER'
                session['company_name'] = 'ARGUS TECHNOLOGIES'
                resp = redirect(url_for('field_worker_module_page'))
                return finalize_login_session(resp, 'super_admin', admin_user)

            err_code = comp_res.get('error') if isinstance(comp_res, dict) else None
            if err_code == 'COMPANY_DEACTIVATED':
                msg = comp_res.get('message') or 'Access Denied: This company account is deactivated.'
            elif err_code == 'PASSWORD_NOT_SET':
                msg = "Password not set yet. Please click 'Continue with Google' below."
            else:
                msg = 'Access Denied: Invalid company admin email or password.'
            return render_template('field_worker_login.html', error=msg, selected_role='company_admin')

    return render_template('field_worker_login.html', selected_role='field_worker')


@app.route('/field-worker')
def field_worker_module_page():
    """Renders the dedicated Field Worker Force Management Module UI with Multi-Tenant Isolation."""
    if not _is_field_authenticated():
        return redirect(url_for('field_worker_login'))

    now_ist = database.get_ist_now()
    today_str = request.args.get('date') or now_ist.strftime('%Y-%m-%d')
    month_year = request.args.get('month') or now_ist.strftime('%Y-%m')

    user_role = session.get('role', 'company_admin')
    logged_in_emp_name = session.get('employee_name', '') if user_role == 'employee' else ''

    comp_res = database.get_all_companies(limit=1000) if user_role == 'super_admin' else None
    companies = comp_res.get('data', []) if isinstance(comp_res, dict) else (comp_res or [])
    active_cid = _resolve_field_company_id(request.args.get('company_id'))

    # Check subscription entitlement for Field Worker module
    if user_role != 'super_admin':
        comp_modules = database.get_company_modules(active_cid)
        if not comp_modules.get('field_worker', False):
            flash('The Field Worker Module is not authorized for your company subscription. Please contact your administrator.', 'warning')
            return redirect(url_for('employee_portal') if user_role == 'employee' else url_for('dashboard'))

    active_comp = database.get_company_by_id(active_cid) or {
        'id': 'ARGUS_MASTER',
        'company_name': 'ARGUS TECHNOLOGIES',
        'latitude': 11.02980,
        'longitude': 76.97400,
        'address': 'SF NO. 515, Bharathiyar Road, Ganapathy, Coimbatore'
    }

    emp_scope = logged_in_emp_name if user_role == 'employee' else None

    # Only load field workers for the active company; exclude standard employees
    if user_role in ['super_admin', 'company_admin']:
        field_workers = database.get_field_workers(company_id=active_cid, include_all_employees=False)
    else:
        field_workers = database.get_field_workers(company_id=active_cid, include_all_employees=False)
        if emp_scope:
            field_workers = [w for w in field_workers if w.get('employee_name') == emp_scope]
    all_employees = field_workers

    branches = database.get_company_branches(active_cid)
    assignments = database.get_field_assignments(company_id=active_cid, employee_name=emp_scope, date_str=today_str)
    tracking = database.get_field_live_tracking(company_id=active_cid, date_str=today_str, employee_name=emp_scope)
    visits = database.get_field_customer_visits(company_id=active_cid, employee_name=emp_scope)
    # Defer Tab 6 payroll calculation on initial page load: user lands on Tab 1
    salary_reports = []

    # Field Worker Dashboard & Leave stats
    dashboard_stats = database.get_field_worker_dashboard_stats(company_id=active_cid, date_str=today_str)
    fw_pending_leaves = database.get_pending_field_worker_leave_requests_count(active_cid)
    fw_emp_id = session.get('employee_id') if user_role == 'employee' else None

    # Fast initial load: leave requests and balances are loaded asynchronously when Tab 6 is activated
    fw_leave_stats = {'pending_count': fw_pending_leaves, 'approved_today': 0, 'rejected_count': 0, 'low_balance_count': 0}
    fw_leave_requests = {'requests': [], 'status_counts': {'pending': fw_pending_leaves}}
    fw_leave_balances = []
    if user_role == 'employee' and fw_emp_id:
        fw_leave_balances = [database.get_employee_leave_balance(fw_emp_id, company_id=active_cid)]

    allowed_tabs = None
    if user_role == 'employee':
        emp_match = next((w for w in field_workers if w.get('employee_name') == logged_in_emp_name), None)
        if not emp_match and logged_in_emp_name:
            emp_match = database.get_employee_by_name(logged_in_emp_name, company_id=active_cid)
        if emp_match:
            allowed_tabs = emp_match.get('allowed_tabs')

    return render_template(
        'field_worker_module.html',
        companies=companies,
        active_company=active_comp,
        active_company_id=active_cid,
        user_role=user_role,
        logged_in_employee_name=logged_in_emp_name,
        field_workers=field_workers,
        all_employees=all_employees,
        branches=branches,
        assignments=assignments,
        tracking=tracking,
        visits=visits,
        salary_reports=salary_reports,
        dashboard_stats=dashboard_stats,
        fw_leave_stats=fw_leave_stats,
        fw_leave_requests=fw_leave_requests.get('requests', []),
        fw_leave_counts=fw_leave_requests.get('status_counts', {}),
        fw_leave_balances=fw_leave_balances,
        fw_pending_leaves=fw_pending_leaves,
        today_str=today_str,
        month_year=month_year,
        allowed_tabs=allowed_tabs,
        company_support_enabled=bool(active_comp.get('support_enabled', True))
    )


@app.route('/api/field-worker/overview', methods=['GET'])
def api_field_worker_overview():
    """Returns live JSON state for the Field Worker Module dashboard (isolated by tenant & role)."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    now_ist = database.get_ist_now()
    cid = _resolve_field_company_id(request.args.get('company_id'))
    date_str = request.args.get('date') or now_ist.strftime('%Y-%m-%d')
    month_year = request.args.get('month') or now_ist.strftime('%Y-%m')
    user_role = session.get('role')
    emp_filter = session.get('employee_name') if user_role == 'employee' else (request.args.get('employee_name') or None)
    include_salary = request.args.get('include_salary') in ['1', 'true', 'True'] or request.args.get('tab') == 'paneFieldPayroll'

    comp = database.get_company_by_id(cid) or {'id': cid, 'company_name': 'ARGUS TECHNOLOGIES', 'latitude': 11.0298, 'longitude': 76.9740}
    
    if user_role in ['super_admin', 'company_admin']:
        workers = database.get_field_workers(company_id=cid, include_all_employees=False)
    else:
        workers = database.get_field_workers(company_id=cid, include_all_employees=False)
        if emp_filter:
            workers = [w for w in workers if w.get('employee_name') == emp_filter]
    all_emps = workers

    branches = database.get_company_branches(cid)
    assignments = database.get_field_assignments(company_id=cid, employee_name=emp_filter, date_str=date_str)
    tracking = database.get_field_live_tracking(company_id=cid, date_str=date_str, employee_name=emp_filter)
    visits = database.get_field_customer_visits(company_id=cid, employee_name=emp_filter)
    salary_reports = database.get_field_worker_salary_report(company_id=cid, month_year=month_year, employee_name=emp_filter) if include_salary else []

    dashboard_stats = database.get_field_worker_dashboard_stats(company_id=cid, date_str=date_str)
    fw_emp_id = session.get('employee_id') if user_role == 'employee' else None
    include_leaves = request.args.get('include_leaves') in ['1', 'true', 'True'] or request.args.get('tab') == 'paneFieldLeaves'
    if include_leaves:
        fw_leave_stats = database.get_admin_leave_stats(company_id=cid, for_field_workers=True)
        fw_leave_res = database.get_leave_requests(company_id=cid, employee_id=fw_emp_id, for_field_workers=True, limit=50)
    else:
        fw_leave_stats = None
        fw_leave_res = {'requests': [], 'status_counts': {}}

    return jsonify({
        'success': True,
        'company': comp,
        'company_id': cid,
        'date': date_str,
        'month_year': month_year,
        'field_workers': workers,
        'all_employees': all_emps,
        'branches': branches,
        'assignments': assignments,
        'tracking': tracking,
        'dashboard_stats': dashboard_stats,
        'fw_leave_stats': fw_leave_stats if include_leaves else None,
        'fw_leave_requests': fw_leave_res.get('requests', []) if include_leaves else None,
        'fw_leave_counts': fw_leave_res.get('status_counts', {}) if include_leaves else None,
        'ai_incidents': tracking.get('ai_incidents', []) if isinstance(tracking, dict) else [],
        'active_ai_incidents': tracking.get('active_ai_incidents', []) if isinstance(tracking, dict) else [],
        'ai_summary': tracking.get('ai_summary', {}) if isinstance(tracking, dict) else {},
        'visits': visits,
        'salary_reports': salary_reports
    })


@app.route('/api/field-worker/register', methods=['POST'])
def api_field_worker_register():
    """Registers a new Field Worker or updates an existing employee with Field Worker configuration (Admin only)."""
    if not session.get('admin_logged_in'):
        return jsonify({'success': False, 'error': 'Only Company Admin or System Admin can register field workers.'}), 403
    try:
        data = request.get_json(silent=True) or request.form.to_dict()
        cid = _resolve_field_company_id(data.get('company_id'))
        data['company_id'] = cid
        data['is_field_worker'] = True
        data['employee_type'] = 'field_worker'

        # Resolve assigned branch location
        bid = str(data.get('branch_id') or '').strip()
        if bid and bid not in ['HOME', 'ALL', 'none', 'None']:
            data['branch_id'] = bid
            b_doc = database.get_branch_by_id(bid, cid)
            data['branch_name'] = b_doc.get('branch_name') if b_doc else 'Home Branch (Head Office)'
        else:
            data['branch_id'] = None
            data['branch_name'] = 'Home Branch (Head Office)'
        data['branch'] = data['branch_name']
        data['territory_zone'] = data['branch_name']
        if 'allowed_tabs' in data and isinstance(data['allowed_tabs'], list):
            data['allowed_tabs'] = [str(t).strip() for t in data['allowed_tabs'] if str(t).strip()]

        # Handle optional face photo for biometric enrollment
        photo_b64 = data.get('photo_data') or ''
        if not photo_b64 and 'photo' in request.files:
            try:
                p_file = request.files['photo']
                if p_file and p_file.filename:
                    p_bytes = p_file.read()
                    if p_bytes:
                        photo_b64 = 'data:image/jpeg;base64,' + base64.b64encode(p_bytes).decode('utf-8')
            except Exception as pf_e:
                print(f"Warning: reading uploaded photo file: {pf_e}")

        if photo_b64 and str(photo_b64).startswith('data:image'):
            try:
                raw_bytes = base64.b64decode(photo_b64.split(',', 1)[1])
                unique_filename, photo_data_url, optimized_bytes = process_uploaded_photo(raw_bytes, 'fw_capture.jpg')
                data['photo'] = unique_filename
                data['photo_filename'] = unique_filename
                data['photo_data'] = photo_data_url
                
                fe = get_face_engine()
                if fe and hasattr(fe, 'extract_face_embedding_from_image'):
                    emb = fe.extract_face_embedding_from_image(optimized_bytes)
                    if emb:
                        data['face_embedding'] = json.dumps(emb)
                        data['embedding_model'] = 'sface_v1'
            except Exception as img_err:
                print(f"Field worker photo processing error: {img_err}")

        # Resolve worker-specific petrol allowance (Unit distance fixed to 1.0 KM)
        unit_km = 1.0
        data['petrol_unit_km'] = 1.0

        p_rate = data.get('petrol_rate_inr')
        if p_rate not in [None, '']:
            p_rate_val = database.safe_float(p_rate, 0.0)
            if p_rate_val > 0:
                data['petrol_rate_inr'] = p_rate_val
                data['has_custom_petrol_rate'] = True
                data['km_allowance_rate'] = round(p_rate_val, 2)
            else:
                data['petrol_rate_inr'] = None
                data['has_custom_petrol_rate'] = False
                data['km_allowance_rate'] = None
        else:
            data['petrol_rate_inr'] = None
            data['has_custom_petrol_rate'] = False
            data['km_allowance_rate'] = None

        existing_emp_id = str(data.get('existing_employee_id') or '').strip()
        if existing_emp_id:
            database.update_employee(existing_emp_id, data, company_id=cid)
            database.invalidate_dashboard_cache(cid)
            return jsonify({
                'success': True,
                'employee_id': existing_emp_id,
                'message': f"Field Worker configuration updated for {data.get('employee_name', 'Employee')}."
            })

        if not data.get('department'):
            data['department'] = 'Field Sales & Operations'
        if not data.get('designation'):
            data['designation'] = 'Field Executive'

        new_id = database.create_employee(data, company_id=cid)
        database.invalidate_dashboard_cache(cid)

        return jsonify({
            'success': True,
            'employee_id': new_id,
            'message': f"Field Worker '{data.get('employee_name')}' registered successfully!"
        })
    except ValueError as ve:
        return jsonify({'success': False, 'error': str(ve)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/field-worker/workers/<worker_id>', methods=['DELETE'])
def api_field_worker_delete(worker_id):
    """Deletes a field worker record and clears associated caches (Admin only)."""
    if not session.get('admin_logged_in'):
        return jsonify({'success': False, 'error': 'Unauthorized: Only Company Admin or System Admin can delete field workers.'}), 403
    try:
        cid = _resolve_field_company_id(request.args.get('company_id'))
        worker = database.get_employee_by_id(worker_id, company_id=cid)
        if not worker:
            return jsonify({'success': False, 'error': 'Field worker not found.'}), 404

        worker_name = worker.get('employee_name', 'Field Worker')
        database.delete_employee(worker_id, company_id=cid)
        database.invalidate_dashboard_cache(cid)
        return jsonify({
            'success': True,
            'message': f"Field worker '{worker_name}' deleted successfully."
        })
    except Exception as e:
        print("Error in api_field_worker_delete:", traceback.format_exc())
        return jsonify({'success': False, 'error': f"Failed to delete worker: {str(e)}"}), 500


@app.route('/api/field-worker/assignments', methods=['GET', 'POST'])
def api_field_worker_assignments():
    """Lists or creates target customer location assignments for field workers."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    if request.method == 'GET':
        cid = _resolve_field_company_id(request.args.get('company_id'))
        emp = session.get('employee_name') if session.get('role') == 'employee' else (request.args.get('employee_name') or None)
        dt = request.args.get('date') or None
        st = request.args.get('status') or None
        items = database.get_field_assignments(company_id=cid, employee_name=emp, date_str=dt, status=st)
        return jsonify({'success': True, 'assignments': items})

    try:
        data = request.get_json(silent=True) or request.form.to_dict()
        cid = _resolve_field_company_id(data.get('company_id'))
        if session.get('role') == 'employee' and session.get('employee_name'):
            data['employee_name'] = session.get('employee_name')
        created = database.create_field_assignment(data, company_id=cid)
        return jsonify({'success': True, 'assignment': created, 'message': 'Visit destination scheduled successfully!'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/field-worker/assignments/<assignment_id>', methods=['PUT', 'DELETE'])
def api_field_worker_assignment_detail(assignment_id):
    """Updates or deletes a specific field location assignment."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    if request.method == 'DELETE':
        if not session.get('admin_logged_in') and session.get('role') != 'employee':
            return jsonify({'success': False, 'error': 'Unauthorized'}), 403
        ok = database.delete_field_assignment(assignment_id, company_id=cid)
        return jsonify({'success': ok, 'message': 'Assignment deleted.' if ok else 'Assignment not found.'})

    data = request.get_json(silent=True) or {}
    ok = database.update_field_assignment(assignment_id, data, company_id=cid)
    return jsonify({'success': ok, 'message': 'Assignment updated.' if ok else 'Assignment not found.'})


@app.route('/api/field-worker/search-destination', methods=['GET'])
def api_search_field_destination():
    """Searches for company, client, and place names with coordinates for destination pinning."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    query = request.args.get('query', '').strip()
    cid = _resolve_field_company_id(request.args.get('company_id'))
    results = database.search_field_destinations(query, company_id=cid)
    return jsonify({'success': True, 'results': results})


@app.route('/api/field-worker/location-ping', methods=['POST'])
def api_field_worker_location_ping():
    """Records a live GPS ping from a field worker and detects Company Site vs Personal Detour."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    try:
        data = request.get_json(silent=True) or request.form.to_dict()
        cid = _resolve_field_company_id(data.get('company_id'))
        if session.get('role') == 'employee' and session.get('employee_name'):
            data['employee_name'] = session.get('employee_name')
        ping = database.record_field_location_ping(data, company_id=cid)
        return jsonify({
            'success': True,
            'ping': ping,
            'ai_incident': ping.get('ai_incident'),
            'message': f"GPS Logged: {ping.get('status_label')} ({ping.get('distance_from_target_m', 0)}m from target)"
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/field-worker/ai-incidents', methods=['GET'])
def api_field_worker_ai_incidents():
    """Returns AI-detected detour and anomaly incidents for the company."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    date_str = request.args.get('date') or database.get_ist_now().strftime('%Y-%m-%d')
    emp_name = request.args.get('employee_name')
    status = request.args.get('status')

    incidents = database.get_field_ai_incidents(company_id=cid, date_str=date_str, employee_name=emp_name, status=status)
    summary = database.get_field_ai_incidents_summary(company_id=cid, date_str=date_str)

    return jsonify({
        'success': True,
        'date': date_str,
        'incidents': incidents,
        'summary': summary
    })


@app.route('/api/field-worker/ai-incidents/<incident_id>/resolve', methods=['POST'])
def api_field_worker_ai_incident_resolve(incident_id):
    """Marks an AI detour incident as Acknowledged or Resolved with admin notes."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    data = request.get_json(silent=True) or request.form.to_dict() or {}
    cid = _resolve_field_company_id(data.get('company_id'))
    status = str(data.get('status') or 'Resolved').strip()
    notes = str(data.get('admin_notes') or '').strip()
    admin_user = session.get('admin_user') or session.get('employee_name') or 'Company Admin'

    ok = database.resolve_field_ai_incident(
        incident_id=incident_id,
        status=status,
        admin_notes=notes,
        resolved_by=admin_user,
        company_id=cid
    )
    if not ok:
        return jsonify({'success': False, 'error': 'Incident not found or could not be updated'}), 404

    return jsonify({
        'success': True,
        'message': f'Incident marked as {status}',
        'incident_id': incident_id
    })


@app.route('/api/field-worker/visits', methods=['GET', 'POST'])
def api_field_worker_visits():
    """Lists or logs Customer Visits (Lead CRM) with automatic geo-verification."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    if request.method == 'GET':
        cid = _resolve_field_company_id(request.args.get('company_id'))
        emp = session.get('employee_name') if session.get('role') == 'employee' else (request.args.get('employee_name') or None)
        start_d = request.args.get('start_date') or None
        end_d = request.args.get('end_date') or None
        status = request.args.get('lead_status') or None
        items = database.get_field_customer_visits(company_id=cid, employee_name=emp, start_date=start_d, end_date=end_d, lead_status=status)
        return jsonify({'success': True, 'visits': items})

    try:
        # Support both multipart form data (with audio_file) and JSON data
        data = request.form.to_dict() if request.form else (request.get_json(silent=True) or {})
        cid = _resolve_field_company_id(data.get('company_id'))
        if session.get('role') == 'employee' and session.get('employee_name'):
            data['employee_name'] = session.get('employee_name')

        audio_file = request.files.get('audio_file') or request.files.get('audio') or request.files.get('recording')
        temp_audio_path = None
        uploaded_gridfs_id = None
        if audio_file and getattr(audio_file, 'filename', None):
            clean_fname = werkzeug.utils.secure_filename(audio_file.filename) or 'recording.m4a'
            ext = os.path.splitext(clean_fname)[1].lower() or '.m4a'
            upload_dir = os.path.join(app.root_path, 'static', 'uploads', 'visit_recordings', cid or 'ARGUS_MASTER')
            os.makedirs(upload_dir, exist_ok=True)
            saved_name = f"visit_audio_{int(time.time() * 1000)}{ext}"
            saved_path = os.path.join(upload_dir, saved_name)

            # Read bytes for persistent GridFS cloud storage across all deploys
            audio_bytes = audio_file.read()
            mime_type = audio_file.content_type or ('audio/webm' if ext == '.webm' else 'audio/mp4')
            try:
                import gridfs
                fs = gridfs.GridFS(database.get_db())
                grid_file = fs.put(audio_bytes, filename=saved_name, content_type=mime_type)
                uploaded_gridfs_id = str(grid_file)
                data['audio_gridfs_id'] = uploaded_gridfs_id
            except Exception as gfs_err:
                print("Error saving audio to GridFS:", gfs_err)

            # Also save to local disk for immediate processing
            try:
                with open(saved_path, 'wb') as f_out:
                    f_out.write(audio_bytes)
            except Exception:
                pass

            rel_path = f"static/uploads/visit_recordings/{cid or 'ARGUS_MASTER'}/{saved_name}"
            data['audio_file_path'] = rel_path
            data['has_audio_recording'] = True
            data['audio_mime_type'] = mime_type
            temp_audio_path = saved_path

        visit = database.create_field_customer_visit(data, company_id=cid)
        visit_id = visit.get('id')

        # If audio was uploaded, process with Tamil Gemini STT and attach analysis to the visit document
        if temp_audio_path and visit_id:
            try:
                dur = float(data.get('duration_seconds') or data.get('audio_duration_seconds') or 0.0)
                mime = data.get('audio_mime_type') or 'audio/mp4'
                audio_url = f"/api/field-worker/visits/{visit_id}/audio"
                updated_visit = database.attach_audio_and_process_visit(
                    visit_id=visit_id,
                    audio_file_path=temp_audio_path,
                    audio_url=audio_url,
                    duration_seconds=dur,
                    mime_type=mime,
                    company_id=cid,
                    gridfs_id=uploaded_gridfs_id
                )
                if updated_visit:
                    visit = updated_visit
            except Exception as stt_err:
                print("Error in background Tamil audio processing:", stt_err)

        return jsonify({
            'success': True,
            'visit': visit,
            'message': f"Customer visit to '{visit.get('customer_name')}' recorded ({visit.get('verification_badge')})!"
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 400


@app.route('/api/field-worker/visits/<visit_id>', methods=['PUT', 'DELETE'])
def api_field_worker_visit_detail(visit_id):
    """Updates or deletes a specific field customer visit / lead."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    if request.method == 'DELETE':
        if not session.get('admin_logged_in'):
            return jsonify({'success': False, 'error': 'Unauthorized'}), 403
        ok = database.delete_field_customer_visit(visit_id, company_id=cid)
        return jsonify({'success': ok, 'message': 'Lead deleted.' if ok else 'Lead not found.'})

    data = request.get_json(silent=True) or request.form.to_dict()
    ok = database.update_field_customer_visit(visit_id, data, company_id=cid)
    return jsonify({'success': ok, 'message': 'Lead updated successfully!' if ok else 'Lead not found.'})


@app.route('/api/field-worker/visits/<visit_id>/audio', methods=['GET'])
def api_field_worker_visit_audio(visit_id):
    """Streams the recorded visit audio file from GridFS cloud database or local filesystem with HTTP Range support."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    visit = database.get_field_customer_visit_by_id(visit_id, company_id=cid)
    if not visit and (session.get('role') == 'super_admin' or cid in ['ALL', 'ARGUS_MASTER', None]):
        visit = database.get_field_customer_visit_by_id(visit_id, company_id=None)

    if not visit or not visit.get('has_audio_recording'):
        return jsonify({'success': False, 'error': 'Audio recording not found for this visit.'}), 404

    mime = visit.get('audio_mime_type') or 'audio/webm'

    # 1. Stream from GridFS (persistent cloud database storage across all deploys)
    grid_id = visit.get('audio_gridfs_id')
    if grid_id:
        try:
            import gridfs, io
            from bson import ObjectId
            fs = gridfs.GridFS(database.get_db())
            if fs.exists(ObjectId(grid_id)):
                grid_out = fs.get(ObjectId(grid_id))
                return send_file(io.BytesIO(grid_out.read()), mimetype=grid_out.content_type or mime, conditional=True)
        except Exception as e:
            print("GridFS Audio Stream Error:", e)

    # 2. Local or dynamic filesystem path
    audio_path = str(visit.get('audio_file_path') or '').strip()
    if audio_path:
        norm_path = audio_path.replace('\\', '/')
        if 'static/uploads/visit_recordings/' in norm_path:
            rel_sub = norm_path.split('static/uploads/visit_recordings/', 1)[1]
            local_cand = os.path.join(app.root_path, 'static', 'uploads', 'visit_recordings', rel_sub)
            if os.path.exists(local_cand):
                audio_path = local_cand

        if not os.path.isabs(audio_path):
            audio_path = os.path.join(app.root_path, audio_path)

        if os.path.exists(audio_path):
            return send_file(audio_path, mimetype=mime, conditional=True)

    return jsonify({
        'success': False,
        'error': 'Audio recording expired during server redeployment (recorded prior to cloud persistence). Please record a new visit.'
    }), 404


@app.route('/api/field-worker/visits/<visit_id>/transcript', methods=['GET'])
def api_field_worker_visit_transcript(visit_id):
    """Returns the Tamil verbatim transcript, English translation, and AI insights for a visit."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    visit = database.get_field_customer_visit_by_id(visit_id, company_id=cid)
    if not visit and (session.get('role') == 'super_admin' or cid in ['ALL', 'ARGUS_MASTER', None]):
        visit = database.get_field_customer_visit_by_id(visit_id, company_id=None)

    if not visit:
        return jsonify({'success': False, 'error': 'Customer visit not found.'}), 404

    # Auto-reprocess on the fly if audio recording exists but transcript is missing or contains an error string
    t_en = str(visit.get('conversation_transcript_english') or '').strip()
    is_err_placeholder = (
        not t_en or
        any(k in t_en.lower() for k in ['error', 'no module', 'fail', 'fallback', 'required'])
    )

    if visit.get('has_audio_recording') and is_err_placeholder:
        audio_path = str(visit.get('audio_file_path') or '').strip()
        grid_id = visit.get('audio_gridfs_id')
        temp_gridfs_file = None
        try:
            if audio_path:
                norm = audio_path.replace('\\', '/')
                if 'static/uploads/visit_recordings/' in norm:
                    sub = norm.split('static/uploads/visit_recordings/', 1)[1]
                    cand = os.path.join(app.root_path, 'static', 'uploads', 'visit_recordings', sub)
                    if os.path.exists(cand):
                        audio_path = cand
                elif not os.path.isabs(audio_path):
                    audio_path = os.path.join(app.root_path, audio_path)

            if (not audio_path or not os.path.exists(audio_path)) and grid_id:
                import gridfs, tempfile
                from bson import ObjectId
                fs = gridfs.GridFS(database.get_db())
                if fs.exists(ObjectId(grid_id)):
                    g_file = fs.get(ObjectId(grid_id))
                    temp_fd, temp_gridfs_file = tempfile.mkstemp(suffix='.webm')
                    with os.fdopen(temp_fd, 'wb') as t_out:
                        t_out.write(g_file.read())
                    audio_path = temp_gridfs_file

            if audio_path and os.path.exists(audio_path):
                mime = visit.get('audio_mime_type') or 'audio/webm'
                dur = float(visit.get('audio_duration_seconds') or 0)
                re_visit = database.attach_audio_and_process_visit(
                    visit_id=visit_id,
                    audio_file_path=audio_path,
                    audio_url=visit.get('audio_url'),
                    duration_seconds=dur,
                    mime_type=mime,
                    company_id=cid,
                    gridfs_id=grid_id
                )
                if re_visit:
                    visit = re_visit
        except Exception as auto_re_err:
            print("Auto-reprocess transcript notice:", auto_re_err)
        finally:
            if temp_gridfs_file and os.path.exists(temp_gridfs_file):
                try:
                    os.remove(temp_gridfs_file)
                except Exception:
                    pass

    resp = jsonify({
        'success': True,
        'visit_id': visit.get('id'),
        'customer_name': visit.get('customer_name'),
        'employee_name': visit.get('employee_name'),
        'has_audio_recording': visit.get('has_audio_recording', False),
        'audio_url': visit.get('audio_url'),
        'audio_duration_seconds': visit.get('audio_duration_seconds', 0),
        'language_detected': visit.get('language_detected', 'Tamil (ta-IN) / Tanglish'),
        'transcript_tamil': visit.get('conversation_transcript_tamil') or visit.get('conversation_transcript', ''),
        'transcript_english': visit.get('conversation_transcript_english', ''),
        'summary_tamil': visit.get('ai_summary_tamil', ''),
        'summary_english': visit.get('ai_summary_english') or visit.get('ai_summary', ''),
        'customer_sentiment': visit.get('ai_sentiment', 'Interested'),
        'customer_sentiment_tamil': visit.get('ai_sentiment_tamil', 'ஆர்வம்'),
        'sentiment_reason': visit.get('ai_sentiment_reason', ''),
        'action_items': visit.get('ai_action_items', []),
        'order_amount': visit.get('order_amount', 0.0),
        'ai_suggested_order_amount': visit.get('ai_suggested_order_amount')
    })
    resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    return resp


@app.route('/api/field-worker/visits/<visit_id>/reprocess-ai', methods=['POST'])
def api_field_worker_visit_reprocess_ai(visit_id):
    """Re-runs Tamil Gemini 1.5 Flash STT on an existing visit audio from GridFS or disk."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    visit = database.get_field_customer_visit_by_id(visit_id, company_id=cid)
    if not visit and (session.get('role') == 'super_admin' or cid in ['ALL', 'ARGUS_MASTER', None]):
        visit = database.get_field_customer_visit_by_id(visit_id, company_id=None)

    if not visit or not visit.get('has_audio_recording'):
        return jsonify({'success': False, 'error': 'No audio recording found for this visit.'}), 404

    mime = visit.get('audio_mime_type') or 'audio/webm'
    audio_path = str(visit.get('audio_file_path') or '').strip()

    # Dynamic path normalization
    if audio_path:
        norm_path = audio_path.replace('\\', '/')
        if 'static/uploads/visit_recordings/' in norm_path:
            rel_sub = norm_path.split('static/uploads/visit_recordings/', 1)[1]
            local_cand = os.path.join(app.root_path, 'static', 'uploads', 'visit_recordings', rel_sub)
            if os.path.exists(local_cand):
                audio_path = local_cand
        elif not os.path.isabs(audio_path):
            audio_path = os.path.join(app.root_path, audio_path)

    # Check GridFS if not on disk
    grid_id = visit.get('audio_gridfs_id')
    temp_gridfs_file = None
    if (not audio_path or not os.path.exists(audio_path)) and grid_id:
        try:
            import gridfs, tempfile
            from bson import ObjectId
            fs = gridfs.GridFS(database.get_db())
            if fs.exists(ObjectId(grid_id)):
                g_file = fs.get(ObjectId(grid_id))
                temp_fd, temp_gridfs_file = tempfile.mkstemp(suffix='.webm')
                with os.fdopen(temp_fd, 'wb') as t_out:
                    t_out.write(g_file.read())
                audio_path = temp_gridfs_file
        except Exception as e:
            print("Error retrieving GridFS audio for reprocess:", e)

    if not audio_path or not os.path.exists(audio_path):
        return jsonify({
            'success': False,
            'error': 'Audio file was stored on ephemeral server storage and expired during server redeployment. Please record a new visit — all new audio files are permanently preserved in Cloud Database!'
        }), 404

    try:
        updated = database.attach_audio_and_process_visit(
            visit_id=visit_id,
            audio_file_path=audio_path,
            audio_url=visit.get('audio_url'),
            duration_seconds=visit.get('audio_duration_seconds', 0),
            mime_type=mime,
            company_id=cid,
            gridfs_id=grid_id
        )
    finally:
        if temp_gridfs_file and os.path.exists(temp_gridfs_file):
            try:
                os.remove(temp_gridfs_file)
            except Exception:
                pass

    return jsonify({
        'success': True,
        'message': 'Tamil AI audio analysis completed successfully!',
        'visit': updated
    })


# =========================================================================
# FIELD WORKER PRODUCT WAREHOUSE & TRANSCRIPT EXPORT ROUTES
# =========================================================================

@app.route('/api/field-worker/products', methods=['GET'])
def api_field_worker_products():
    '''Returns dynamic Product Warehouse items and pricing for the active company.'''
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    force_sync = request.args.get('sync') == '1' or request.args.get('refresh') == '1'
    
    settings = database.get_warehouse_settings(company_id=cid)
    sheet_url = str(settings.get('sheet_url') or '').strip()
    
    # Auto-sync from Google Sheets if explicitly requested or if products lack qty field
    existing_prods = settings.get('products') or []
    needs_qty_sync = bool(sheet_url and existing_prods and any('qty' not in p or p.get('qty') is None for p in existing_prods))
    
    if sheet_url and (force_sync or needs_qty_sync):
        try:
            database.sync_warehouse_products_from_sheet(sheet_url=sheet_url, company_id=cid)
            settings = database.get_warehouse_settings(company_id=cid)
        except Exception as sync_err:
            print(f"Auto-sync Google Sheets products warning for {cid}: {sync_err}")

    products = database.get_warehouse_products(company_id=cid)
    resp = jsonify({
        'success': True,
        'company_id': cid,
        'products': products,
        'product_count': len(products),
        'sheet_url': settings.get('sheet_url', ''),
        'last_synced_at': settings.get('last_synced_at', ''),
        'sync_status': settings.get('sync_status', 'Ready')
    })
    resp.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate, max-age=0'
    resp.headers['Pragma'] = 'no-cache'
    resp.headers['Expires'] = '0'
    return resp


@app.route('/api/field-worker/products/sync', methods=['POST'])
def api_field_worker_products_sync():
    '''Manually or automatically refreshes products from the configured Google Sheet.'''
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    data = request.get_json(silent=True) or request.form.to_dict() or {}
    sheet_url = data.get('sheet_url')
    res = database.sync_warehouse_products_from_sheet(sheet_url=sheet_url, company_id=cid)
    return jsonify(res)


@app.route('/api/field-worker/products/sheet-url', methods=['POST'])
def api_field_worker_products_save_url():
    '''Saves or updates the Google Sheets URL for the Product Warehouse.'''
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    data = request.get_json(silent=True) or request.form.to_dict() or {}
    sheet_url = data.get('sheet_url', '').strip()
    res = database.save_warehouse_sheet_url(sheet_url=sheet_url, company_id=cid)
    return jsonify({
        'success': True,
        'message': 'Google Sheets Product Warehouse URL saved and synced successfully!',
        'result': res
    })


@app.route('/api/field-worker/visits/<visit_id>/transcript/export', methods=['GET'])
def api_field_worker_visit_transcript_export(visit_id):
    '''Generates and downloads a clean Speech-to-Text audit document (.txt).'''
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    doc = database.generate_visit_transcript_document(visit_id, company_id=cid)
    if not doc and (session.get('role') == 'super_admin' or cid in ['ALL', 'ARGUS_MASTER', None]):
        doc = database.generate_visit_transcript_document(visit_id, company_id=None)
    if not doc:
        return jsonify({'success': False, 'error': 'Customer visit transcript not found.'}), 404

    from flask import Response
    return Response(
        doc['content'],
        mimetype='text/plain; charset=utf-8',
        headers={
            'Content-Disposition': f'attachment; filename="{doc["filename"]}"',
            'Cache-Control': 'no-cache'
        }
    )


@app.route('/api/field-worker/visits/<visit_id>/transcript/docx', methods=['GET'])
def api_field_worker_visit_transcript_docx(visit_id):
    '''Generates and downloads a clean, professional Speech-to-Text audit document in Microsoft Word (.docx) format only.'''
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    doc_res = database.generate_visit_transcript_docx(visit_id, company_id=cid)
    if not doc_res and (session.get('role') == 'super_admin' or cid in ['ALL', 'ARGUS_MASTER', None]):
        doc_res = database.generate_visit_transcript_docx(visit_id, company_id=None)
    if not doc_res:
        return jsonify({'success': False, 'error': 'Customer visit transcript not found.'}), 404

    return send_file(
        doc_res['buffer'],
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=doc_res['filename']
    )


@app.route('/api/system/gemini-key', methods=['GET', 'POST'])
def api_system_gemini_key():
    '''Allows admin to view/set the Google Gemini API Key for Indian Tamil audio transcription.'''
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    db = database.get_db()
    if request.method == 'GET':
        env_key = os.environ.get('GEMINI_API_KEY') or os.environ.get('GOOGLE_API_KEY')
        db_key_doc = db.system_settings.find_one({'key': 'gemini_api_key'})
        saved_key = db_key_doc.get('value') if db_key_doc else None
        active_key = env_key or saved_key
        has_key = bool(active_key)
        masked_key = f"{active_key[:4]}...{active_key[-4:]}" if (active_key and len(active_key) > 8) else ("Configured" if has_key else "Not Set")
        return jsonify({
            'success': True,
            'has_key': has_key,
            'masked_key': masked_key,
            'source': 'Environment' if env_key else ('Database' if saved_key else 'None')
        })

    data = request.get_json(silent=True) or {}
    new_key = str(data.get('gemini_api_key') or '').strip()
    if not new_key:
        return jsonify({'success': False, 'error': 'Gemini API Key cannot be empty.'}), 400

    db.system_settings.update_one(
        {'key': 'gemini_api_key'},
        {'$set': {'key': 'gemini_api_key', 'value': new_key, 'updated_at': database.get_ist_now()}},
        upsert=True
    )
    os.environ['GEMINI_API_KEY'] = new_key

    return jsonify({
        'success': True,
        'message': 'Google Gemini API Key saved successfully! AI speech-to-text is now fully active.'
    })

@app.route('/api/ai/teacher/explain', methods=['POST'])
def api_teacher_ai_explain():
    """Contextual explanation endpoint for Teacher AI powered by Google Gemini (with instant offline fallback)."""
    try:
        data = request.get_json(silent=True) or {}
        element_title = str(data.get('title') or data.get('label') or '').strip()
        element_tag = str(data.get('tag') or '').strip()
        category = str(data.get('category') or 'Feature').strip()
        page_context = str(data.get('context') or data.get('page') or '').strip()

        if not element_title:
            return jsonify({'success': False, 'error': 'Element title or label required'}), 400

        # Check in-memory cache to avoid duplicate API calls
        cache_key = f"{page_context}::{element_tag}::{element_title}".lower()
        if hasattr(app, '_teacher_ai_cache') and cache_key in app._teacher_ai_cache:
            return jsonify({'success': True, 'cached': True, **app._teacher_ai_cache[cache_key]})

        # If Gemini API key is configured, query Gemini for a crisp, concise 1-2 sentence explanation
        gemini_key = os.environ.get('GEMINI_API_KEY') or os.environ.get('GOOGLE_API_KEY')
        if not gemini_key:
            db = database.get_db()
            setting = db.system_settings.find_one({'key': 'gemini_api_key'})
            if setting and setting.get('value'):
                gemini_key = str(setting['value']).strip()

        if gemini_key:
            try:
                import requests
                headers = {'Content-Type': 'application/json'}
                prompt = (
                    "You are Teacher AI for the Argus Technologies AI Attendance & Operations Management Web Application. "
                    "Explain clearly and concisely in exactly 1 or 2 sentences what this UI control does and how it affects attendance or operations.\n"
                    f"Element Name: {element_title}\n"
                    f"Element Type: {element_tag}\n"
                    f"Screen/Context: {page_context}\n"
                    f'Respond ONLY with a valid JSON object in this format: {{"title": "{element_title}", "category": "{category}", "text": "...", "tip": "..."}}'
                )
                model_name = 'gemini-2.0-flash'
                url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={gemini_key}"
                payload = {
                    "contents": [{"parts": [{"text": prompt}]}],
                    "generationConfig": {"temperature": 0.2, "maxOutputTokens": 200}
                }
                resp = requests.post(url, headers=headers, json=payload, timeout=4)
                if resp.status_code == 200:
                    res_json = resp.json()
                    candidates = res_json.get('candidates', [])
                    if candidates:
                        text_resp = candidates[0].get('content', {}).get('parts', [{}])[0].get('text', '')
                        import re
                        m = re.search(r'\{.*\}', text_resp, re.DOTALL)
                        if m:
                            parsed = json.loads(m.group(0))
                            if not hasattr(app, '_teacher_ai_cache'):
                                app._teacher_ai_cache = {}
                            app._teacher_ai_cache[cache_key] = parsed
                            return jsonify({'success': True, 'ai_generated': True, **parsed})
            except Exception as gemini_err:
                print(f"Warning: Teacher AI Gemini generation notice: {gemini_err}")

        # Fallback explanation
        fallback_res = {
            'title': element_title,
            'category': category,
            'text': f"{element_title} is an interactive control in the Argus Attendance system.",
            'tip': None
        }
        return jsonify({'success': True, 'ai_generated': False, **fallback_res})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500



# =========================================================================
# ARGUS CO-WORKER AI (AGENT 2) - AUTONOMOUS EXECUTIVE ASSISTANT
# =========================================================================
@app.route('/api/ai/coworker/transcribe', methods=['POST'])
def api_coworker_ai_transcribe():
    """Transcribes natural human voice notes using Google Gemini Multimodal Audio."""
    try:
        audio_file = request.files.get('audio_file') or request.files.get('audio')
        if not audio_file:
            return jsonify({'success': False, 'error': 'No audio file provided'}), 400

        audio_bytes = audio_file.read()
        if not audio_bytes:
            return jsonify({'success': False, 'error': 'Empty audio recording'}), 400

        import base64
        import requests
        audio_b64 = base64.b64encode(audio_bytes).decode('utf-8')
        mime_type = audio_file.content_type or 'audio/webm'

        gemini_key = os.environ.get('GEMINI_API_KEY') or os.environ.get('GOOGLE_API_KEY')
        if not gemini_key:
            db = database.get_db()
            setting = db.system_settings.find_one({'key': 'gemini_api_key'})
            if setting and setting.get('value'):
                gemini_key = str(setting['value']).strip()

        if not gemini_key:
            return jsonify({'success': False, 'error': 'Gemini API Key is not configured for voice transcription'}), 400

        prompt = (
            "You are an expert AI speech transcriptionist for the Argus Web App. "
            "Listen to this natural voice message carefully. The speaker may speak English, Indian English, Tamil, Hindi, or Tanglish at variable speeds. "
            "Transcribe verbatim what the user said into clear text. Do NOT add preamble or meta-commentary. Return ONLY the transcribed text."
        )

        candidate_models = ['gemini-2.0-flash', 'gemini-1.5-flash', 'gemini-flash-latest']
        for model in candidate_models:
            try:
                url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={gemini_key}"
                payload = {
                    "contents": [{
                        "parts": [
                            {"text": prompt},
                            {"inline_data": {"mime_type": mime_type, "data": audio_b64}}
                        ]
                    }],
                    "generationConfig": {"temperature": 0.1, "maxOutputTokens": 250}
                }
                resp = requests.post(url, headers={'Content-Type': 'application/json'}, json=payload, timeout=25)
                if resp.status_code == 200:
                    res_json = resp.json()
                    candidates = res_json.get('candidates', [])
                    if candidates:
                        transcript = candidates[0].get('content', {}).get('parts', [{}])[0].get('text', '').strip()
                        if transcript:
                            return jsonify({'success': True, 'transcript': transcript})
            except Exception as e:
                print(f"Notice: Gemini transcribe candidate {model} error: {e}")

        return jsonify({'success': False, 'error': 'Could not transcribe audio. Please try again or type directly.'}), 500
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/ai/coworker/chat', methods=['POST'])
def api_coworker_ai_chat():
    """Autonomous executive assistant endpoint supporting all Argus operations with multi-turn conversational memory & task execution."""
    try:
        data = request.get_json(silent=True) or {}
        user_message = str(data.get('message') or '').strip()
        history = data.get('history') or []
        context = data.get('context') or {}
        pending_task = data.get('pending_task') or session.get('coworker_pending_task') or None

        if not user_message:
            return jsonify({'success': False, 'error': 'Message cannot be empty'}), 400

        # Check for explicit cancellation
        msg_lower = user_message.lower().strip()
        if msg_lower in ['cancel', 'cancel task', 'never mind', 'stop', 'abort', 'leave it', 'dont do it', "don't do it", 'forget it']:
            session['coworker_pending_task'] = None
            return jsonify({
                'success': True,
                'reply': "Understood, I've cancelled that pending request. How else can I assist you?",
                'task_executed': False,
                'pending_task': None
            })

        db = database.get_db()
        cid = str(get_current_company_id() or 'ARGUS_MASTER')

        # Load active registered employees for entity matching (strictly tenant-isolated)
        emp_filter = {'status': {'$ne': 'Terminated'}}
        if cid and cid != 'ARGUS_MASTER':
            emp_filter['$or'] = [{'company_id': cid}, {'company_id': {'$in': [None, '', 'ARGUS_MASTER']}}]
        emp_cursor = db.employees.find(emp_filter, {
            'employee_name': 1, 'id': 1, 'department': 1, 'phone': 1, 'mobile': 1,
            'designation': 1, 'basic_salary': 1, 'salary_type': 1, 'shift_hours': 1,
            'branch_id': 1, 'company_id': 1, 'created_at': 1, 'status': 1
        })
        employees = list(emp_cursor)
        emp_names = [e.get('employee_name') for e in employees if e.get('employee_name')]

        def resolve_employee(name):
            if not name:
                return None, []
            import difflib
            raw = str(name).strip()
            # Clean common trailing/leading prepositions or punctuation
            cleaned = re.sub(r'^(?:for|to|of|about)\s+', '', raw, flags=re.IGNORECASE).strip()
            cleaned = re.sub(r'\s+(?:for|to|of|as|with|at|in|on|rs|inr)$', '', cleaned, flags=re.IGNORECASE).strip()
            if not cleaned:
                return None, []

            target = cleaned.lower()

            # 1. Exact match (case-insensitive)
            for e in employees:
                e_name = str(e.get('employee_name') or '').strip()
                if e_name.lower() == target:
                    return e, []

            # 2. Token match / Substring match
            target_tokens = set(target.split())
            token_candidates = []
            for e in employees:
                e_name = str(e.get('employee_name') or '').strip()
                e_lower = e_name.lower()
                e_tokens = set(e_lower.split())
                if target_tokens and (target_tokens.issubset(e_tokens) or any(t in e_tokens for t in target_tokens if len(t) >= 4)):
                    token_candidates.append(e)
                elif len(target) >= 3 and (target in e_lower or e_lower in target):
                    token_candidates.append(e)

            if len(token_candidates) == 1:
                return token_candidates[0], []
            elif len(token_candidates) > 1:
                best_match = None
                best_score = 0.0
                for c in token_candidates:
                    score = difflib.SequenceMatcher(None, target, str(c.get('employee_name') or '').lower()).ratio()
                    if score > best_score:
                        best_score = score
                        best_match = c
                if best_score >= 0.80:
                    return best_match, []
                return None, token_candidates

            # 3. Fuzzy similarity match (handles spelling variations like Mathew vs Matthew)
            fuzzy_matches = []
            for e in employees:
                e_name = str(e.get('employee_name') or '').strip()
                sim = difflib.SequenceMatcher(None, target, e_name.lower()).ratio()
                if sim >= 0.76:
                    fuzzy_matches.append((sim, e))

            if fuzzy_matches:
                fuzzy_matches.sort(key=lambda x: x[0], reverse=True)
                if len(fuzzy_matches) == 1 or (fuzzy_matches[0][0] - fuzzy_matches[1][0] >= 0.12):
                    return fuzzy_matches[0][1], []
                return None, [m[1] for m in fuzzy_matches]

            return None, []

        def find_context_employee():
            """Recovers the most relevant employee from pending task or recent message turns."""
            if pending_task and isinstance(pending_task, dict):
                p_emp = (pending_task.get('parameters') or {}).get('employee_name')
                if p_emp:
                    doc, _ = resolve_employee(p_emp)
                    if doc:
                        return doc.get('employee_name')
            if isinstance(history, list):
                for h in reversed(history[-6:]):
                    task_info = h.get('task')
                    if isinstance(task_info, dict) and task_info.get('details'):
                        d_emp = task_info['details'].get('Employee')
                        if d_emp:
                            doc, _ = resolve_employee(d_emp)
                            if doc:
                                return doc.get('employee_name')
                    h_text = str(h.get('text') or '')
                    for en in emp_names:
                        if re.search(r'\b' + re.escape(en) + r'\b', h_text, re.IGNORECASE):
                            return en
            return None

        parsed_action = None

        # Tier 2: Fast Gemini NLU Attempt with strict 3.5s timeout guard
        gemini_key = os.environ.get('GEMINI_API_KEY') or os.environ.get('GOOGLE_API_KEY')
        if not gemini_key:
            setting = db.system_settings.find_one({'key': 'gemini_api_key'})
            if setting and setting.get('value'):
                gemini_key = str(setting['value']).strip()

        if gemini_key:
            try:
                import requests
                emp_list_str = ", ".join(emp_names[:60])
                today_str = datetime.now().strftime('%Y-%m-%d')
                clean_msg = user_message.replace('"', '\"').replace('\n', ' ')

                history_lines = []
                if isinstance(history, list):
                    for h in history[-6:]:
                        sender = 'User' if h.get('sender') == 'user' else 'Assistant'
                        txt = str(h.get('text') or '').replace('\n', ' ')
                        history_lines.append(f"{sender}: {txt}")
                transcript_str = "\n".join(history_lines) if history_lines else "None"
                pending_str = json.dumps(pending_task) if pending_task else "None"

                prompt = (
                    "You are Co-Worker AI, an operations executive assistant for the Argus Attendance Web App.\n"
                    f"Current Date: {today_str}\n"
                    f"Known Employees in this Company: [{emp_list_str}]\n\n"
                    f"Recent Conversation Transcript:\n{transcript_str}\n\n"
                    f"Active In-Progress Pending Task:\n{pending_str}\n\n"
                    f'Latest User Message: "{clean_msg}"\n\n'
                    "Actions available in Argus:\n"
                    "1. 'create_manual_attendance': Parameters: { 'employee_name': string, 'date': string, 'hours': string, 'status': string, 'reason': string }\n"
                    "2. 'create_payment': Parameters: { 'employee_name': string, 'amount': number, 'status': string, 'payment_type': string, 'date': string, 'reason': string }\n"
                    "3. 'create_advance': Parameters: { 'employee_name': string, 'amount': number, 'date': string, 'reason': string }\n"
                    "4. 'check_payments': Parameters: { 'employee_name': string, 'query_type': 'balance'|'history'|'summary' }\n"
                    "5. 'check_attendance': Parameters: { 'employee_name': string, 'date': string }\n"
                    "6. 'apply_leave': Parameters: { 'employee_name': string, 'leave_type': string, 'from_date': string, 'to_date': string, 'reason': string }\n"
                    "7. 'get_employee_info': Parameters: { 'employee_name': string, 'field': string }\n"
                    "8. 'get_salary_info': Parameters: { 'employee_name': string }\n"
                    "9. 'check_field_worker': Parameters: { 'target': string }\n"
                    "10. 'navigate': Parameters: { 'target': 'manual_entry'|'payment_entry'|'salary_report'|'field_worker'|'employees'|'leaves'|'live_report'|'attendance_report'|'payslips'|'dashboard' }\n"
                    "11. 'general_chat': Conversational response.\n\n"
                    "Respond ONLY with a valid JSON object matching this schema:\n"
                    "{\n"
                    '  "action": string,\n'
                    '  "parameters": { ... },\n'
                    '  "reply_text": string\n'
                    "}"
                )

                url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-flash-latest:generateContent?key={gemini_key}"
                payload = {
                    "contents": [{"parts": [{"text": prompt}]}],
                    "generationConfig": {"temperature": 0.1, "maxOutputTokens": 350}
                }
                resp = requests.post(url, headers={'Content-Type': 'application/json'}, json=payload, timeout=3.5)
                if resp.status_code == 200:
                    cand = resp.json().get('candidates', [])
                    if cand:
                        t = cand[0].get('content', {}).get('parts', [{}])[0].get('text', '')
                        m = re.search(r'\{.*\}', t, re.DOTALL)
                        if m:
                            parsed_action = json.loads(m.group(0))
            except Exception as e:
                pass

        # Tier 1: Deterministic Multi-Turn & Feature NLU Engine (Zero Latency, Complete Offline Support)
        if not parsed_action or not parsed_action.get('action') or parsed_action.get('action') == 'general_chat':
            msg_lower = user_message.lower().strip()

            # A. Check active in-progress pending tasks
            if pending_task and isinstance(pending_task, dict):
                p_action = pending_task.get('action')
                p_params = pending_task.get('parameters') or {}

                # 1. Pending Manual Attendance Continuation
                if p_action == 'create_manual_attendance':
                    # Extract employee if provided now or already saved
                    curr_emp = p_params.get('employee_name')
                    if not curr_emp:
                        # Check if user mentioned an employee
                        for e in employees:
                            en = str(e.get('employee_name') or '')
                            if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                                curr_emp = en
                                break
                        if not curr_emp:
                            # Try resolving from raw message tokens
                            clean_candidate = re.sub(r'\b(?:hours?|hrs?|proper|improper|half\s*day|full\s*day|for|to|of|add|log|today|yesterday)\b', '', msg_lower, flags=re.IGNORECASE).strip()
                            if clean_candidate:
                                c_doc, _ = resolve_employee(clean_candidate)
                                if c_doc:
                                    curr_emp = c_doc.get('employee_name')

                    # Extract hours
                    hours_val = p_params.get('hours') or '08:00'
                    hr_m = re.search(r'(\d{1,2}(?::\d{2})?)\s*(?:hrs?|hours?)?', msg_lower)
                    if 'half day' in msg_lower or '4 hour' in msg_lower or '4 hr' in msg_lower:
                        hours_val = '04:00'
                    elif 'full day' in msg_lower or '8 hour' in msg_lower or '8 hr' in msg_lower:
                        hours_val = '08:00'
                    elif hr_m:
                        raw_h = hr_m.group(1)
                        hours_val = f"{int(raw_h):02d}:00" if ':' not in raw_h else raw_h

                    # Extract status
                    st_val = p_params.get('status') or 'Proper'
                    if 'improper' in msg_lower:
                        st_val = 'Improper'
                    elif 'half day' in msg_lower:
                        st_val = 'Half Day'
                    elif 'late' in msg_lower:
                        st_val = 'Late'
                    elif 'overtime' in msg_lower or 'ot' in msg_lower:
                        st_val = 'Overtime'
                    elif 'proper' in msg_lower:
                        st_val = 'Proper'

                    # Extract date
                    date_val = p_params.get('date') or datetime.now().strftime('%Y-%m-%d')
                    if 'yesterday' in msg_lower:
                        date_val = (datetime.now() - timedelta(days=1)).strftime('%Y-%m-%d')
                    elif 'tomorrow' in msg_lower:
                        date_val = (datetime.now() + timedelta(days=1)).strftime('%Y-%m-%d')

                    parsed_action = {
                        'action': 'create_manual_attendance',
                        'parameters': {
                            'employee_name': curr_emp,
                            'hours': hours_val,
                            'status': st_val,
                            'date': date_val,
                            'reason': p_params.get('reason') or 'Recorded by Co-Worker AI'
                        },
                        'reply_text': ''
                    }

                # 2. Pending Payment Continuation
                elif p_action == 'create_payment':
                    amt_match = re.search(r'(?:rs\.?|inr|rupees|₹)?\s*(\d+(?:\.\d+)?)', user_message, re.IGNORECASE)
                    st_val = None
                    for cand_st in ['advance repayment', 'advance', 'salary', 'incentive', 'bonus', 'allowance', 'petrol', 'other']:
                        if cand_st in msg_lower:
                            st_val = cand_st.title()
                            break

                    curr_emp = p_params.get('employee_name')
                    if not curr_emp:
                        for e in employees:
                            en = str(e.get('employee_name') or '')
                            if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                                curr_emp = en
                                break

                    amt_val = float(amt_match.group(1)) if amt_match else p_params.get('amount')
                    parsed_action = {
                        'action': 'create_payment',
                        'parameters': {
                            'employee_name': curr_emp,
                            'amount': amt_val,
                            'status': st_val or p_params.get('status') or 'Advance Repayment',
                            'payment_type': p_params.get('payment_type') or 'UPI',
                            'date': p_params.get('date') or datetime.now().strftime('%Y-%m-%d'),
                            'reason': p_params.get('reason') or ''
                        },
                        'reply_text': ''
                    }

                # 3. Pending Advance Continuation
                elif p_action == 'create_advance':
                    amt_match = re.search(r'(?:rs\.?|inr|rupees|₹)?\s*(\d+(?:\.\d+)?)', user_message, re.IGNORECASE)
                    amt_val = float(amt_match.group(1)) if amt_match else p_params.get('amount')
                    curr_emp = p_params.get('employee_name')
                    if not curr_emp:
                        for e in employees:
                            en = str(e.get('employee_name') or '')
                            if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                                curr_emp = en
                                break
                    parsed_action = {
                        'action': 'create_advance',
                        'parameters': {
                            'employee_name': curr_emp,
                            'amount': amt_val,
                            'date': p_params.get('date') or datetime.now().strftime('%Y-%m-%d'),
                            'reason': p_params.get('reason') or 'Staff Advance'
                        },
                        'reply_text': ''
                    }

                # 4. Pending Leave Continuation
                elif p_action == 'apply_leave':
                    curr_emp = p_params.get('employee_name')
                    if not curr_emp:
                        for e in employees:
                            en = str(e.get('employee_name') or '')
                            if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                                curr_emp = en
                                break
                    leave_type = p_params.get('leave_type') or 'Casual Leave'
                    if 'sick' in msg_lower:
                        leave_type = 'Sick Leave'
                    elif 'earned' in msg_lower:
                        leave_type = 'Earned Leave'
                    parsed_action = {
                        'action': 'apply_leave',
                        'parameters': {
                            'employee_name': curr_emp,
                            'leave_type': leave_type,
                            'from_date': p_params.get('from_date') or datetime.now().strftime('%Y-%m-%d'),
                            'to_date': p_params.get('to_date') or datetime.now().strftime('%Y-%m-%d'),
                            'reason': p_params.get('reason') or 'Personal'
                        },
                        'reply_text': ''
                    }

            # B. Single-Turn Intent Parsing (If no pending task matched or continues)
            if not parsed_action:
                is_nav = any(w in msg_lower for w in ['open', 'go to', 'take me', 'navigate', 'show page', 'view page'])

                # 1. BALANCE & ADVANCE INQUIRY
                is_balance_inq = any(phrase in msg_lower for phrase in ['check balance', 'advance balance', 'payment balance', 'balance for', 'advance summary', 'what is the balance', 'outstanding advance', 'dues for']) or (msg_lower == 'check balance') or (msg_lower == 'balance')
                if is_balance_inq and not is_nav:
                    # Check if employee specified in text
                    target_emp = None
                    for e in employees:
                        en = str(e.get('employee_name') or '')
                        if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                            target_emp = en
                            break
                    if not target_emp:
                        bal_m = re.search(r'(?:balance|advance|dues)\s+(?:for|of)\s+([A-Za-z\s]+)', user_message, re.IGNORECASE)
                        if bal_m:
                            b_doc, _ = resolve_employee(bal_m.group(1))
                            if b_doc:
                                target_emp = b_doc.get('employee_name')
                    if not target_emp:
                        # Fallback to context employee
                        target_emp = find_context_employee()

                    parsed_action = {
                        'action': 'check_payments',
                        'parameters': {'employee_name': target_emp, 'query_type': 'balance'},
                        'reply_text': ''
                    }

                # 2. MANUAL ATTENDANCE ENTRY
                elif ('manual' in msg_lower and any(w in msg_lower for w in ['entry', 'attendance', 'create', 'add', 'log', 'record'])) and not is_nav:
                    # Check for employee
                    target_emp = None
                    for e in employees:
                        en = str(e.get('employee_name') or '')
                        if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                            target_emp = en
                            break
                    if not target_emp:
                        emp_m = re.search(r'(?:for|to)\s+([A-Za-z\s]+?)(?:\s+(?:\d+|proper|improper|half|today|yesterday)|$)', user_message, re.IGNORECASE)
                        if emp_m:
                            m_doc, _ = resolve_employee(emp_m.group(1))
                            if m_doc:
                                target_emp = m_doc.get('employee_name')

                    # Hours & Status
                    hours_val = '08:00'
                    if 'half day' in msg_lower or '4 hour' in msg_lower or '4 hr' in msg_lower:
                        hours_val = '04:00'
                    elif 'full day' in msg_lower or '8 hour' in msg_lower or '8 hr' in msg_lower:
                        hours_val = '08:00'
                    else:
                        hr_m = re.search(r'(\d{1,2}(?::\d{2})?)\s*(?:hrs?|hours?)?', msg_lower)
                        if hr_m:
                            raw_h = hr_m.group(1)
                            hours_val = f"{int(raw_h):02d}:00" if ':' not in raw_h else raw_h

                    st_val = 'Proper'
                    if 'improper' in msg_lower:
                        st_val = 'Improper'
                    elif 'half day' in msg_lower:
                        st_val = 'Half Day'
                    elif 'late' in msg_lower:
                        st_val = 'Late'
                    elif 'overtime' in msg_lower:
                        st_val = 'Overtime'

                    date_val = datetime.now().strftime('%Y-%m-%d')
                    if 'yesterday' in msg_lower:
                        date_val = (datetime.now() - timedelta(days=1)).strftime('%Y-%m-%d')

                    parsed_action = {
                        'action': 'create_manual_attendance',
                        'parameters': {
                            'employee_name': target_emp,
                            'hours': hours_val,
                            'status': st_val,
                            'date': date_val,
                            'reason': 'Recorded by Co-Worker AI'
                        },
                        'reply_text': ''
                    }

                # 3. CREATE ADVANCE ENTRY
                elif re.search(r'(?:give|create|add|record|issue)?\s*advance\s+(?:of\s+)?(?:rs\.?|inr|₹)?\s*(\d+(?:\.\d+)?)\s*(?:to|for)\s*([A-Za-z\s]+)', user_message, re.IGNORECASE) or \
                     re.search(r'advance\s+(?:for|to)\s*([A-Za-z\s]+)\s*(?:of\s+)?(?:rs\.?|inr|₹)?\s*(\d+(?:\.\d+)?)', user_message, re.IGNORECASE):
                    adv_m1 = re.search(r'(?:give|create|add|record|issue)?\s*advance\s+(?:of\s+)?(?:rs\.?|inr|₹)?\s*(\d+(?:\.\d+)?)\s*(?:to|for)\s*([A-Za-z\s]+)', user_message, re.IGNORECASE)
                    adv_m2 = re.search(r'advance\s+(?:for|to)\s*([A-Za-z\s]+)\s*(?:of\s+)?(?:rs\.?|inr|₹)?\s*(\d+(?:\.\d+)?)', user_message, re.IGNORECASE)
                    if adv_m1:
                        amt = float(adv_m1.group(1))
                        emp_raw = adv_m1.group(2).strip()
                    else:
                        emp_raw = adv_m2.group(1).strip()
                        amt = float(adv_m2.group(2))
                    parsed_action = {
                        'action': 'create_advance',
                        'parameters': {'employee_name': emp_raw, 'amount': amt, 'date': datetime.now().strftime('%Y-%m-%d')},
                        'reply_text': ''
                    }

                # 4. CREATE PAYMENT ENTRY
                elif any(w in msg_lower for w in ['payment', 'pay ', 'paid ']) and not is_nav:
                    pay_match = re.search(r'(?:create|add|make|give|record)?\s*payment(?:\s+entry)?\s+(?:for\s+)?([A-Za-z\s]+?)\s+(?:for\s+|of\s+|rs\.?|inr|₹)?\s*(\d+(?:\.\d+)?)(?:\s+(?:as|set|with|status)?\s*(.*))?', user_message, re.IGNORECASE)
                    alt_pay = re.search(r'(?:give|pay)\s+([A-Za-z\s]+?)\s+(?:for\s+|of\s+|rs\.?|inr|₹)?\s*(\d+(?:\.\d+)?)', user_message, re.IGNORECASE)
                    partial_pay = re.search(r'(?:create|add|make|record)?\s*payment(?:\s+entry)?(?:\s+(?:for|to)\s+([A-Za-z\s]+))?', user_message, re.IGNORECASE)

                    if pay_match:
                        emp = pay_match.group(1).strip()
                        emp = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', emp, flags=re.IGNORECASE).strip()
                        amt = float(pay_match.group(2))
                        st = (pay_match.group(3) or 'Advance Repayment').strip()
                        parsed_action = {
                            'action': 'create_payment',
                            'parameters': {'employee_name': emp, 'amount': amt, 'status': st},
                            'reply_text': ''
                        }
                    elif alt_pay:
                        emp = alt_pay.group(1).strip()
                        emp = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', emp, flags=re.IGNORECASE).strip()
                        amt = float(alt_pay.group(2))
                        parsed_action = {
                            'action': 'create_payment',
                            'parameters': {'employee_name': emp, 'amount': amt, 'status': 'Advance Repayment'},
                            'reply_text': ''
                        }
                    elif partial_pay:
                        emp = partial_pay.group(1)
                        if emp:
                            emp = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', emp.strip(), flags=re.IGNORECASE).strip()
                        parsed_action = {
                            'action': 'create_payment',
                            'parameters': {'employee_name': emp, 'amount': None},
                            'reply_text': ''
                        }

                # 5. ATTENDANCE & PRESENCE INQUIRY
                elif any(p in msg_lower for p in ['present today', 'who is present', 'who is absent', 'attendance today', 'how many present', 'punch in today', 'is present']) and not is_nav:
                    target_emp = None
                    for e in employees:
                        en = str(e.get('employee_name') or '')
                        if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                            target_emp = en
                            break
                    parsed_action = {
                        'action': 'check_attendance',
                        'parameters': {'employee_name': target_emp},
                        'reply_text': ''
                    }

                # 6. APPLY LEAVE
                elif ('leave' in msg_lower or 'permission' in msg_lower) and any(w in msg_lower for w in ['apply', 'request', 'take', 'for', 'sick', 'casual']) and not is_nav:
                    target_emp = None
                    for e in employees:
                        en = str(e.get('employee_name') or '')
                        if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                            target_emp = en
                            break
                    l_type = 'Casual Leave'
                    if 'sick' in msg_lower:
                        l_type = 'Sick Leave'
                    elif 'earned' in msg_lower:
                        l_type = 'Earned Leave'
                    elif 'permission' in msg_lower:
                        l_type = 'Permission'
                    parsed_action = {
                        'action': 'apply_leave',
                        'parameters': {'employee_name': target_emp, 'leave_type': l_type},
                        'reply_text': ''
                    }

                # 7. EMPLOYEE DIRECTORY & CONTACT
                elif any(p in msg_lower for p in ['phone number', 'mobile', 'contact', 'email', 'department', 'designation', 'who is', 'details of', 'info of', 'profile of']) and not is_nav:
                    target_emp = None
                    for e in employees:
                        en = str(e.get('employee_name') or '')
                        if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                            target_emp = en
                            break
                    if not target_emp:
                        e_m = re.search(r'(?:of|for|is)\s+([A-Za-z\s]+)', user_message, re.IGNORECASE)
                        if e_m:
                            e_doc, _ = resolve_employee(e_m.group(1))
                            if e_doc:
                                target_emp = e_doc.get('employee_name')
                    parsed_action = {
                        'action': 'get_employee_info',
                        'parameters': {'employee_name': target_emp},
                        'reply_text': ''
                    }

                # 8. SALARY & PAYSLIP INQUIRY
                elif any(p in msg_lower for p in ['salary', 'payslip', 'pay slip', 'wages', 'earnings']) and not is_nav:
                    target_emp = None
                    for e in employees:
                        en = str(e.get('employee_name') or '')
                        if re.search(r'\b' + re.escape(en.lower()) + r'\b', msg_lower):
                            target_emp = en
                            break
                    if target_emp:
                        parsed_action = {
                            'action': 'get_salary_info',
                            'parameters': {'employee_name': target_emp},
                            'reply_text': ''
                        }

                # 9. NAVIGATION
                if not parsed_action and is_nav:
                    target = 'dashboard'
                    if 'manual' in msg_lower:
                        target = 'manual_entry'
                    elif 'payment' in msg_lower or 'ledger' in msg_lower or 'advance' in msg_lower:
                        target = 'payment_entry'
                    elif 'salary' in msg_lower:
                        target = 'salary_report'
                    elif 'payslip' in msg_lower:
                        target = 'payslips'
                    elif 'leave' in msg_lower:
                        target = 'leaves'
                    elif 'live' in msg_lower:
                        target = 'live_report'
                    elif 'report' in msg_lower or 'attendance' in msg_lower:
                        target = 'attendance_report'
                    elif 'employee' in msg_lower or 'staff' in msg_lower:
                        target = 'employees'
                    elif 'field' in msg_lower:
                        target = 'field_worker'
                    elif 'profile' in msg_lower or 'company' in msg_lower:
                        target = 'company_profile'
                    elif 'support' in msg_lower or 'help' in msg_lower:
                        target = 'support'
                    parsed_action = {
                        'action': 'navigate',
                        'parameters': {'target': target},
                        'reply_text': f"Navigating to {target.replace('_', ' ').title()}."
                    }

                # 10. Fallback General Operations Chat
                if not parsed_action:
                    parsed_action = {
                        'action': 'general_chat',
                        'parameters': {},
                        'reply_text': f"I understand you said: '{user_message}'. As your Co-Worker AI assistant, I can log manual attendance, record payments or advances, check advance balances, inspect today's attendance, apply leaves, look up employee contacts, and navigate anywhere in Argus."
                    }

        # Task Execution Dispatcher
        action_name = parsed_action.get('action') or 'general_chat'
        params = parsed_action.get('parameters') or {}
        reply_text = parsed_action.get('reply_text') or ''

        # 1. CREATE MANUAL ATTENDANCE
        if action_name == 'create_manual_attendance':
            raw_name = params.get('employee_name')
            if raw_name:
                raw_name = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', str(raw_name).strip(), flags=re.IGNORECASE).strip()

            emp_doc = None
            if raw_name:
                emp_doc, candidates = resolve_employee(raw_name)
                if not emp_doc:
                    if candidates:
                        cand_names = ", ".join([f"**{c.get('employee_name')}**" for c in candidates[:4]])
                        return jsonify({
                            'success': True,
                            'reply': f"I found multiple employees matching '**{raw_name}**': {cand_names}. Which one would you like to log attendance for?",
                            'task_executed': False,
                            'pending_task': None
                        })
                    else:
                        sample_staff = ", ".join([f"**{e.get('employee_name')}**" for e in employees[:5] if e.get('employee_name')])
                        staff_hint = f" Active staff in your company: {sample_staff}." if sample_staff else ""
                        return jsonify({
                            'success': True,
                            'reply': f"I couldn't find an employee named '**{raw_name}**' registered in your company.{staff_hint} Please check the spelling.",
                            'task_executed': False,
                            'pending_task': None
                        })

            canonical_emp_name = emp_doc.get('employee_name') if emp_doc else raw_name

            # If employee is missing, initiate multi-turn state!
            if not canonical_emp_name:
                new_pending = {
                    'action': 'create_manual_attendance',
                    'parameters': {
                        'hours': params.get('hours') or '08:00',
                        'status': params.get('status') or 'Proper',
                        'date': params.get('date') or datetime.now().strftime('%Y-%m-%d')
                    }
                }
                session['coworker_pending_task'] = new_pending
                return jsonify({
                    'success': True,
                    'reply': "I'd be glad to record a manual entry! Which **employee** is this for, and what are the **hours** / **status**? (e.g., *'Add manual attendance for Philip Matthew today for 8 hours Proper'*).",
                    'task_executed': False,
                    'pending_task': new_pending
                })

            emp_name = canonical_emp_name
            date_str = params.get('date') or datetime.now().strftime('%Y-%m-%d')
            hours_str = params.get('hours') or '08:00'
            status = params.get('status') or 'Proper'
            reason = params.get('reason') or 'Recorded by Co-Worker AI'

            entry_payload = {
                'employee_name': emp_name,
                'date': date_str,
                'hours': hours_str,
                'mode': 'Hours',
                'status': status,
                'reason': reason
            }
            inserted_id = database.create_manual_entry(entry_payload, company_id=cid)
            session['coworker_pending_task'] = None

            return jsonify({
                'success': True,
                'reply': f"I've successfully logged manual attendance for **{emp_name}** on **{date_str}** with status **{status}** ({hours_str} hrs). Entry ID is **#{inserted_id}**.",
                'task_executed': True,
                'task_data': {
                    'type': 'attendance',
                    'details': {
                        'Employee': emp_name,
                        'Date': date_str,
                        'Hours': hours_str,
                        'Status': status,
                        'Entry ID': str(inserted_id)
                    },
                    'action_label': 'View Manual Entries',
                    'action_url': '/manual-entry'
                },
                'pending_task': None
            })

        # 2. CHECK PAYMENTS / BALANCE / ADVANCE LEDGER
        elif action_name == 'check_payments':
            raw_name = params.get('employee_name')
            if raw_name:
                raw_name = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', str(raw_name).strip(), flags=re.IGNORECASE).strip()

            emp_doc = None
            if raw_name:
                emp_doc, _ = resolve_employee(raw_name)

            if emp_doc or raw_name:
                emp_name = emp_doc.get('employee_name') if emp_doc else raw_name
                # Calculate advances and repayments
                adv_f = {'employee_name': emp_name}
                if cid and cid != 'ARGUS_MASTER':
                    adv_f['company_id'] = cid
                advances = list(db.advances.find(adv_f))
                total_adv = sum(float(a.get('amount') or 0.0) for a in advances)

                pay_f = {'employee_name': emp_name}
                if cid and cid != 'ARGUS_MASTER':
                    pay_f['company_id'] = cid
                payments = list(db.payments.find(pay_f))
                total_repaid = sum(float(p.get('amount') or 0.0) for p in payments if 'repayment' in str(p.get('status') or '').lower())
                total_other = sum(float(p.get('amount') or 0.0) for p in payments if 'repayment' not in str(p.get('status') or '').lower())
                outstanding_bal = max(0.0, total_adv - total_repaid)

                return jsonify({
                    'success': True,
                    'reply': f"Here is the advance & payment statement for **{emp_name}**:\n"
                             f"• Total Advance Issued: **₹{total_adv:,.2f}**\n"
                             f"• Total Advance Repaid: **₹{total_repaid:,.2f}**\n"
                             f"• **Outstanding Advance Balance: ₹{outstanding_bal:,.2f}**\n"
                             f"• Other Payouts/Disbursements: **₹{total_other:,.2f}**",
                    'task_executed': True,
                    'task_data': {
                        'type': 'balance',
                        'details': {
                            'Employee': emp_name,
                            'Total Advance': f"₹{total_adv:,.2f}",
                            'Total Repaid': f"₹{total_repaid:,.2f}",
                            'Outstanding Balance': f"₹{outstanding_bal:,.2f}",
                            'Other Payments': f"₹{total_other:,.2f}"
                        },
                        'action_label': 'View Payment Ledger',
                        'action_url': '/payment-entry'
                    },
                    'pending_task': None
                })
            else:
                # Company-wide advance summary
                adv_f = {}
                if cid and cid != 'ARGUS_MASTER':
                    adv_f['company_id'] = cid
                advances = list(db.advances.find(adv_f))
                tot_comp_adv = sum(float(a.get('amount') or 0.0) for a in advances)

                pay_f = {}
                if cid and cid != 'ARGUS_MASTER':
                    pay_f['company_id'] = cid
                payments = list(db.payments.find(pay_f))
                tot_comp_rep = sum(float(p.get('amount') or 0.0) for p in payments if 'repayment' in str(p.get('status') or '').lower())
                net_comp_bal = max(0.0, tot_comp_adv - tot_comp_rep)

                sample_staff = ", ".join([f"**{e.get('employee_name')}**" for e in employees[:4] if e.get('employee_name')])
                return jsonify({
                    'success': True,
                    'reply': f"Here is your company-wide Advance Summary:\n"
                             f"• Total Company Advances: **₹{tot_comp_adv:,.2f}**\n"
                             f"• Total Repayments: **₹{tot_comp_rep:,.2f}**\n"
                             f"• **Net Outstanding Advances: ₹{net_comp_bal:,.2f}**\n\n"
                             f"To check an individual's balance, you can say *'Check balance for Philip Matthew'* or ask about: {sample_staff}.",
                    'task_executed': True,
                    'task_data': {
                        'type': 'balance',
                        'details': {
                            'Scope': 'Company Total',
                            'Total Advance': f"₹{tot_comp_adv:,.2f}",
                            'Total Repaid': f"₹{tot_comp_rep:,.2f}",
                            'Net Outstanding': f"₹{net_comp_bal:,.2f}"
                        },
                        'action_label': 'View Payment Ledger',
                        'action_url': '/payment-entry'
                    },
                    'pending_task': None
                })

        # 3. CREATE ADVANCE ENTRY
        elif action_name == 'create_advance':
            raw_name = params.get('employee_name')
            raw_amt = params.get('amount')
            if raw_name:
                raw_name = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', str(raw_name).strip(), flags=re.IGNORECASE).strip()

            emp_doc = None
            if raw_name:
                emp_doc, candidates = resolve_employee(raw_name)
                if not emp_doc and candidates:
                    cand_names = ", ".join([f"**{c.get('employee_name')}**" for c in candidates[:4]])
                    return jsonify({
                        'success': True,
                        'reply': f"Multiple employees matched '**{raw_name}**': {cand_names}. Which one should receive the advance?",
                        'task_executed': False,
                        'pending_task': None
                    })

            canonical_emp = emp_doc.get('employee_name') if emp_doc else raw_name
            if not canonical_emp or raw_amt is None:
                new_pending = {
                    'action': 'create_advance',
                    'parameters': {'employee_name': canonical_emp, 'amount': raw_amt}
                }
                session['coworker_pending_task'] = new_pending
                return jsonify({
                    'success': True,
                    'reply': f"Please specify the employee name and advance amount (e.g., *'Give advance of ₹2,000 to Philip Matthew'*).",
                    'task_executed': False,
                    'pending_task': new_pending
                })

            amt = float(raw_amt)
            adv_date = params.get('date') or datetime.now().strftime('%Y-%m-%d')
            reason = params.get('reason') or 'Staff Advance'

            adv_payload = {
                'employee_name': canonical_emp,
                'amount': amt,
                'advance_date': adv_date,
                'reason': reason
            }
            adv_id = database.create_advance(adv_payload, company_id=cid)
            session['coworker_pending_task'] = None

            return jsonify({
                'success': True,
                'reply': f"I've recorded an advance of **₹{amt:,.2f}** for **{canonical_emp}**. Advance ID is **#{adv_id}**.",
                'task_executed': True,
                'task_data': {
                    'type': 'advance',
                    'details': {
                        'Employee': canonical_emp,
                        'Amount': f"₹{amt:,.2f}",
                        'Date': adv_date,
                        'Entry ID': str(adv_id)
                    },
                    'action_label': 'View Advance Ledger',
                    'action_url': '/payment-entry'
                },
                'pending_task': None
            })

        # 4. CREATE PAYMENT ENTRY
        elif action_name == 'create_payment':
            raw_name = params.get('employee_name')
            raw_amt = params.get('amount')
            if raw_name:
                raw_name = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', str(raw_name).strip(), flags=re.IGNORECASE).strip()

            emp_doc = None
            if raw_name:
                emp_doc, candidates = resolve_employee(raw_name)
                if not emp_doc:
                    if candidates:
                        cand_names = ", ".join([f"**{c.get('employee_name')}**" for c in candidates[:4]])
                        return jsonify({
                            'success': True,
                            'reply': f"I found multiple employees matching '**{raw_name}**': {cand_names}. Which one would you like to record the payment for?",
                            'task_executed': False,
                            'pending_task': None
                        })
                    else:
                        sample_staff = ", ".join([f"**{e.get('employee_name')}**" for e in employees[:5] if e.get('employee_name')])
                        staff_hint = f" Active staff: {sample_staff}." if sample_staff else ""
                        return jsonify({
                            'success': True,
                            'reply': f"I couldn't find an employee named '**{raw_name}**' registered in your company.{staff_hint} Please check the name.",
                            'task_executed': False,
                            'pending_task': None
                        })

            canonical_emp_name = emp_doc.get('employee_name') if emp_doc else raw_name

            if not raw_name or raw_amt is None:
                new_pending = {
                    'action': 'create_payment',
                    'parameters': {
                        'employee_name': canonical_emp_name,
                        'status': params.get('status')
                    }
                }
                session['coworker_pending_task'] = new_pending
                return jsonify({
                    'success': True,
                    'reply': f"I'd be glad to record a payment for **{canonical_emp_name}**. Could you please specify the **amount (₹)** and payment status (e.g. *Advance Repayment*, *Salary*, *Incentive*)?",
                    'task_executed': False,
                    'pending_task': new_pending
                })

            emp_name = canonical_emp_name
            amt = float(raw_amt)
            status = params.get('status') or 'Advance Repayment'
            pay_type = params.get('payment_type') or 'UPI'
            pay_date = params.get('date') or datetime.now().strftime('%Y-%m-%d')
            reason = params.get('reason') or ''

            payment_payload = {
                'employee_name': emp_name,
                'amount': amt,
                'status': status,
                'payment_type': pay_type,
                'payment_date': pay_date,
                'reason': reason,
                'bank': ''
            }
            inserted_id = database.create_payment(payment_payload, company_id=cid)
            session['coworker_pending_task'] = None

            return jsonify({
                'success': True,
                'reply': f"I've recorded the payment entry of **₹{amt:,.2f}** for **{emp_name}** ({status}). Entry ID is **#{inserted_id}**.",
                'task_executed': True,
                'task_data': {
                    'type': 'payment',
                    'details': {
                        'Employee': emp_name,
                        'Amount': f"₹{amt:,.2f}",
                        'Status': status,
                        'Date': pay_date,
                        'Mode': pay_type,
                        'Entry ID': str(inserted_id)
                    },
                    'action_label': 'View Payment Ledger',
                    'action_url': '/payment-entry'
                },
                'pending_task': None
            })

        # 5. CHECK ATTENDANCE & LIVE STATUS
        elif action_name == 'check_attendance':
            today_str = datetime.now().strftime('%Y-%m-%d')
            raw_name = params.get('employee_name')
            if raw_name:
                emp_doc, candidates = resolve_employee(raw_name)
                if not emp_doc:
                    return jsonify({
                        'success': True,
                        'reply': f"I couldn't find an employee named '**{raw_name}**' in your registered company staff.",
                        'task_executed': False,
                        'pending_task': None
                    })
                emp_name = emp_doc.get('employee_name')
                rec = db.attendance.find_one({'employee_name': emp_name, 'date': today_str})
                if rec:
                    in_t = rec.get('in_time', 'N/A')
                    out_t = rec.get('out_time', 'Active')
                    st = rec.get('status', 'Present')
                    return jsonify({
                        'success': True,
                        'reply': f"**{emp_name}** is marked **{st}** today ({today_str}). In-time: *{in_t}*, Out-time: *{out_t}*.",
                        'task_executed': True,
                        'task_data': {
                            'type': 'attendance',
                            'details': {'Employee': emp_name, 'Date': today_str, 'Status': st, 'In Time': in_t, 'Out Time': out_t}
                        },
                        'pending_task': None
                    })
                else:
                    return jsonify({
                        'success': True,
                        'reply': f"**{emp_name}** has not punched in today yet ({today_str}). Would you like me to log a manual entry for them?",
                        'task_executed': False,
                        'pending_task': {'action': 'create_manual_attendance', 'parameters': {'employee_name': emp_name}}
                    })
            else:
                total_emps = len(employees)
                present_count = db.attendance.count_documents({'date': today_str})
                absent_count = max(0, total_emps - present_count)
                return jsonify({
                    'success': True,
                    'reply': f"Attendance Summary for today (**{today_str}**):\n"
                             f"• Present Today: **{present_count}**\n"
                             f"• Absent Today: **{absent_count}**\n"
                             f"• Total Registered Staff: **{total_emps}**",
                    'task_executed': True,
                    'task_data': {
                        'type': 'attendance',
                        'details': {'Date': today_str, 'Present': str(present_count), 'Absent': str(absent_count), 'Total Staff': str(total_emps)}
                    },
                    'action_label': 'View Live Report',
                    'action_url': '/live-report',
                    'pending_task': None
                })

        # 6. APPLY LEAVE
        elif action_name == 'apply_leave':
            raw_name = params.get('employee_name')
            if raw_name:
                raw_name = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', str(raw_name).strip(), flags=re.IGNORECASE).strip()

            emp_doc = None
            if raw_name:
                emp_doc, candidates = resolve_employee(raw_name)

            if not emp_doc:
                new_pending = {
                    'action': 'apply_leave',
                    'parameters': params
                }
                session['coworker_pending_task'] = new_pending
                return jsonify({
                    'success': True,
                    'reply': "Please specify the employee name and leave type (e.g. *'Apply casual leave for Suresh tomorrow'*).",
                    'task_executed': False,
                    'pending_task': new_pending
                })

            emp_name = emp_doc.get('employee_name')
            leave_type = params.get('leave_type') or 'Casual Leave'
            from_d = params.get('from_date') or datetime.now().strftime('%Y-%m-%d')
            to_d = params.get('to_date') or from_d
            reason = params.get('reason') or 'Personal leave'

            leave_payload = {
                'employee_name': emp_name,
                'employee_id': emp_doc.get('id'),
                'leave_type': leave_type,
                'from_date': from_d,
                'to_date': to_d,
                'reason': reason,
                'session': 'Full Day',
                'company_id': cid
            }
            database.submit_leave_request(leave_payload)
            session['coworker_pending_task'] = None

            return jsonify({
                'success': True,
                'reply': f"Leave application for **{emp_name}** ({leave_type}) from **{from_d}** to **{to_d}** has been submitted.",
                'task_executed': True,
                'task_data': {
                    'type': 'leave',
                    'details': {
                        'Employee': emp_name,
                        'Leave Type': leave_type,
                        'From': from_d,
                        'To': to_d,
                        'Reason': reason
                    },
                    'action_label': 'View Leave Portal',
                    'action_url': '/employee-portal#leave'
                },
                'pending_task': None
            })

        # 7. GET EMPLOYEE INFO & DIRECTORY
        elif action_name == 'get_employee_info':
            raw_name = params.get('employee_name')
            if raw_name:
                raw_name = re.sub(r'\s+(for|to|of|as|with|at|in|on|rs|inr)$', '', str(raw_name).strip(), flags=re.IGNORECASE).strip()
            emp_doc = None
            if raw_name:
                emp_doc, candidates = resolve_employee(raw_name)

            if not emp_doc:
                sample_staff = ", ".join([f"**{e.get('employee_name')}**" for e in employees[:5] if e.get('employee_name')])
                return jsonify({
                    'success': True,
                    'reply': f"Which employee would you like details for? Active staff: {sample_staff}.",
                    'task_executed': False,
                    'pending_task': None
                })

            e_name = emp_doc.get('employee_name')
            dept = emp_doc.get('department') or 'General'
            desig = emp_doc.get('designation') or 'Staff'
            phone = emp_doc.get('phone') or emp_doc.get('mobile') or 'Not listed'
            status_e = emp_doc.get('status') or 'Active'

            return jsonify({
                'success': True,
                'reply': f"**{e_name}** ({desig} - {dept}):\n• Phone / Mobile: **{phone}**\n• Status: **{status_e}**",
                'task_executed': True,
                'task_data': {
                    'type': 'employee',
                    'details': {
                        'Employee': e_name,
                        'Department': dept,
                        'Designation': desig,
                        'Phone': phone,
                        'Status': status_e
                    },
                    'action_label': 'View Employee Profile',
                    'action_url': '/employees'
                },
                'pending_task': None
            })

        # 8. GET SALARY INFO
        elif action_name == 'get_salary_info':
            raw_name = params.get('employee_name')
            emp_doc = None
            if raw_name:
                emp_doc, _ = resolve_employee(raw_name)
            if not emp_doc:
                return jsonify({
                    'success': True,
                    'reply': "Please specify the employee name to look up salary details.",
                    'task_executed': False,
                    'pending_task': None
                })

            e_name = emp_doc.get('employee_name')
            sal_type = emp_doc.get('salary_type') or 'monthly'
            basic = float(emp_doc.get('basic_salary') or 0.0)
            shift_hrs = emp_doc.get('shift_hours') or '08:00'

            return jsonify({
                'success': True,
                'reply': f"Salary Structure for **{e_name}**:\n• Type: **{sal_type.title()}**\n• Basic Pay: **₹{basic:,.2f}**\n• Standard Shift: **{shift_hrs} hrs**",
                'task_executed': True,
                'task_data': {
                    'type': 'salary',
                    'details': {
                        'Employee': e_name,
                        'Salary Type': sal_type.title(),
                        'Basic Pay': f"₹{basic:,.2f}",
                        'Shift Hours': shift_hrs
                    },
                    'action_label': 'View Payslip Preview',
                    'action_url': '/payslip-preview'
                },
                'pending_task': None
            })

        # 9. FIELD WORKER MODULE
        elif action_name == 'check_field_worker':
            fw_count = db.employees.count_documents({'is_field_worker': True, 'status': {'$ne': 'Terminated'}})
            return jsonify({
                'success': True,
                'reply': f"There are **{fw_count}** active field workers registered. You can review location logs, visits, and field disbursements below.",
                'task_executed': True,
                'task_data': {
                    'type': 'attendance',
                    'details': {'Field Workers': str(fw_count), 'Module': 'Field Worker Operations'}
                },
                'action_label': 'Open Field Worker Module',
                'action_url': '/field-worker-module',
                'pending_task': None
            })

        # 10. NAVIGATION
        elif action_name == 'navigate':
            target = params.get('target', '').lower()
            url_map = {
                'manual_entry': '/manual-entry',
                'payment_entry': '/payment-entry',
                'salary_report': '/salary-report',
                'payslips': '/payslip-preview',
                'field_worker': '/field-worker-module',
                'employees': '/employees',
                'leaves': '/leaves',
                'live_report': '/live-report',
                'attendance_report': '/attendance-report',
                'dashboard': '/dashboard',
                'company_profile': '/company-profile',
                'support': '/support'
            }
            dest_url = url_map.get(target, '/dashboard')
            return jsonify({
                'success': True,
                'reply': reply_text or f"Navigating to {target.replace('_', ' ').title()}...",
                'action': 'navigate',
                'task_data': {'url': dest_url},
                'pending_task': None
            })

        # 11. GENERAL CHAT
        else:
            return jsonify({
                'success': True,
                'reply': reply_text or "How can I assist you with attendance, payments, leaves, or employee operations today?",
                'task_executed': False,
                'pending_task': None
            })

    except Exception as e:
        return jsonify({'success': False, 'error': f"Co-Worker AI error: {str(e)}"}), 500



@app.route('/api/field-worker/salary-report', methods=['GET'])
def api_field_worker_salary_report():
    """Returns monthly Field Worker Salary & Activity report."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    month_year = request.args.get('month') or database.get_ist_now().strftime('%Y-%m')
    emp_name = session.get('employee_name') if session.get('role') == 'employee' else (request.args.get('employee_name') or None)
    reports = database.get_field_worker_salary_report(company_id=cid, month_year=month_year, employee_name=emp_name)
    return jsonify({'success': True, 'month_year': month_year, 'reports': reports, 'salary_reports': reports})


@app.route('/api/field-worker/payslip-pdf', methods=['GET'])
def api_field_worker_payslip_pdf():
    """Generates and downloads an individual Field Worker's branded monthly Payslip PDF supporting 8 templates and color themes."""
    if not _is_field_authenticated():
        return redirect(url_for('field_worker_login'))

    cid = _resolve_field_company_id(request.args.get('company_id'))
    month_year = request.args.get('month') or database.get_ist_now().strftime('%Y-%m')
    emp_name = session.get('employee_name') if session.get('role') == 'employee' else request.args.get('employee_name')
    if not emp_name:
        return jsonify({'error': 'employee_name is required'}), 400

    template_param = request.args.get('template') or request.args.get('template_id') or 'template_1'
    theme_param = request.args.get('theme') or 'navy'

    # Strict isolation: ensure the worker is in get_field_workers of cid
    fw_list = database.get_field_workers(company_id=cid, include_all_employees=False)
    fw_names = [w.get('employee_name') for w in fw_list if w.get('employee_name')]
    if emp_name not in fw_names:
        return jsonify({'error': f"Access Denied: '{emp_name}' is not a registered field worker of this company"}), 403

    # Generate standard payslip data + field worker salary extras
    payslip = database.get_payslip_data(emp_name, month_year, company_id=cid)
    salary_rep = database.get_field_worker_salary_report(company_id=cid, month_year=month_year, employee_name=emp_name)
    fw_extras = salary_rep[0] if salary_rep else {}

    # Merge field worker specific fields
    payslip['is_field_worker'] = True
    payslip['total_assigned_locations'] = fw_extras.get('total_assigned_locations', 0)
    payslip['completed_assignments'] = fw_extras.get('completed_assignments', 0)
    payslip['total_visits'] = fw_extras.get('total_visits', 0)
    payslip['verified_visits'] = fw_extras.get('verified_visits', 0)
    payslip['unverified_visits'] = fw_extras.get('unverified_visits', 0)
    payslip['orders_booked'] = fw_extras.get('orders_booked', 0)
    payslip['orders_value'] = fw_extras.get('orders_value', 0.0)
    payslip['total_km_traveled'] = fw_extras.get('total_km_traveled', 0.0)
    payslip['detour_alerts'] = fw_extras.get('detour_alerts', 0)
    payslip['per_visit_incentive_rate'] = fw_extras.get('per_visit_incentive_rate', 0.0)
    payslip['km_allowance_rate'] = fw_extras.get('km_allowance_rate', 5.0)
    payslip['visit_incentive_pay'] = 0.0  # User Requirement 2: Do NOT add automated visit incentives to payslip
    payslip['travel_allowance_pay'] = fw_extras.get('travel_allowance_pay', 0.0)
    payslip['travel_allowance_calculated'] = fw_extras.get('travel_allowance_calculated', fw_extras.get('petrol_allowance', 0.0))
    payslip['travel_allowance_paid'] = fw_extras.get('travel_allowance_paid', 0.0)
    payslip['travel_allowance_remaining'] = fw_extras.get('travel_allowance_remaining', fw_extras.get('travel_allowance_pay', 0.0))
    payslip['petrol_allowance'] = payslip.get('petrol_allowance', 0.0)
    payslip['petrol_allowance_calculated'] = fw_extras.get('petrol_allowance_calculated', fw_extras.get('petrol_allowance', 0.0))
    payslip['petrol_allowance_paid'] = payslip.get('petrol_allowance_paid', fw_extras.get('petrol_allowance_paid', 0.0))
    payslip['petrol_allowance_deduction'] = payslip.get('petrol_allowance_deduction', payslip.get('petrol_allowance_paid', 0.0))
    payslip['petrol_allowance_remaining'] = fw_extras.get('petrol_allowance_remaining', 0.0)
    payslip['petrol_allowance_status'] = fw_extras.get('petrol_allowance_status', 'UNPAID')
    payslip['base_attendance_pay'] = payslip.get('basic_salary', 0.0)
    payslip['gross_field_salary'] = payslip.get('total_earnings', 0.0)
    payslip['net_field_payable'] = payslip.get('net_pay', 0.0)
    payslip['net_pay_in_words'] = payslip.get('net_pay_in_words', '')

    # Generate PDF with requested template and theme
    pdf_buf = pdf_generator.generate_payslip_pdf(payslip, template_id=template_param, theme=theme_param)
    safe_name = re.sub(r'[^A-Za-z0-9_-]', '_', emp_name)
    view_inline = (request.args.get('view') == '1' or request.args.get('inline') == '1')
    return send_file(
        pdf_buf,
        mimetype='application/pdf',
        as_attachment=not view_inline,
        download_name=f"Field_Worker_Payslip_{safe_name}_{month_year}.pdf"
    )


@app.route('/api/field-worker/export', methods=['GET'])
def api_field_worker_export():
    """Exports Field Worker Visits/Leads, Live Route Tracking, or Salary Reports to PDF or Excel."""
    if not _is_field_authenticated():
        return redirect(url_for('field_worker_login'))

    cid = _resolve_field_company_id(request.args.get('company_id'))
    report_type = (request.args.get('type') or 'visits').strip().lower()
    export_format = (request.args.get('format') or 'pdf').strip().lower()
    date_str = request.args.get('date') or database.get_ist_now().strftime('%Y-%m-%d')
    month_year = request.args.get('month') or database.get_ist_now().strftime('%Y-%m')
    emp_name = session.get('employee_name') if session.get('role') == 'employee' else (request.args.get('employee_name') or None)
    comp_info = database.get_company_by_id(cid)

    if report_type == 'visits':
        visits = database.get_field_customer_visits(company_id=cid, employee_name=emp_name)
        headers = ['Date & Time', 'Field Worker', 'Customer / Shop', 'Contact Person', 'Phone', 'Product Pitched', 'Lead Status', 'Order Value (Rs.)', 'Geo-Verification', 'Location Address']
        rows = [
            [
                v.get('timestamp', ''),
                v.get('employee_name', ''),
                v.get('customer_name', ''),
                v.get('contact_person', ''),
                v.get('customer_phone', ''),
                v.get('product_pitched', ''),
                v.get('lead_status', ''),
                f"{float(v.get('order_amount', 0)):.2f}",
                v.get('verification_badge', ''),
                v.get('address', '')
            ]
            for v in visits
        ]
        title = "FIELD WORKER CUSTOMER VISITS & SALES LEADS REPORT"
        fname = f"Field_Customer_Visits_{date_str}"
    elif report_type == 'tracking':
        tracking = database.get_field_live_tracking(company_id=cid, date_str=date_str, employee_name=emp_name)
        logs = tracking.get('logs', [])
        headers = ['Timestamp', 'Field Worker', 'Assigned Target', 'Distance (m)', 'Step (KM)', 'Route Status', 'Live Street Address']
        rows = [
            [
                lg.get('timestamp', ''),
                lg.get('employee_name', ''),
                lg.get('nearest_target_name', ''),
                f"{float(lg.get('distance_from_target_m', 0)):.1f}m",
                f"{float(lg.get('step_km', 0)):.3f} km",
                lg.get('status_label', ''),
                lg.get('address', '')
            ]
            for lg in logs
        ]
        title = f"FIELD WORKER LIVE ROUTE & DETOUR TRACKING LOG ({date_str})"
        fname = f"Field_Route_Tracking_{date_str}"
    else:
        reports = database.get_field_worker_salary_report(company_id=cid, month_year=month_year, employee_name=emp_name)
        headers = ['Field Worker', 'Zone', 'Days / Hrs', 'Verified Visits', 'Distance (KM)', 'Detours', 'Base Pay', 'Visit Incentive', 'KM Allowance', 'Net Payable (Rs.)']
        rows = [
            [
                r.get('employee_name', ''),
                r.get('territory_zone', ''),
                f"{r.get('working_days', 0)}d / {r.get('total_hours', '00:00')}h",
                f"{r.get('verified_visits', 0)} / {r.get('total_visits', 0)}",
                f"{float(r.get('total_km_traveled', 0)):.2f} km",
                str(r.get('detour_alerts', 0)),
                f"{float(r.get('base_attendance_pay', 0)):.2f}",
                f"{float(r.get('visit_incentive_pay', 0)):.2f}",
                f"{float(r.get('travel_allowance_pay', 0)):.2f}",
                f"{float(r.get('net_field_payable', 0)):.2f}"
            ]
            for r in reports
        ]
        title = f"FIELD WORKER MONTHLY SALARY & INCENTIVE REPORT ({month_year})"
        fname = f"Field_Salary_Report_{month_year}"

    if export_format in ['excel', 'xlsx', 'csv']:
        return export_table_data(headers, rows, fname, export_format=export_format, sheet_name="Field Report")

    pdf_buf = pdf_generator.generate_field_worker_report_pdf(
        headers, rows, title=title, subtitle=f"Generated on {database.get_ist_now().strftime('%d/%m/%Y %I:%M %p')}", company_info=comp_info
    )
    return send_file(pdf_buf, mimetype='application/pdf', as_attachment=True, download_name=f"{fname}.pdf")


@app.route('/api/field-worker/dashboard-stats', methods=['GET'])
def api_field_worker_dashboard_stats():
    """Returns real-time operations dashboard stats for the Field Worker module."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    date_str = request.args.get('date') or database.get_ist_now().strftime('%Y-%m-%d')
    stats = database.get_field_worker_dashboard_stats(company_id=cid, date_str=date_str)
    return jsonify({'success': True, 'stats': stats, 'company_id': cid, 'date': date_str})


@app.route('/api/field-worker/leave-requests', methods=['GET'])
def api_field_worker_leave_requests():
    """Returns leave and permission requests isolated for field workers."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    user_role = session.get('role')
    status = request.args.get('status', 'ALL')
    search = request.args.get('search', '').strip()
    limit = int(request.args.get('limit', 50))
    page = int(request.args.get('page', 1))

    # If employee logged in, only fetch their own requests
    emp_id = None
    if user_role == 'employee':
        emp_id = session.get('employee_id')
    else:
        req_emp = request.args.get('employee_id')
        if req_emp:
            emp_id = req_emp

    res = database.get_leave_requests(
        company_id=cid,
        employee_id=emp_id,
        status=status,
        limit=limit,
        page=page,
        for_field_workers=True,
        search=search
    )
    return jsonify({
        'success': True,
        'requests': res.get('requests', []),
        'total': res.get('total', 0),
        'page': res.get('page', 1),
        'limit': res.get('limit', limit),
        'status_counts': res.get('status_counts', {'all': 0, 'pending': 0, 'approved': 0, 'rejected': 0})
    })


@app.route('/api/field-worker/leave-requests/apply', methods=['POST'])
def api_field_worker_apply_leave():
    """Submits a leave request for a field worker with strict field worker isolation."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    data = request.form.to_dict() or request.get_json(silent=True) or {}
    user_role = session.get('role')

    emp_id = session.get('employee_id') if user_role == 'employee' else (data.get('employee_id') or data.get('existing_employee_id'))
    emp_name = session.get('employee_name') if user_role == 'employee' else (data.get('employee_name') or '')

    # If admin applying on behalf of a field worker
    if not emp_id and emp_name:
        emp_doc = database.get_db().employees.find_one({'employee_name': emp_name})
        if emp_doc:
            emp_id = str(emp_doc.get('id') or emp_doc.get('_id'))

    emp = database.get_employee_by_id(emp_id) if emp_id else None
    if not emp and emp_name:
        emp = database.get_db().employees.find_one({'employee_name': emp_name})
    if not emp:
        return jsonify({'success': False, 'error': 'Field worker profile not found'}), 404

    leave_type = data.get('leave_type', 'Casual Leave (CL)').strip()
    from_date = data.get('from_date', '').strip()
    to_date = data.get('to_date', '').strip() or from_date
    session_type = data.get('session', 'Full Day').strip()
    reason = data.get('reason', '').strip()

    if not leave_type or not from_date or not reason:
        return jsonify({'success': False, 'error': 'Please fill in all required fields (Leave Type, Dates, Reason).'}), 400

    cid = _resolve_field_company_id(data.get('company_id') or emp.get('company_id'))

    attachments = []
    if 'attachment' in request.files:
        file = request.files['attachment']
        if file and file.filename:
            fname = werkzeug.utils.secure_filename(file.filename) or 'fw_leave_doc'
            base_name, ext = os.path.splitext(fname)
            unique_fname = f"fw_{base_name}_{int(time.time())}{ext}"
            os.makedirs(LEAVES_FOLDER, exist_ok=True)
            dest_path = os.path.join(LEAVES_FOLDER, unique_fname)
            file.save(dest_path)
            size_kb = round(os.path.getsize(dest_path) / 1024, 1)
            attachments.append({
                'filename': unique_fname,
                'original_name': fname,
                'file_url': f"/uploads/leaves/{unique_fname}",
                'size_str': f"{size_kb} KB",
                'mime_type': file.mimetype or 'application/octet-stream',
                'uploaded_at': database.get_ist_now().strftime('%d %b %Y %I:%M %p')
            })

    req_doc = {
        'company_id': cid,
        'employee_id': str(emp.get('id') or emp.get('_id')),
        'employee_code': str(emp.get('id') or emp.get('employee_id') or emp.get('emp_id') or ''),
        'employee_name': emp.get('employee_name', ''),
        'department': emp.get('department', 'Field Sales & Operations'),
        'designation': emp.get('designation', 'Field Executive'),
        'mobile_number': emp.get('mobile_number', ''),
        'branch_id': emp.get('branch_id'),
        'branch_name': emp.get('branch_name') or emp.get('branch') or 'Home Branch (Head Office)',
        'leave_type': leave_type,
        'from_date': from_date,
        'to_date': to_date,
        'session': session_type,
        'reason': reason,
        'attachments': attachments,
        'is_field_worker': True,
        'employee_type': 'field_worker'
    }

    new_req = database.submit_leave_request(req_doc)
    return jsonify({
        'success': True,
        'message': f'Leave request for {emp.get("employee_name")} submitted successfully!',
        'request': new_req
    })


@app.route('/api/field-worker/permission-requests/apply', methods=['POST'])
def api_field_worker_apply_permission():
    """Submits an hourly permission request for a field worker."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    data = request.form.to_dict() or request.get_json(silent=True) or {}
    user_role = session.get('role')

    emp_id = session.get('employee_id') if user_role == 'employee' else (data.get('employee_id') or data.get('existing_employee_id'))
    emp_name = session.get('employee_name') if user_role == 'employee' else (data.get('employee_name') or '')

    if not emp_id and emp_name:
        emp_doc = database.get_db().employees.find_one({'employee_name': emp_name})
        if emp_doc:
            emp_id = str(emp_doc.get('id') or emp_doc.get('_id'))

    emp = database.get_employee_by_id(emp_id) if emp_id else None
    if not emp and emp_name:
        emp = database.get_db().employees.find_one({'employee_name': emp_name})
    if not emp:
        return jsonify({'success': False, 'error': 'Field worker profile not found'}), 404

    perm_type = data.get('permission_type', 'Short Leave').strip()
    perm_date = data.get('date', '').strip()
    from_time = data.get('from_time', '').strip()
    to_time = data.get('to_time', '').strip()
    duration = data.get('duration', '02:00').strip()
    reason = data.get('reason', '').strip()

    if not perm_type or not perm_date or not from_time or not to_time or not reason:
        return jsonify({'success': False, 'error': 'Please fill in all required fields (Date, Times, Reason).'}), 400

    cid = _resolve_field_company_id(data.get('company_id') or emp.get('company_id'))

    req_doc = {
        'company_id': cid,
        'employee_id': str(emp.get('id') or emp.get('_id')),
        'employee_code': str(emp.get('id') or emp.get('employee_id') or emp.get('emp_id') or ''),
        'employee_name': emp.get('employee_name', ''),
        'department': emp.get('department', 'Field Sales & Operations'),
        'designation': emp.get('designation', 'Field Executive'),
        'mobile_number': emp.get('mobile_number', ''),
        'branch_id': emp.get('branch_id'),
        'branch_name': emp.get('branch_name') or emp.get('branch') or 'Home Branch (Head Office)',
        'permission_type': perm_type,
        'date': perm_date,
        'from_time': from_time,
        'to_time': to_time,
        'duration': duration,
        'reason': reason,
        'attachments': [],
        'is_field_worker': True,
        'employee_type': 'field_worker'
    }

    new_req = database.submit_permission_request(req_doc)
    return jsonify({
        'success': True,
        'message': f'Permission request for {emp.get("employee_name")} submitted successfully!',
        'request': new_req
    })


@app.route('/api/field-worker/leave-requests/<request_id>/action', methods=['POST'])
def api_field_worker_leave_request_action(request_id):
    """Admin approves or rejects a field worker's leave or permission request."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    if session.get('role') == 'employee':
        return jsonify({'success': False, 'error': 'Unauthorized: Only Company Administrators can approve or reject leaves.'}), 403

    data = request.get_json(silent=True) or request.form.to_dict() or {}
    action = str(data.get('action') or data.get('status') or '').strip()
    if action.lower() in ['approve', 'approved']:
        action = 'Approved'
    elif action.lower() in ['reject', 'rejected']:
        action = 'Rejected'
    else:
        return jsonify({'success': False, 'error': "Action must be 'Approved' or 'Rejected'"}), 400

    remark = data.get('admin_remark', '').strip()
    reviewer = session.get('user') or session.get('admin_username') or 'Field Operations Admin'
    res = database.update_leave_request_status(request_id, action, admin_remark=remark, reviewer=reviewer)
    return jsonify(res)


@app.route('/api/field-worker/leave-balances', methods=['GET'])
def api_field_worker_leave_balances():
    """Returns annual leave quotas and live balances for field workers."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    user_role = session.get('role')

    if user_role == 'employee':
        emp_id = session.get('employee_id')
        bal = database.get_employee_leave_balance(emp_id, company_id=cid)
        return jsonify({'success': True, 'balance': bal, 'balances': [bal] if bal else []})

    search = request.args.get('search', '').strip()
    bals = database.get_all_employees_leave_balances(company_id=cid, search=search, for_field_workers=True)
    return jsonify({'success': True, 'balances': bals})


# ------------------------------------------------------------------------------
# FIELD WORKER MODULE: ISOLATED MANUAL ENTRY, PAYMENT ENTRY & PAYSLIP PREVIEW
# ------------------------------------------------------------------------------

@app.route('/api/field-worker/manual-entries', methods=['GET'])
def api_field_worker_get_manual_entries():
    """Returns manual entries filtered strictly to field workers of the active company."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401
    
    cid = _resolve_field_company_id(request.args.get('company_id'))
    from_date = request.args.get('from_date', '').strip()
    to_date = request.args.get('to_date', '').strip()
    status = request.args.get('status', 'All').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))

    user_role = session.get('role')
    if user_role == 'employee':
        emp_name = session.get('employee_name', '')
        search = emp_name

    res = database.get_field_worker_manual_entries(
        company_id=cid,
        from_date=from_date,
        to_date=to_date,
        status=status,
        search=search,
        page=page,
        limit=limit
    )
    res['success'] = True
    return jsonify(res)


@app.route('/api/field-worker/manual-entries', methods=['POST'])
def api_field_worker_create_manual_entry():
    """Creates a manual entry strictly for a registered field worker."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401
    if session.get('role') not in ['super_admin', 'company_admin', 'admin']:
        return jsonify({'success': False, 'error': 'Admin privileges required'}), 403

    data = request.get_json(silent=True) or request.form.to_dict() or {}
    cid = _resolve_field_company_id(data.get('company_id'))

    try:
        inserted_id = database.save_field_worker_manual_entry(data, company_id=cid)
        return jsonify({
            'success': True,
            'id': inserted_id,
            'message': 'Field worker manual entry saved successfully'
        })
    except ValueError as ve:
        return jsonify({'success': False, 'error': str(ve)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': f'Failed to save manual entry: {str(e)}'}), 500


@app.route('/api/field-worker/manual-entries/<entry_id>', methods=['DELETE'])
def api_field_worker_delete_manual_entry(entry_id):
    """Deletes a field worker manual entry with strict company and identity verification."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401
    if session.get('role') not in ['super_admin', 'company_admin', 'admin']:
        return jsonify({'success': False, 'error': 'Admin privileges required'}), 403

    cid = _resolve_field_company_id()
    try:
        ok = database.delete_field_worker_manual_entry(entry_id, company_id=cid)
        if ok:
            return jsonify({'success': True, 'message': 'Field worker manual entry deleted successfully'})
        return jsonify({'success': False, 'error': 'Entry not found or unauthorized'}), 404
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/field-worker/payments', methods=['GET'])
def api_field_worker_get_payments():
    """Returns payment records filtered strictly to field workers of the active company."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    employee = request.args.get('employee', 'All').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()
    bank = request.args.get('bank', 'All').strip()
    payment_type = request.args.get('payment_type', 'All').strip()
    status = (request.args.get('status') or 'All').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))

    user_role = session.get('role')
    if user_role == 'employee':
        employee = session.get('employee_name', '')

    res = database.get_field_worker_payments(
        company_id=cid,
        employee=employee,
        start_date=start_date,
        end_date=end_date,
        bank=bank,
        payment_type=payment_type,
        status=status,
        search=search,
        page=page,
        limit=limit
    )
    res['success'] = True
    return jsonify(res)


@app.route('/api/field-worker/payment-stats', methods=['GET'])
def api_field_worker_get_payment_stats():
    """Returns aggregated payment and advance statistics strictly for field workers."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    month_year = request.args.get('month', '').strip() or None
    stats = database.get_field_worker_payment_stats(company_id=cid, month_year=month_year)
    return jsonify({'success': True, 'stats': stats})


@app.route('/api/field-worker/payments', methods=['POST'])
def api_field_worker_create_payment():
    """Creates a payment entry strictly for a field worker."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401
    if session.get('role') not in ['super_admin', 'company_admin', 'admin']:
        return jsonify({'success': False, 'error': 'Admin privileges required'}), 403

    data = {}
    if request.is_json:
        data = request.get_json(silent=True) or {}
    else:
        data = request.form.to_dict()

    if 'receipt' in request.files:
        file = request.files['receipt']
        if file and file.filename:
            fname = werkzeug.utils.secure_filename(file.filename)
            uniq = f"{int(database.time.time())}_{fname}"
            file.save(os.path.join(RECEIPTS_FOLDER, uniq))
            data['receipt_filename'] = uniq

    cid = _resolve_field_company_id(data.get('company_id'))
    try:
        inserted_id = database.save_field_worker_payment(data, company_id=cid)
        return jsonify({
            'success': True,
            'id': inserted_id,
            'message': 'Field worker payment recorded successfully'
        })
    except ValueError as ve:
        return jsonify({'success': False, 'error': str(ve)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': f'Failed to record payment: {str(e)}'}), 500


@app.route('/api/field-worker/advances', methods=['POST'])
def api_field_worker_create_advance():
    """Creates a salary advance record strictly for a field worker."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401
    if session.get('role') not in ['super_admin', 'company_admin', 'admin']:
        return jsonify({'success': False, 'error': 'Admin privileges required'}), 403

    data = request.get_json(silent=True) or request.form.to_dict() or {}
    cid = _resolve_field_company_id(data.get('company_id'))
    try:
        adv = database.save_field_worker_advance(data, company_id=cid)
        return jsonify({
            'success': True,
            'advance': adv,
            'message': 'Field worker advance recorded successfully'
        })
    except ValueError as ve:
        return jsonify({'success': False, 'error': str(ve)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': f'Failed to record advance: {str(e)}'}), 500


@app.route('/api/field-worker/payments/<payment_id>', methods=['DELETE'])
def api_field_worker_delete_payment(payment_id):
    """Deletes a field worker payment entry with tenant and worker isolation."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401
    if session.get('role') not in ['super_admin', 'company_admin', 'admin']:
        return jsonify({'success': False, 'error': 'Admin privileges required'}), 403

    cid = _resolve_field_company_id()
    try:
        ok = database.delete_field_worker_payment(payment_id, company_id=cid)
        if ok:
            return jsonify({'success': True, 'message': 'Payment record deleted successfully'})
        return jsonify({'success': False, 'error': 'Payment not found or unauthorized'}), 404
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/field-worker/balance-report', methods=['GET'])
def api_field_worker_get_balance_report():
    """Returns advance and repayment running ledger strictly isolated for field workers."""
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    employee = request.args.get('employee', 'All').strip()
    search = request.args.get('search', '').strip()
    page = int(request.args.get('page', 1))
    limit = int(request.args.get('limit', 10))

    user_role = session.get('role')
    if user_role == 'employee':
        employee = session.get('employee_name', '')

    res = database.get_field_worker_balance_report(
        company_id=cid,
        employee_name=employee,
        search=search,
        page=page,
        limit=limit
    )
    res['success'] = True
    return jsonify(res)


@app.route('/api/field-worker/payslip-preview-data', methods=['GET'])
def api_field_worker_get_payslip_preview_data():
    """
    Returns complete, isolated payslip preview data for a Field Worker.
    Guarantees standard employees are rejected with 403/404.
    """
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    cid = _resolve_field_company_id(request.args.get('company_id'))
    employee = (request.args.get('employee_name') or request.args.get('employee') or '').strip()
    month = request.args.get('month', '').strip()
    if not month:
        month = database.get_ist_now().strftime("%Y-%m")

    # Strict isolation: ensure the worker is in get_field_workers of cid
    fw_list = database.get_field_workers(company_id=cid, include_all_employees=False)
    fw_names = [w.get('employee_name') for w in fw_list if w.get('employee_name')]

    user_role = session.get('role')
    if user_role == 'employee':
        employee = session.get('employee_name', '')

    if not employee:
        if fw_names:
            employee = fw_names[0]
        else:
            return jsonify({'success': False, 'error': 'No registered field workers found for this company'}), 404

    if employee not in fw_names:
        return jsonify({'success': False, 'error': f"Access Denied: '{employee}' is not a registered field worker of this company"}), 403

    # Generate standard payslip data + field worker salary extras
    payslip = database.get_payslip_data(employee, month, company_id=cid)
    salary_rep = database.get_field_worker_salary_report(company_id=cid, month_year=month, employee_name=employee)
    fw_extras = salary_rep[0] if salary_rep else {}

    # Merge field worker specific fields into preview data
    payslip['is_field_worker'] = True
    payslip['total_assigned_locations'] = fw_extras.get('total_assigned_locations', 0)
    payslip['completed_assignments'] = fw_extras.get('completed_assignments', 0)
    payslip['total_visits'] = fw_extras.get('total_visits', 0)
    payslip['verified_visits'] = fw_extras.get('verified_visits', 0)
    payslip['unverified_visits'] = fw_extras.get('unverified_visits', 0)
    payslip['orders_booked'] = fw_extras.get('orders_booked', 0)
    payslip['orders_value'] = fw_extras.get('orders_value', 0.0)
    payslip['total_km_traveled'] = fw_extras.get('total_km_traveled', 0.0)
    payslip['detour_alerts'] = fw_extras.get('detour_alerts', 0)
    payslip['per_visit_incentive_rate'] = fw_extras.get('per_visit_incentive_rate', 0.0)
    payslip['km_allowance_rate'] = fw_extras.get('km_allowance_rate', 5.0)
    payslip['visit_incentive_pay'] = fw_extras.get('visit_incentive_pay', 0.0)
    payslip['verified_visits_incentive'] = fw_extras.get('verified_visits_incentive', fw_extras.get('visit_incentive_pay', 0.0))
    payslip['travel_allowance_pay'] = fw_extras.get('travel_allowance_pay', 0.0)
    payslip['travel_allowance_calculated'] = fw_extras.get('travel_allowance_calculated', fw_extras.get('petrol_allowance', 0.0))
    payslip['travel_allowance_paid'] = fw_extras.get('travel_allowance_paid', 0.0)
    payslip['travel_allowance_remaining'] = fw_extras.get('travel_allowance_remaining', fw_extras.get('travel_allowance_pay', 0.0))
    payslip['petrol_allowance'] = fw_extras.get('petrol_allowance', payslip.get('petrol_allowance', 0.0))
    payslip['petrol_allowance_calculated'] = fw_extras.get('petrol_allowance_calculated', fw_extras.get('petrol_allowance', 0.0))
    payslip['petrol_allowance_paid'] = fw_extras.get('petrol_allowance_paid', 0.0)
    payslip['petrol_allowance_remaining'] = fw_extras.get('petrol_allowance_remaining', 0.0)
    payslip['petrol_allowance_status'] = fw_extras.get('petrol_allowance_status', 'UNPAID')
    payslip['base_attendance_pay'] = fw_extras.get('base_attendance_pay', payslip.get('basic_salary', 0.0))
    payslip['gross_field_salary'] = fw_extras.get('gross_field_salary', payslip.get('total_earnings', 0.0))
    payslip['net_field_payable'] = fw_extras.get('net_field_payable', payslip.get('net_pay', 0.0))
    payslip['net_pay_in_words'] = fw_extras.get('net_pay_in_words', payslip.get('net_pay_in_words', ''))

    return jsonify({'success': True, 'payslip': payslip, 'field_extras': fw_extras})


@app.route('/api/field-worker/petrol-rate-config', methods=['GET', 'POST'])
def api_field_worker_petrol_rate_config():
    """
    GET: Retrieves current petrol/diesel allowance rate configuration for the company.
    POST: Updates petrol/diesel allowance rate configuration and syncs across all field workers.
    """
    if not _is_field_authenticated():
        return jsonify({'success': False, 'error': 'Authentication required'}), 401

    user_role = session.get('role', '')
    req_cid = request.args.get('company_id') if request.method == 'GET' else (request.get_json(silent=True) or {}).get('company_id')
    cid = _resolve_field_company_id(req_cid)

    if request.method == 'GET':
        config = database.get_company_petrol_allowance_config(cid)
        return jsonify({'success': True, **config, 'company_id': str(cid)})

    # POST - Requires admin
    if user_role not in ['super_admin', 'company_admin', 'admin']:
        return jsonify({'success': False, 'error': 'Unauthorized: Admin privileges required to edit petrol configuration'}), 403

    payload = request.get_json(silent=True) or {}
    rate_inr = payload.get('rate_inr')

    if rate_inr is None:
        return jsonify({'success': False, 'error': 'Rate in Rupees is required'}), 400

    try:
        unit_km = 1.0  # Fixed reference unit distance
        rate_inr = float(rate_inr)
        if rate_inr < 0:
            raise ValueError()
    except (ValueError, TypeError):
        return jsonify({'success': False, 'error': 'Invalid numeric value for rate in rupees'}), 400

    result = database.set_company_petrol_allowance_config(cid, 1.0, rate_inr)
    return jsonify({
        'success': True,
        'message': f"Petrol/Diesel allowance rate updated to Rs. {result['per_km_rate']:.2f} per {result['unit_km']} KM successfully!",
        **result,
        'company_id': str(cid)
    })



# ==============================================================================
# MULTI-BRANCH ENTERPRISE MODULE (DISTRICT HUBS & EMPLOYEE ALLOCATION)
# ==============================================================================

def _is_branch_admin_authenticated():
    """Checks if the current session is an authenticated admin (Super Admin, Company Admin, or Admin)."""
    if not (session.get('admin_logged_in') or session.get('logged_in')):
        return False
    role = session.get('role', 'company_admin')
    return role in ['super_admin', 'company_admin', 'admin']


@app.route('/manage-branches')
@app.route('/branches')
def branch_management_page():
    """Renders the Multi-Branch Enterprise Hub UI with district mapping & employee allocation."""
    if not (session.get('admin_logged_in') or session.get('logged_in')):
        return redirect(url_for('login'))

    user_role = session.get('role', 'company_admin')
    if user_role not in ['super_admin', 'company_admin', 'admin']:
        flash('Admin authorization required to access the Multi-Branch Hub.', 'warning')
        return redirect(url_for('dashboard'))

    active_cid = _resolve_field_company_id(request.args.get('company_id'))

    # Check subscription entitlement
    if user_role not in ['super_admin', 'admin']:
        comp_modules = database.get_company_modules(active_cid)
        if not comp_modules.get('branch_company', True):
            flash('The Multi-Branch Module is not authorized for your company subscription. Please contact your administrator.', 'warning')
            return redirect(url_for('dashboard'))

    active_comp = database.get_company_by_id(active_cid) or {
        'id': 'ARGUS_MASTER',
        'company_name': 'ARGUS TECHNOLOGIES',
        'latitude': 11.02980,
        'longitude': 76.97400,
        'address': 'SF NO. 515, Bharathiyar Road, Ganapathy, Coimbatore'
    }

    comp_res = database.get_all_companies(limit=1000) if user_role == 'super_admin' else None
    companies = comp_res.get('data', []) if isinstance(comp_res, dict) else (comp_res or [])
    company_map = {str(c['id']): c.get('company_name', 'Company') for c in companies}
    company_map['ARGUS_MASTER'] = 'ARGUS TECHNOLOGIES'

    req_cid = request.args.get('company_id')
    if user_role == 'super_admin':
        if req_cid and req_cid not in ['ALL', 'all', 'None', 'null']:
            active_cid = str(req_cid).strip()
            branches = database.get_company_branches(active_cid)
        else:
            active_cid = 'ALL'
            branches = database.get_company_branches('ALL')

        for b in branches:
            cid = str(b.get('company_id', ''))
            b['company_name'] = company_map.get(cid, 'ARGUS TECHNOLOGIES')

        # Pre-sort by company name then branch name so Jinja2 groupby functions correctly without disjoint chunks
        branches.sort(key=lambda b: (b.get('company_name', '').lower(), b.get('branch_name', '').lower()))
    else:
        branches = database.get_company_branches(active_cid)
        for b in branches:
            b['company_name'] = active_comp.get('company_name', 'Company')

    return render_template(
        'branch_management_module.html',
        branches=branches,
        active_company=active_comp,
        active_company_id=active_cid,
        companies=companies,
        user_role=user_role,
        total_branches=len(branches),
        active_tab='MANAGE BRANCHES'
    )


@app.route('/branches/kiosk')
def branch_kiosk_page():
    """Dedicated Branch Attendance Kiosk with branch branding & geofence tracking."""
    active_cid = _resolve_field_company_id(request.args.get('company_id'))
    active_bid = request.args.get('branch_id')
    active_branch = None
    if active_bid:
        active_branch = database.get_branch_by_id(active_bid, active_cid)

    active_comp = database.get_company_by_id(active_cid) or {
        'id': 'ARGUS_MASTER',
        'company_name': 'ARGUS TECHNOLOGIES',
        'latitude': 11.02980,
        'longitude': 76.97400,
        'address': 'SF NO. 515, Bharathiyar Road, Ganapathy, Coimbatore'
    }
    branches = database.get_company_branches(active_cid)
    return render_template(
        'attendance.html',
        is_branch_kiosk=True,
        active_branch=active_branch,
        active_branch_id=active_bid,
        branches=branches,
        active_company=active_comp,
        active_company_id=active_cid
    )


@app.route('/api/branches', methods=['GET'])
def api_get_branches():
    """Returns JSON list of branches for the active company."""
    if not _is_branch_admin_authenticated():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    active_cid = _resolve_field_company_id(request.args.get('company_id'))
    branches = database.get_company_branches(active_cid)
    return jsonify({'success': True, 'company_id': active_cid, 'branches': branches})


@app.route('/api/branches', methods=['POST'])
def api_create_branch():
    """Creates a new branch for the active company."""
    if not _is_branch_admin_authenticated():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    active_cid = _resolve_field_company_id(request.args.get('company_id'))
    data = request.get_json() or {}
    if session.get('role') == 'super_admin' and data.get('company_id'):
        active_cid = str(data.get('company_id')).strip()

    b_name = (data.get('branch_name') or '').strip()
    if not b_name:
        return jsonify({'success': False, 'error': 'Branch name is required'}), 400

    try:
        branch_id = database.create_company_branch(data, active_cid)
        return jsonify({'success': True, 'branch_id': branch_id, 'message': f'Branch "{b_name}" registered successfully!'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/branches/<branch_id>', methods=['PUT'])
def api_update_branch(branch_id):
    """Updates an existing branch location."""
    if not _is_branch_admin_authenticated():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    active_cid = _resolve_field_company_id(request.args.get('company_id'))
    data = request.get_json() or {}
    if session.get('role') == 'super_admin':
        if data.get('company_id'):
            active_cid = str(data.get('company_id')).strip()
        else:
            b_doc = database.get_branch_by_id(branch_id, company_id='ALL')
            if b_doc and b_doc.get('company_id'):
                active_cid = str(b_doc.get('company_id'))

    try:
        success = database.update_company_branch(branch_id, data, active_cid)
        return jsonify({'success': True, 'message': 'Branch updated successfully'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/branches/<branch_id>', methods=['DELETE'])
def api_delete_branch(branch_id):
    """Deletes a branch and safely unassigns its employees."""
    if not _is_branch_admin_authenticated():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    active_cid = _resolve_field_company_id(request.args.get('company_id'))
    if session.get('role') == 'super_admin':
        req_cid = request.args.get('company_id')
        if req_cid and req_cid not in ['ALL', 'all']:
            active_cid = str(req_cid).strip()
        else:
            b_doc = database.get_branch_by_id(branch_id, company_id='ALL')
            if b_doc and b_doc.get('company_id'):
                active_cid = str(b_doc.get('company_id'))
    try:
        success = database.delete_company_branch(branch_id, active_cid)
        return jsonify({'success': True, 'message': 'Branch deleted and employees unassigned successfully'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/branches/assign-employee', methods=['POST'])
def api_assign_employee_branch():
    """Assigns one or more employees to a branch (or unassigns if branch_id is empty)."""
    if not _is_branch_admin_authenticated():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    active_cid = _resolve_field_company_id(request.args.get('company_id'))
    data = request.get_json() or {}

    branch_id = data.get('branch_id')
    emp_ids = data.get('employee_ids') or []
    single_emp_id = data.get('employee_id')
    if single_emp_id and single_emp_id not in emp_ids:
        emp_ids.append(single_emp_id)

    if not emp_ids:
        return jsonify({'success': False, 'error': 'No employee selected for assignment'}), 400

    try:
        for eid in emp_ids:
            database.assign_employee_to_branch(eid, branch_id, active_cid)
        return jsonify({'success': True, 'message': f'Successfully updated branch assignment for {len(emp_ids)} employee(s)'})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/branches/stats', methods=['GET'])
def api_branch_stats():
    """Returns attendance summary, optionally filtered by branch_id."""
    if not _is_branch_admin_authenticated():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    active_cid = _resolve_field_company_id(request.args.get('company_id'))
    branch_id = request.args.get('branch_id')
    stats = database.get_dashboard_stats(active_cid, branch_id=branch_id)
    summary = database.get_branch_attendance_summary(active_cid)
    return jsonify({'success': True, 'stats': stats, 'data': summary})


@app.route('/api/branches/set-active', methods=['POST'])
def api_set_active_branch():
    """Sets or clears the active branch filter in session for top navigation."""
    if not _is_branch_admin_authenticated():
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    bid = data.get('branch_id')
    if not bid or str(bid).strip().upper() in ['ALL', 'NONE', '']:
        session.pop('active_branch_id', None)
        session.pop('active_branch_name', None)
        return jsonify({'success': True, 'active_branch_id': None, 'active_branch_name': 'All Branches'})

    if str(bid).strip().upper() in ['HOME', 'HQ', 'MAIN', 'UNASSIGNED']:
        session['active_branch_id'] = 'HOME'
        session['active_branch_name'] = 'Home Branch'
        session.modified = True
        return jsonify({'success': True, 'active_branch_id': 'HOME', 'active_branch_name': 'Home Branch'})

    if session.get('role') == 'super_admin':
        branch = database.get_branch_by_id(bid, company_id='ALL')
    else:
        active_cid = _resolve_field_company_id(None)
        branch = database.get_branch_by_id(bid, active_cid)

    if branch:
        session['active_branch_id'] = str(bid)
        session['active_branch_name'] = branch.get('branch_name')
        session.modified = True
        return jsonify({'success': True, 'active_branch_id': str(bid), 'active_branch_name': branch.get('branch_name')})

    return jsonify({'success': False, 'error': 'Branch not found'}), 404




if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)


